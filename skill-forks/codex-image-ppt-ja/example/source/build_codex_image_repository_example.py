"""Prepare/record native image jobs and call ningzimu's unmodified packager.

Image generation is performed by the host's image_gen.imagegen tool, never by
this script. Every final slide is the original raster returned by that tool.
"""
from pathlib import Path
from datetime import datetime, timezone
import argparse, hashlib, json, shutil, subprocess

ROOT=Path(__file__).resolve().parents[1]
UP=ROOT/'.build/actual-upstreams/ningzimu__codex-ppt-skill'
SKILL=UP/'skills/codex-ppt'
OUT=ROOT/'examples/repositories/codex-image-ppt'
PY=ROOT/'.build/ppt-master-venv/Scripts/python.exe'
BACKEND='built-in image_gen.imagegen'
AGENT='/root/astra_spec'
TITLES=['Gitの保存・記録・共有','変更は3つの場所を進む','確認してから、チームへ共有する']
NOTES=[
'エディターでファイルを保存すること、Gitで履歴に記録すること、リモートへ共有することを分けて覚えます。保存しただけではコミットされません。コミットしても、まだ他の人へ送ったことにはなりません。まず手元で変更を確かめ、小さな単位で記録してから共有しましょう。',
'作業ツリーは編集中のファイルです。git addは次のコミットに含める内容をステージへ選びます。git commitはそのステージの内容をローカルの履歴に記録します。addした後にさらに編集した内容は自動でステージへ入りません。必要ならもう一度addし、git diff --stagedで内容を確かめます。',
'最初にgit statusとgit diffで変更の範囲を確認します。対象ファイルだけをaddし、git diff --stagedで記録される差分を読みます。秘密情報や不要なファイルがないことを確かめてcommitします。共有先と作業ブランチを確認し、チームのルールに従ってpushとレビュー依頼を行います。ここでは共有の手順を説明しているだけで、実際のpushは実行しません。',
]
SOURCES=['https://git-scm.com/book/en/v2/Git-Basics-Recording-Changes-to-the-Repository',
         'https://git-scm.com/book/en/v2/Git-Basics-Working-with-Remotes']
def write(p,s):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(s,encoding='utf-8')
def dump(p,d):write(p,json.dumps(d,ensure_ascii=False,indent=2)+'\n')
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def rel(p):return p.relative_to(ROOT).as_posix()
def run(script,*args):
    cmd=[str(PY),str(SKILL/'scripts'/script),*[str(a) for a in args]]
    r=subprocess.run(cmd,cwd=ROOT,capture_output=True,text=True,encoding='utf-8',timeout=90)
    ledger=OUT/'source/commands.json';rows=load(ledger) if ledger.exists() else []
    log=OUT/'source/logs'/f'{len(rows)+1:02d}-{script}.txt';write(log,r.stdout+'\n'+r.stderr)
    rows.append({'argv':cmd,'exit_code':r.returncode,'log':rel(log),'log_sha256':sha(log),'at':datetime.now(timezone.utc).isoformat()});dump(ledger,rows)
    if r.returncode:raise RuntimeError(r.stdout+r.stderr)
    print(r.stdout)

def prepare():
    if (OUT/'slide_jobs.json').exists():raise RuntimeError('Existing jobs preserved; do not overwrite')
    OUT.mkdir(parents=True,exist_ok=True)
    style={'name':'Japanese clean vector explainer','visual_direction':'Pure white background; crisp flat vector-like diagrams, deep navy headings, teal accents, pale blue panels; no photo, no texture, no gradients, no 3D. Generous whitespace. All visible copy Japanese except Git commands.','color_palette':'white #FFFFFF, navy #172B4D, teal #168B87, pale blue #EAF3F8','typography':'Large legible Japanese sans-serif, bold title, high contrast; all exact text is readable at presentation size.'}
    dump(OUT/'style.json',style)
    write(OUT/'outline.md','# Gitの保存・記録・共有\n\n3ページ・16:9・日本語。写真と外部画像なし。\n\n'+ '\n\n'.join(f'## Slide {i+1}: {t}\n\n{NOTES[i]}' for i,t in enumerate(TITLES)))
    dump(OUT/'input/data.json',{'title':TITLES[0],'audience':'Gitを使い始めた人','facts':NOTES,'sources':SOURCES,'source_checked_at':'2026-10-14','actual_push_executed':False})
    write(OUT/'input/BRIEF.md','# Git入門・3ページ\n\n保存、ステージへの選択、ローカル記録、リモート共有を混同せず説明する。画像型PPTXとして制作。\n')
    method={'backend_used':BACKEND,'tool_name':'image_gen.imagegen','mode':'generate','model':None,'model_evidence':'Tool does not expose an exact model identifier.','size':'16:9 requested in prompt; inspect actual raster dimensions','input_context_preparation':'View generated sample before referencing it as style only','approval_authority':'production agent under user-authorized autonomous generation; no human visual approval claimed'}
    slides=[
      {'number':1,'title':TITLES[0],'role':'cover','intent':'Teach the distinction before commands','key_points':['保存：ファイルに反映','記録：手元の履歴に残す','共有：リモートへ送る'],'layout':{'composition':'Large bold title upper-left, subtitle directly below. Lower half has three clear horizontally aligned simple vector icons: file, commit/history, shared cloud. Each has its exact label. Minimal cover.'},'visual_elements':{'main_visual':'Three distinct concepts, clean flat icons with no extra labels'},'constraints':['Visible subtitle exactly: はじめてのGit・3分入門','Visible bottom takeaway exactly: 保存しただけでは、履歴にも共有先にも届かない','No slide number, no extra English slogans.']},
      {'number':2,'title':TITLES[1],'role':'concept diagram','intent':'Show working tree, staging area and local repository accurately','key_points':['作業ツリー','ステージ','ローカル履歴'],'layout':{'composition':'Large title top. Three large panels left-to-right at same height, connected by rightward arrows. Left panel exact body 編集中のファイル; middle 次に記録する内容; right コミットの記録. Arrow left→middle labeled git add; middle→right labeled git commit. Bottom two short callouts.'},'visual_elements':{'main_visual':'Three-stage state diagram with local computer boundary enclosing ALL three panels. No cloud on this page.'},'constraints':['Bottom exact text: git add は記録する内容を選ぶ','Bottom exact text: git commit で手元の履歴に残す','Show only these texts and title; do not invent commands or labels.']},
      {'number':3,'title':TITLES[2],'role':'procedure','intent':'Check changes, select, inspect staged content, commit and share','key_points':['変更を確認','記録する内容を選ぶ','差分を確認して記録','共有先を確認して送る'],'layout':{'composition':'Four numbered steps in a spacious vertical flow on the left two-thirds; right third a simple local-history arrow to team-cloud illustration. Step1 command git status / git diff. Step2 command git add README.md. Step3 command git diff --staged → git commit. Step4 command git push. Keep all commands visibly aligned and large.'},'visual_elements':{'main_visual':'Readable 4-step checklist and small share diagram. No dense terminal screenshot.'},'constraints':['Bottom exact takeaway: 小さく記録して、レビューを依頼する','Small but readable note exactly: pushの前に、共有先と作業ブランチを確認','The numbered four steps are process steps, not slide numbers.']},
    ]
    spec={'title':TITLES[0],'language':'Japanese','aspect_ratio':'16:9','selected_image_backend':BACKEND,'max_concurrent_slides':1,'sample_generation_method':method,'style':style,'deck_context':{'source_summary':'Git basics; saving files, selecting staged content, committing locally, pushing to remote are distinct actions.','canonical_terms':['作業ツリー','ステージ','ローカル履歴','リモート'],'core_claim':'Confirm the exact changes before recording and sharing.'},'slides':slides}
    dump(OUT/'deck_spec.json',spec)
    dump(OUT/'workflow-decisions.json',{'authority':'User requested actual slides and autonomous production; parent assigned this three-page native image example.','outline':'production-agent accepted','style':'production-agent selected from specified brief','backend':BACKEND,'sample':'pending production visual review','human_visual_approval':False,'workflow_variation':'One assigned production agent handles three pages; no extra child agents. Original CLI dispatch/result records use the real production agent id. No per-page child spawn is claimed.','paid_api_calls':0,'model_identity':None})
    write(OUT/'speech.md','\n\n'.join(f'## Slide {i+1}: {t}\n\n{NOTES[i]}' for i,t in enumerate(TITLES))+'\n')
    run('prepare_slide_prompts.py','--spec',OUT/'deck_spec.json','--out-dir',OUT)
    for path in ['LICENSE','skills/codex-ppt/SKILL.md','skills/codex-ppt/scripts/assemble_ppt.py','skills/codex-ppt/scripts/prepare_slide_prompts.py','skills/codex-ppt/scripts/record_slide_dispatch.py','skills/codex-ppt/scripts/record_slide_result.py','skills/codex-ppt/scripts/slide_run_state.py','skills/codex-ppt/scripts/slide_job_status.py']:
        dst=OUT/'source/upstream'/path;dst.parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(UP/path,dst)
    shutil.copyfile(UP/'LICENSE',OUT/'LICENSE.upstream')

def dispatch(number):
    run('record_slide_dispatch.py',OUT,'--slide',f'slide_{number:02d}','--agent-id',AGENT,'--agent-nickname','Astra production agent','--prompt-file',f'prompts/slide_{number:02d}.json')

def record(number,source,note):
    run('record_slide_result.py',OUT,'--slide',f'slide_{number:02d}','--agent-id',AGENT,'--backend-used',BACKEND,'--selected-source',source,'--qa-note',note)
    if number==1:
        decisions=load(OUT/'workflow-decisions.json');decisions['sample']='production-agent visually inspected and selected; not human-approved';dump(OUT/'workflow-decisions.json',decisions)
        spec=load(OUT/'deck_spec.json');sample=str((OUT/'origin_image/slide_01.png').resolve())
        spec['approved_style_reference']={'path':sample,'role':'production-agent selected sample; style only','fidelity':'Do not copy layout/content'};dump(OUT/'deck_spec.json',spec)
        for n in (2,3):
            p=OUT/f'prompts/slide_{n:02d}.json';job=load(p)
            job['input_images']=[{'path':sample,'role':'style-only sample selected by production agent; not human-approved'}]
            job['prompt']+='\n\nUse the attached sample ONLY for the white/navy/teal visual style and typography. Use this page\'s exact text and composition, not the sample content. The sample was selected by the production agent.'
            dump(p,job)

def assemble():
    run('slide_job_status.py',OUT)
    jobs=load(OUT/'slide_jobs.json')
    if not all(s['status'] in ('recorded','accepted') for s in jobs['slides']):raise RuntimeError('Unrecorded image jobs')
    run('assemble_ppt.py',OUT.parent,OUT.name+'.pptx','--aspect-ratio','16:9')

def finalize(reviewed):
    if not reviewed:raise RuntimeError('Pass --reviewed only after inspecting all source images and PowerPoint renders')
    from PIL import Image
    from pptx import Presentation
    import zipfile, importlib.metadata
    ppt=OUT/(OUT.name+'.pptx');receipt=load(OUT/'render/powerpoint-render.json')
    if receipt['source_sha256']!=sha(ppt) or receipt['slide_count']!=3:raise RuntimeError('Stale PowerPoint rendering')
    prs=Presentation(ppt)
    if len(prs.slides)!=3 or any(len(s.shapes)!=1 or s.shapes[0].shape_type!=13 for s in prs.slides):raise RuntimeError('Expected one full-slide raster per page')
    if any(not s.notes_slide.notes_text_frame.text.strip() for s in prs.slides):raise RuntimeError('Missing notes')
    images=[]
    with zipfile.ZipFile(ppt) as z:
        embedded={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('ppt/media/')}
        for i in range(1,4):
            source=OUT/f'origin_image/slide_{i:02d}.png';render=OUT/f'render/slide-{i:02d}.png'
            im=Image.open(source);im.verify()
            if sha(source) not in embedded:raise RuntimeError('Original image not embedded byte-identically')
            images.append({'page':i,'source':rel(source),'source_sha256':sha(source),'dimensions':list(Image.open(source).size),'render':rel(render),'render_sha256':sha(render),'visually_reviewed':True})
    pin=load(UP/'ACQUISITION.json')
    refs=[]
    for f in sorted((OUT/'source/upstream').rglob('*')):
        if f.is_file():
            path=f.relative_to(OUT/'source/upstream').as_posix();refs.append({'path':path,'sha256':sha(f),'url':f"https://github.com/ningzimu/codex-ppt-skill/blob/{pin['commit']}/{path}"})
    method='原典prepare_slide_prompts/dispatch/result/state CLIを使用し、Codex native image_gen.imagegenで全3ページを生成。原典assemble_ppt.pyが画像とspeech.mdからPPTXを構築。'
    dump(OUT/'PROVENANCE.json',{'repository':'ningzimu/codex-ppt-skill','upstream_commit':pin['commit'],'acquisition':pin,'references':refs,'method':method,'execution_type':'native_image_generation_and_original_packager','commands_executed':load(OUT/'source/commands.json'),'artifact':rel(ppt),'artifact_sha256':sha(ppt),'input_sha256':sha(OUT/'input/data.json'),'image_backend':BACKEND,'model':None,'model_note':'実ツールはモデル名を公開しない。GPT Image2利用とは断定しない。','editability':'raster_slides','native_requests':'source/native-imagegen-requests.json','image_records':images,'workflow_decisions':'workflow-decisions.json','upstream_source_modified':False,'raster_content_postprocessing':False,'generator':'tools/build_codex_image_repository_example.py','checked_at':datetime.now(timezone.utc).isoformat()})
    dump(OUT/'QA.json',{'status':'qa_complete','artifact_sha256':sha(ppt),'page_count':3,'visual_review':{'status':'passed','reviewer':'production agent /root/astra_spec','all_pages_reviewed':True,'pages':[1,2,3],'source_and_powerpoint_reviewed':True,'findings':['全3頁の日本語本文とGitコマンドに読めない文字・欠けなし。','状態図は3領域をローカルPC枠内に置き、add/commitの矢印方向を確認。','3頁目のgit diff --staged、README.md、共有先/作業ブランチの注記を視認。']},'images':images,'powerpoint_render':{'status':'passed','receipt':'render/powerpoint-render.json','source_sha256':sha(ppt)},'structure':{'picture_objects':3,'native_text_objects':0,'speaker_notes':3,'original_images_embedded_byte_identically':True},'editability':'raster_slides','not_claimed':['本文/図の個別編集','人による採用','実モデル名の確認','Git push実行','外部公開']})
    dump(OUT/'environment.json',{'python':__import__('sys').version,'dependencies':{n:importlib.metadata.version(n) for n in ('python-pptx','Pillow','filelock')},'runtime':'project-local .build/ppt-master-venv; no common skill bootstrap','image_backend':BACKEND,'model':None,'powerpoint':receipt['version']})
    write(OUT/'PROMPT.txt','Gitの保存・記録・共有を説明する日本語3ページ。表紙、作業ツリー/ステージ/ローカル履歴の3段図、確認して共有する手順。16:9、写真なし、読みやすいベクトル風図解。ningzimu/codex-ppt-skillのnative画像生成と原典packagerを使う。\n')
    shutil.copyfile(OUT/'speech.md',OUT/'NOTES.md')
    guide='''# Gitの保存・記録・共有 — 画像型PPTX

Codex内蔵の `image_gen.imagegen` で生成した日本語3ページを、ningzimu/codex-ppt-skill の原典 `assemble_ppt.py` でPowerPointにしました。全ページの元画像とPowerPoint 16.0のPNG/PDFを制作担当が確認済みです。

**各ページは1枚の画像です。本文・図・コマンドをPowerPoint上で個別編集できません。** 発表者ノートは日本語で3ページ分入っています。内蔵ツールの実モデル名は公開されないため、GPT Image2使用とは記載しません。

## 再現する

カタログrootからproject-local Python環境を使います。必要な組版依存は `python-pptx`, `Pillow`, `filelock`。共通skill runtimeのbootstrap、APIキー設定、外部API呼び出しは行いません。

1. `source/upstream/` の原典コードとLICENSE、`deck_spec.json`、`outline.md`、`style.json`を確認します。
2. 新しい制作フォルダーで原典 `prepare_slide_prompts.py --spec <deck_spec.json> --out-dir <deck>` を実行します。既存の生成記録を上書きしないでください。
3. `prompts/slide_XX.json` を読み、実際に利用可能なnative画像ツールで1頁ずつ生成します。実際に送った全文と引数は `source/native-imagegen-requests.json` にあります。2・3頁は1頁目を作風だけの参考にしています。
4. 原典 `record_slide_dispatch.py` / `record_slide_result.py` で本当の担当者ID・生成元パス・目視結果を記録します。生成画像を描画コードで置換しません。
5. 全画像を視認し、`speech.md` を用意して原典 `assemble_ppt.py <deckの親> <deck名>.pptx --aspect-ratio 16:9` を実行します。
6. PowerPointで全ページを再描画し、画像とノートを確認します。このカタログでは `pwsh -File tools/render_powerpoint.ps1 -PresentationPath <pptx> -OutputDirectory <render>` を使用しました。

既存画像からの再組版は `python tools/build_codex_image_repository_example.py assemble` で可能です。新規画像の生成はホストツールが必要で、Pythonだけでは実行できません。再組版後はPowerPointレンダーも更新します。

## 原典手順との差分

ユーザーから自律制作を依頼された範囲で、outline・style・sampleを制作担当が判断しました。人が見て承認したという記録ではありません。また、3頁を1人の割り当て済み制作agentが担当し、1頁ごとの追加子agentは作っていません。原典のdispatch/result CLIには実担当者IDを記録しています。`workflow-decisions.json` に差分を明示しました。

原典prompt compilerには言語共通のChinese-text文が残るため、実送信prompt末尾でJapanese指定を明確化しました。原典ソースは無改変です。元画像1672×941を原典packagerが16:9ページへ配置しています。画像の文字を手作業で描き足す処理はしていません。

## 出典と権利

Gitの説明は [Pro Git: Recording Changes](https://git-scm.com/book/en/v2/Git-Basics-Recording-Changes-to-the-Repository) と [Working with Remotes](https://git-scm.com/book/en/v2/Git-Basics-Working-with-Remotes) を確認した独自の日本語要約です。実際のpushは行っていません。

原典コードはMIT。`LICENSE.upstream` と `source/upstream/LICENSE` の著作権・許諾表示を保持しています。画像は今回の生成物で、原典MITコードと生成画像の扱いを同一視しません。
'''
    guide+='''
## 再現キットに含めるファイル

自動の再現スクリプトは**既存3画像の原典組版だけ**を実行します。画像再生成は画像利用枠を消費する別操作です。同梱promptをnative画像ツールへ依頼する手順として案内し、APIキーを探す処理や自動画像生成は入れません。

最小のオフライン組版キットは、次の配置にします。

```text
reproduce/
  source/upstream/LICENSE
  source/upstream/skills/codex-ppt/scripts/assemble_ppt.py
  workspace/codex-image-ppt/origin_image/slide_01.png
  workspace/codex-image-ppt/origin_image/slide_02.png
  workspace/codex-image-ppt/origin_image/slide_03.png
  workspace/codex-image-ppt/speech.md
```

キット内の隔離Pythonへ `python-pptx` と `Pillow` を導入済みなら、キットrootから次を実行できます。原典packagerは画像の原本と3頁のノートを読み、同名フォルダー内にPPTXを出力します。

```powershell
& ./.venv/Scripts/python.exe ./source/upstream/skills/codex-ppt/scripts/assemble_ppt.py ./workspace codex-image-ppt.pptx --aspect-ratio 16:9
```

配布の再現証拠には上記に加え、`LICENSE.upstream`、`PROVENANCE.json`、`QA.json`、`environment.json`、`input/`、`PROMPT.txt`、`outline.md`、`style.json`、`deck_spec.json`、`workflow-decisions.json`、`slide_jobs.json`、`slide_run_state.json`、`prompts/`、`source/native-imagegen-requests.json`、`source/commands.json`、`source/logs/`、`render/powerpoint-render.json` を含めます。原典のjobs/state経路も再現する場合は `source/upstream/skills/codex-ppt/scripts/` の同梱Pythonスクリプト一式と `filelock` が必要です。

`tools/build_codex_image_repository_example.py` はカタログ用の補助で、単体の画像生成器ではありません。カタログroot、`.build/actual-upstreams/ningzimu__codex-ppt-skill`、`.build/ppt-master-venv/Scripts/python.exe`、`examples/repositories/codex-image-ppt` という配置に依存します。持ち運ぶキットでは、このhelperをそのまま動かすより上記の原典packagerを直接呼んでください。PowerPointによる再QAはWindowsとPowerPointを別途要します。
'''
    write(OUT/'README.md',guide)
    # Viewer only displays actual PowerPoint renders; it never redraws slides.
    write(OUT/'index.html','''<!doctype html><html lang="ja"><meta charset="utf-8"><title>Gitの保存・記録・共有</title><style>html,body{margin:0;background:#182231;height:100%;overflow:hidden}img{width:100%;height:100%;object-fit:contain}</style><img id="slide" src="render/slide-01.png" alt="Gitの保存・記録・共有"><script>let n=1;function go(v){n=Math.max(1,Math.min(3,v));document.querySelector('img').src='render/slide-'+String(n).padStart(2,'0')+'.png';parent.postMessage({type:'yebisu-page',page:n,total:3},location.origin)}addEventListener('keydown',e=>{if(e.key==='ArrowRight')go(n+1);if(e.key==='ArrowLeft')go(n-1)});addEventListener('message',e=>{if(e.source!==parent||e.origin!==location.origin)return;if(e.data?.type==='yebisu-slide')go(e.data.page)});go(1);</script></html>''')
    base=rel(OUT);card={'id':'codex-image-ppt','title':TITLES[0],'repository':'ningzimu/codex-ppt-skill','skill_repository':'ningzimu/codex-ppt-skill','upstream_commit':pin['commit'],'fork_id':None,'page_count':3,'format':'pptx','preview_kind':'images','editability':'raster_slides','use_case':'Git入門・図解','style':'native画像生成・白/紺/青緑の図解','cover':base+'/render/slide-01.png','interior':base+'/render/slide-02.png','preview':base+'/index.html','pages':[base+f'/render/slide-{i:02d}.png' for i in range(1,4)],'artifact_download':rel(ppt),'source_download':base+'/source/build_codex_image_repository_example.py','pdf_download':base+'/render/deck.pdf','input':base+'/input/BRIEF.md','input_data':base+'/input/data.json','prompt':base+'/PROMPT.txt','prompt_text':(OUT/'PROMPT.txt').read_text(encoding='utf-8').strip(),'provenance':base+'/PROVENANCE.json','qa':base+'/QA.json','guide':base+'/README.md','fork_download':None,'titles':TITLES,'notes':NOTES,'method':method,'image_backend':BACKEND,'model':None,'sample_notice':'実際のGit送信を行わない入門教材。本文・図は画像であり、個別編集不可。'}
    dump(ROOT/'data/codex-image-repository-examples.json',[card])
    dump(ROOT/'data/codex-image-repository-guides.json',[{'repository':'ningzimu/codex-ppt-skill','execution_type':'native_image_generation_and_original_packager','summary_ja':method,'case_ids':['codex-image-ppt'],'requirements':['Codex native image_gen.imagegen','Python / python-pptx / Pillow / filelock','PowerPointとPowerShell 7（実描画QA）'],'steps':['原典spec/prompt作成','native画像3頁生成・視認','原典dispatch/resultで記録','原典assemble_ppt.py','PowerPoint全頁確認'],'commands':['python tools/build_codex_image_repository_example.py assemble'],'source_refs':refs,'limitations':['本文と図はraster_slidesで個別編集不可','実モデル名はtool非公開','sampleは制作agent判断で人の承認ではない','3頁を同一担当者が制作した運用差分を記録'],'guide':base+'/README.md'}])
    shutil.copyfile(Path(__file__),OUT/'source/build_codex_image_repository_example.py')
    print('Finalized 3 raster slides, original images unchanged, 3 speaker notes, PowerPoint PNG/PDF reviewed.')

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('action',choices=['prepare','dispatch','record','assemble','finalize']);p.add_argument('--slide',type=int);p.add_argument('--source',type=Path);p.add_argument('--qa-note');p.add_argument('--reviewed',action='store_true');a=p.parse_args()
    if a.action=='prepare':prepare()
    elif a.action=='dispatch':dispatch(a.slide)
    elif a.action=='record':record(a.slide,a.source,a.qa_note)
    elif a.action=='assemble':assemble()
    else:finalize(a.reviewed)
