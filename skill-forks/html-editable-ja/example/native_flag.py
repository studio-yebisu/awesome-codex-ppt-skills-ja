"""Case adapter: the known simple five-vertex SVG flag becomes one native shape.

Runs after unchanged original assemble and before original Stage5. This is not
an arbitrary SVG-to-native capability of the upstream converter.
"""
from pathlib import Path
import json,hashlib
from pptx import Presentation
from pptx.dml.color import RGBColor
PX_TO_EMU=6350
def replace_simple_flag(pptx):
    path=Path(pptx);prs=Presentation(path);slide=prs.slides[0]
    expected=(1440,285,315,245)
    candidates=[s for s in slide.shapes if s.shape_type==13 and all(abs(a/PX_TO_EMU-b)<0.1 for a,b in zip([s.left,s.top,s.width,s.height],expected))]
    if len(candidates)!=1:raise RuntimeError('Expected exactly one known flag picture at its measured source bounds.')
    pic=candidates[0];old_blob_sha=hashlib.sha256(pic.image.blob).hexdigest();before_count=len(slide.shapes)
    # Scale by half-pixels so the 122.5px notch is exact integer geometry.
    builder=slide.shapes.build_freeform(0,0,scale=PX_TO_EMU/2)
    builder.add_line_segments([(630,0),(470,245),(630,490),(0,490)],close=True)
    shape=builder.convert_to_shape(origin_x=pic.left,origin_y=pic.top)
    shape.name='NativeFlag:five-vertex-cloth'
    shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string('38271F');shape.line.fill.background()
    parent=pic._element.getparent();index=parent.index(pic._element);parent.remove(shape._element);parent.insert(index,shape._element)
    rid=pic._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip').get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed');parent.remove(pic._element);slide.part.drop_rel(rid)
    assert len(slide.shapes)==before_count and sum(s.shape_type==13 for s in slide.shapes)==0
    prs.save(path)
    report={'scope':'Case-specific known simple flag only, not general SVG conversion. Original converter unchanged; applied before original Stage5a/5b.','original_svg_path':'M0 0 H315 L235 122.5 L315 245 H0 Z','measured_bounds_css_px':list(expected),'native_vertices_css_px':[[0,0],[315,0],[235,122.5],[315,245],[0,245]],'native_shape_name':shape.name,'shape_type':'native_freeform_custom_geometry','source_picture_sha256':old_blob_sha,'slide1_pictures_after':0,'slide1_shape_count_before':before_count,'slide1_shape_count_after':len(slide.shapes),'other_slide_content_intentionally_changed':False,'font_embedding_changed':False}
    out=Path(__file__).parent/'qa/native-flag.json';out.parent.mkdir(exist_ok=True);out.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print('[case-adapter] Known five-point flag converted to editable native freeform before original Stage5.')
