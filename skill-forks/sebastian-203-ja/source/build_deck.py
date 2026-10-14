from pathlib import Path
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
import json
ROOT=Path(__file__).resolve().parents[1]
p=Presentation(ROOT/'source/academic-template.pptx')
base=p.slides[1]
for _ in range(6):
 s=p.slides.add_slide(base.slide_layout)
 for sh in base.shapes:s.shapes._spTree.insert_element_before(deepcopy(sh.element),'p:extLst')
INK='0B3A5B';GRAY='5D6B78';BLUE='DDEAF1';TEAL='347B7F';WHITE='FFFFFF';DARK='202A33'
def text(s,x,y,w,h,t,size=20,color=INK,bold=False,font='Noto Sans CJK JP'):
 sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=sh.text_frame;tf.word_wrap=True
 tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(t.split('\n')):
  pp=tf.paragraphs[0] if i==0 else tf.add_paragraph();pp.text=line
  for r in pp.runs:r.font.name=font;r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
 return sh
def rect(s,x,y,w,h,color=BLUE,line=WHITE):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.color.rgb=RGBColor.from_string(line);return sh
def line(s,x,y,x2,y2,color=GRAY,width=1):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x2), Inches(y2));sh.line.color.rgb=RGBColor.from_string(color);sh.line.width=Pt(width)
def clear(s,cover=False):
 # Preserve content template rule, footer bar and slide layout/master, replacing specimen strings.
 keep=[] if cover else [3,4]
 for j,sh in reversed(list(enumerate(s.shapes))):
  if j not in keep:s.shapes._spTree.remove(sh.element)
def page(s,n,title,section,claim,note):
 clear(s)
 text(s,.45,.2,8,.25,'静けさを選べる図書館',10,GRAY)
 text(s,10,.2,2.5,.25,section,9,GRAY)
 text(s,.45,.62,12,.5,title,25,INK,True)
 text(s,.65,1.42,11.8,.45,claim,15,DARK)
 text(s,.45,7.2,10,.2,'架空の研究提案・未実施　／　効果を示す実測値はありません',7.2,WHITE)
 text(s,12.35,7.2,.45,.2,f'{n:02d}',8,WHITE)
 text(s,.75,6.3,11.6,.55,note,10,GRAY)
 s.notes_slide.notes_text_frame.text=f'{title}\n{claim}\n{note}\n架空の未実施研究提案。数値結果・実在の被験者・外部引用は含まない。'
titles=['静けさを選べる\n図書館','利用目的と音環境のずれ','ゾーンの境界と音の伝わり方','検証する関係','配置案の比較','試行と観察の順序','評価項目と判断の条件','実施前に確認すること']
notes=[]
s=p.slides[0];clear(s,True)
text(s,.75,1.25,10,1.55,titles[0],38,INK,True)
text(s,.75,3.25,9,.6,'図書館の音環境を、利用者が選べるか',21,DARK)
text(s,.75,4.25,9,.8,'ゾーニングと案内表示の研究計画\n架空の提案・未実施',15,GRAY)
line(s,.75,5.3,12.2,5.3,INK,2)
text(s,.75,5.55,10,.5,'静かな閲覧　／　小声の相談　／　会話のある活動',15,INK)
text(s,12.35,7.2,.45,.2,'01',8,INK)
notes.append('表紙。架空の未実施研究提案。')
s=p.slides[1];page(s,2,titles[1],'背景','同じ図書館でも、必要とする音環境は異なる。','読図：利用目的と期待する音環境の対応を読む。分類はこの提案の設計仮説。')
for i,(a,b) in enumerate([('読む・集中する','静かな席を探す'),('短く相談する','小声で話せる席を探す'),('活動を共有する','会話できる場所を探す')]):
 y=2.2+i*1.05;text(s,.8,y,4,.45,a,23,INK,True);line(s,5,y+.24,6.2,y+.24,TEAL,2);text(s,6.8,y,5,.6,b,20,DARK)
notes.append('利用目的と期待する音環境の対応。分類は設計仮説。')
s=p.slides[2];page(s,3,titles[2],'空間モデル','静かな席へ向かう動線と、会話のある活動を分けて考える。','読図：入口から静穏ゾーンへ進む経路と、隣接する境界を見る。縮尺のない概念図。')
rect(s,.8,2.3,2.5,2.9,INK);rect(s,3.45,2.3,3.1,2.9,TEAL);rect(s,6.7,2.3,5.5,2.9,BLUE)
for x,w,a,b,c in [(.95,2.2,'会話','活動・相談',WHITE),(3.65,2.7,'移行','通路・書架',WHITE),(7,4.9,'静穏','閲覧・集中',INK)]:text(s,x,2.75,w,.6,a,29,c,True);text(s,x,3.7,w,.7,b,16,c)
text(s,.8,5.45,11,.6,'入口　　　　　　　　　　　案内表示　　　　　　　　　　静かな席',14,GRAY)
notes.append('縮尺のないゾーニング概念図。建築図や実測音分布ではない。')
s=p.slides[3];page(s,4,titles[3],'研究仮説','配置と表示が、席の選択と音環境の評価に関係するか。','読図：左は介入する条件、中央は行動、右は評価。線は検証候補で、因果関係の実証ではない。')
for x,t,b in [(.8,'配置・表示','境界の配置\n音のルールの説明'),(5,'席の選択','希望に合う席を選べるか\n迷いや移動が起きるか'),(9.3,'音環境の評価','集中しやすさ\n会話のしやすさ')]:text(s,x,2.55,3.3,.65,t,24,INK,True);text(s,x,3.65,3.2,1.1,b,17,DARK)
line(s,3.75,3,4.65,3,TEAL,3);line(s,8.1,3,9,3,TEAL,3)
text(s,.8,5.15,11,.5,'交絡の候補：混雑、時間帯、利用目的、活動内容',16,GRAY)
notes.append('配置・表示と席選択・評価の仮説モデル。実証済み因果関係ではない。')
s=p.slides[4];page(s,5,titles[4],'比較条件','案Aと案Bを、同じ観察項目で比較する。','読表：境界と表示だけを先に比較する。席数や面積の最終設定は現地条件の確認後に決める。')
rows=[['比較する条件','案A：境界を明示','案B：ゆるやかな移行'],['ゾーン境界','静穏と会話の境界を示す','書架・通路を移行帯にする'],['案内表示','各ゾーンの音のルール','入口で利用目的別に案内'],['検証したい点','ルールの分かりやすさ','移動と音の受け止め方']]
for r,row in enumerate(rows):
 for c,tx in enumerate(row):
  x=.8+[0,3.25,7.25][c];w=[3.05,3.8,4][c];y=2.15+r*.8
  if r==0:rect(s,x,y,w,.68,INK)
  text(s,x+.1,y+.1,w-.2,.6,tx,14,WHITE if r==0 else DARK,r==0)
  if r>0:line(s,x,y+.74,x+w,y+.74,BLUE,1)
notes.append('二つの未実施配置案の比較。席数・面積の実績を含まない。')
s=p.slides[5];page(s,6,titles[5],'調査方法','現状を観察し、小さく試し、同じ項目で再び観察する。','読図：順序は研究計画。試行前に施設の許可と倫理面の確認を行う。日程は未設定。')
for i,(a,b) in enumerate([('現状把握','利用目的と席の選択\n音への困りごと'),('案の試行','配置と表示を仮設\n条件変更を記録'),('再観察','同じ項目を観察\n利用者の評価を確認'),('比較・修正','混雑などの違いを確認\n案の修正点を整理')]):
 x=.8+i*3.05;text(s,x,2.35,2.8,.7,f'{i+1:02d}',32,TEAL,True);text(s,x,3.3,2.8,.6,a,21,INK,True);text(s,x,4.25,2.7,1.1,b,15,DARK)
line(s,.8,3.1,12.1,3.1,BLUE,2)
notes.append('現状把握・仮設試行・再観察・比較修正の研究手順。未実施。')
s=p.slides[6];page(s,7,titles[6],'評価設計','席を選べることと、必要な活動を続けられることを評価する。','読表：評価項目と取得方法の対応。採否の基準値は事前に定め、調査後に都合よく変えない。')
for i,(a,b) in enumerate([('席の選びやすさ','利用者への任意の短い質問'),('集中・会話のしやすさ','利用目的と合わせた自己評価'),('席の移動・混雑','個人を特定しない観察記録'),('音環境の状態','許可を得た測定と条件記録')]):
 y=2.2+i*.78;text(s,.8,y,4.5,.5,a,20,INK,True);text(s,6,y,6,.5,b,17,DARK);line(s,.8,y+.65,12,y+.65,BLUE)
notes.append('評価項目と予定取得方法。調査結果・基準値を捏造しない。')
s=p.slides[7];page(s,8,titles[7],'実施の条件','運用とプライバシーを確認してから、試行の範囲を決める。','提案の位置づけ：設計と評価の計画。効果、利用者の支持、実施可否はまだ確認していない。')
text(s,.8,2.25,5.8,2.1,'静かな席も、\n話せる場所も。',32,INK,True)
for i,(a,b) in enumerate([('施設との合意','避難・動線・運用への影響'),('参加者への説明','任意参加、同意、断る自由'),('記録の取り扱い','識別情報を集めない方針')]):text(s,7,2.25+i*1.05,5,.45,a,19,TEAL,True);text(s,7,2.78+i*1.05,5,.45,b,14,DARK)
notes.append('施設合意、任意参加、個人識別情報を集めない方針。効果未確認。')
for s,n in zip(p.slides,notes):s.notes_slide.notes_text_frame.text=n+'\n架空の未実施研究提案。'
p.save(ROOT/'deck.pptx')
(ROOT/'input/data.json').write_text(json.dumps({'status':'架空・未実施研究提案','titles':[t.replace('\n','') for t in titles],'notes':notes,'numeric_results':None,'citations':[]},ensure_ascii=False,indent=2))
