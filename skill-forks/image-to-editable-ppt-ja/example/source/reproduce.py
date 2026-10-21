"""Replay accepted worker manifests through the original finalize CLI, then add notes.

This replays already authored page manifests; it does not reconstruct new pages,
generate images, perform OCR, or fabricate fresh worker dispatches.
"""
from pathlib import Path
import argparse,hashlib,json,os,shutil,subprocess,sys,zipfile
CASE=Path(__file__).resolve().parents[1]
RUN_NAME='run-repair-r4'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def dump(p,d):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def add_notes(out,source):
 from pptx import Presentation
 data=load(out/'input/data.json');dest=out/'deck.pptx'
 def stable(path):
  with zipfile.ZipFile(path) as z:return {n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if (n.startswith('ppt/slides/slide') and n.endswith('.xml')) or n.startswith('ppt/media/')}
 before=stable(source);prs=Presentation(source)
 for slide,page in zip(prs.slides,data['pages']):slide.notes_slide.notes_text_frame.text=page['note']
 prs.save(dest);tmp=dest.with_suffix('.notes-tmp.pptx')
 with zipfile.ZipFile(source) as a,zipfile.ZipFile(dest) as b,zipfile.ZipFile(tmp,'w') as z:
  for info in b.infolist():z.writestr(info,a.read(info.filename) if info.filename in before else b.read(info.filename))
 tmp.replace(dest);assert before==stable(dest)
 assert [s.notes_slide.notes_text_frame.text for s in Presentation(dest).slides]==[p['note'] for p in data['pages']]
 dump(out/'source/notes-addition.json',{'source_sha256':sha(source),'input_sha256':sha(out/'input/data.json'),'artifact_sha256':sha(dest),'operation':'Newly authored catalog speaker notes; not extracted from source images','slide_xml_and_media_unchanged':True,'verified_parts':len(before)})
def main():
 ap=argparse.ArgumentParser();ap.add_argument('--output-dir',type=Path,default=CASE);a=ap.parse_args();out=a.output_dir.resolve()
 if out!=CASE:
  if out.exists():raise SystemExit('A separate replay destination must not already exist')
  out.mkdir(parents=True)
  for name in ['input',RUN_NAME,'source/upstream']:shutil.copytree(CASE/name,out/name,ignore=shutil.ignore_patterns('__pycache__','*.pyc'))
 run=out/RUN_NAME;cli=out/'source/upstream/skills/image-to-editable-ppt/cli/editppt/cli.py'
 # job_dir belongs to orchestrator run metadata. Worker page manifests stay byte-identical.
 manifests={p:sha(p) for p in (run/'pages').glob('*/manifest.json')}
 deck=load(run/'deck_manifest.json');deck['job_dir']=str(run);dump(run/'deck_manifest.json',deck)
 env=os.environ.copy()
 for k in ['OPENAI_API_KEY','OPENAI_BASE_URL','PADDLE_OCR_TOKEN']:env.pop(k,None)
 cfg=out/'source/offline-config';cfg.mkdir(parents=True,exist_ok=True)
 env.update(EDITPPT_CONFIG_HOME=str(cfg),CODEX_AUTH_FILE=str(cfg/'no-auth.json'),CODEX_PPT_HOME=str(cfg/'no-codex-ppt'),PYTHONIOENCODING='utf-8',PYTHONUTF8='1',PYTHONPATH=os.pathsep.join([str(cli.parents[1]),str(cli.parent/'runtime')]))
 args=[sys.executable,str(cli),'run','finalize',str(run)];r=subprocess.run(args,cwd=out,env=env,capture_output=True,text=True,encoding='utf-8',errors='replace')
 (out/'source/current-finalize.log').write_text(r.stdout+r.stderr,encoding='utf-8')
 if r.returncode:raise SystemExit(r.returncode)
 assert all(sha(p)==h for p,h in manifests.items()),'Worker manifests must not be changed by replay'
 add_notes(out,run/deck['output'])
 dump(out/'source/current-finalize.json',{'command':['python','source/upstream/skills/image-to-editable-ppt/cli/editppt/cli.py','run','finalize',RUN_NAME],'exit_code':r.returncode,'worker_manifests_unchanged':True,'original_validation_passed':load(run/'final/validation.json')['passed'],'artifact_sha256':sha(out/'deck.pptx'),'new_worker_dispatches':False,'new_image_generation':False,'external_ocr':False,'recorded_worker_execution':'source/round4-orchestration.json; replay does not repeat those real worker dispatches'})
 dump(out/'QA.json',{'status':'generated_awaiting_review','artifact_sha256':sha(out/'deck.pptx'),'previous_review_not_transferred':True})
 print(r.stdout)
if __name__=='__main__':main()
