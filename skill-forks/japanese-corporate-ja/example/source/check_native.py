"""Audit actual editable objects, template parts, data and package integrity."""
from pathlib import Path
import hashlib, json, zipfile, posixpath
from collections import Counter
from lxml import etree
from pptx import Presentation
from PIL import Image, ImageOps, ImageDraw

CASE = Path(__file__).resolve().parents[1]
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def dump(p, value): p.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')

def audit():
    deck = CASE/'deck.pptx'
    template = CASE/'source/upstream/skills/japanese-corporate-pptx/assets/japanese-corporate-ponchi-template.pptx'
    p = Presentation(deck); ec = json.loads((CASE/'input/data.json').read_text('utf-8'))
    rows=[]; out=[]; placeholders=[]; counts=Counter()
    for i,s in enumerate(p.slides):
        row={'index':i,'objects':len(s.shapes),'native_shapes':len(s.shapes),'pictures':0,'native_charts':0,'native_tables':0,'native_connectors':0,'texts':[]}
        for sh in s.shapes:
            row['pictures'] += int(sh.shape_type == 13)
            row['native_charts'] += int(sh.has_chart)
            row['native_tables'] += int(sh.has_table)
            row['native_connectors'] += int(sh._element.tag.endswith('}cxnSp'))
            if sh.has_text_frame: row['texts'].append(sh.text)
            if sh.has_table: row['texts'] += [c.text for r in sh.table.rows for c in r.cells]
            if sh.left < -1270 or sh.top < -1270 or sh.left+sh.width > p.slide_width+1270 or sh.top+sh.height > p.slide_height+1270:
                out.append({'page':i+1,'name':sh.name,'box_emu':[sh.left,sh.top,sh.width,sh.height]})
        for text in row['texts']:
            if '{{' in text or '}}' in text: placeholders.append({'page':i+1,'text':text})
        for k in ['objects','native_shapes','pictures','native_charts','native_tables','native_connectors']: counts[k]+=row[k]
        rows.append(row)
    chart=next(sh.chart for sh in p.slides[2].shapes if sh.has_chart)
    actual={'categories':[c.label for c in chart.plots[0].categories], 'series':[{'name':s.name,'values':list(s.values)} for s in chart.series]}
    expected={'categories':[c['name'] for c in ec['channels']], 'series':[{'name':'2025年9月','values':[c['previous_man_yen'] for c in ec['channels']]},{'name':'2026年9月','values':[c['sales_man_yen'] for c in ec['channels']]}]}
    same=[]; serialization_only=[]; different=[]
    def normalize(data,name):
        root=etree.fromstring(data)
        if name.endswith('.rels'):
            base=posixpath.dirname(posixpath.dirname(name))
            for rel in root:
                if rel.get('TargetMode')!='External':
                    value=rel.get('Target','')
                    rel.set('Target',posixpath.normpath(value.lstrip('/') if value.startswith('/') else posixpath.join(base,value)))
        return etree.tostring(root,method='c14n')
    with zipfile.ZipFile(deck) as z, zipfile.ZipFile(template) as t:
        crc=z.testzip()
        for n in t.namelist():
            if n.startswith(('ppt/slideMasters/','ppt/slideLayouts/','ppt/theme/')) and n.endswith(('.xml','.rels')):
                if n in z.namelist() and z.read(n)==t.read(n): same.append(n)
                elif n in z.namelist() and normalize(z.read(n),n)==normalize(t.read(n),n): serialization_only.append(n)
                else: different.append(n)
        media=[n for n in z.namelist() if n.startswith('ppt/media/')]
    result={'pptx_sha256':sha(deck),'slides':rows,'totals':dict(counts),'chart_data':actual,'chart_matches_input':actual==expected,'package_crc_error':crc,'outside_canvas':out,'unresolved_placeholders':placeholders,'template_parts':{'byte_identical':same,'serialization_only_xml_or_relative_relationship_paths':serialization_only,'semantic_changes_or_missing':different},'media_parts':media,'scope':'OOXML/native-object audit; edit/save/reopen is separately recorded in edit-reopen.json; no aesthetic score.'}
    dump(CASE/'source/native-elements.json',result)
    thumbs=[]
    for i in range(6):
        im=Image.open(CASE/f'render/slide-{i+1:02}.png').convert('RGB');im.thumbnail((640,360))
        tile=Image.new('RGB',(660,400),'#e8edf2');tile.paste(im,((660-im.width)//2,10));ImageDraw.Draw(tile).text((12,378),f'{i+1:02}',fill='#111111');thumbs.append(tile)
    sheet=Image.new('RGB',(1320,1200),'white')
    for i,im in enumerate(thumbs): sheet.paste(im,((i%2)*660,(i//2)*400))
    sheet.save(CASE/'render/contact.png')
    assert len(p.slides)==6 and counts['native_charts']==1 and counts['native_tables']==2 and counts['pictures']==0
    assert actual==expected and crc is None and not out and not placeholders and not different
    assert '10/09の目標' in '\n'.join(rows[4]['texts']) and '分類率100%' in '\n'.join(rows[4]['texts'])
    print(json.dumps({'native':dict(counts),'chart_matches':True,'unchanged_template_parts':len(same),'bounds_violations':len(out),'crc':'passed'}))

if __name__ == '__main__': audit()
