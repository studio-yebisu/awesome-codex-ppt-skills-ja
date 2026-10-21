"""Aham原典の関数・変換器・lintで「朝20分の作業ワークショップ案」を作る。

  python generate.py --sample   P6見本：表紙とP2だけを source/sample/ へ（SVG・PPTX・検査結果）
  python generate.py --final    全4ページを source/final/ へ。input/phase-06-decision.json の
                                decision が "approve" のときだけ動く

ページ内容は input/phase-05-layout-plan.json、数値は input/data.json から読む。
upstream/ は変更しない。原典の出力に加える局所adapterは ADAPTERS に挙げた3件だけで、
色・座標・字号は変えない。検査に1件でも失敗すると終了コード1で止まる。
"""
from pathlib import Path
import argparse, contextlib, datetime, html, io, json, os, re, subprocess, sys

HERE = Path(__file__).resolve().parent
UP = HERE / 'upstream'
sys.path.insert(0, str(UP / 'assets/components'))
sys.path.insert(0, str(UP / 'assets'))
import components as c, themes, layouts, structures, charts
import svg_to_pptx.drawingml_utils as du
from svg_to_pptx_wrapper import svg_to_native_pptx


def load(name):
    return json.loads((HERE / 'input' / name).read_text(encoding='utf-8'))


data = load('data.json')
plan = load('phase-05-layout-plan.json')
PAGES = {p['page']: p for p in plan['pages']}
TOTAL = len(PAGES)
BRAND = plan['brand']

# ---- 局所adapter（記録は source/layout-map.json と同じ） ----
JA_FONT = "Manrope, 'Zen Kaku Gothic New', sans-serif"
FIXED_LABELS = {'>目录　CONTENTS<': '>目次<', '>来源：': '>出典：'}
ADAPTERS = [
    {'id': 'font-family', 'what': '原典の SANS / MONO 字体列を完全一致で JA_FONT に置き換える', 'value': JA_FONT},
    {'id': 'ea-font-registration', 'what': '変換器の和文字体集合 EA_FONTS に実行時だけ Zen Kaku Gothic New を加え、PPTX で latin=Manrope / ea=Zen Kaku Gothic New にする（ファイルは変更しない）'},
    {'id': 'fixed-labels', 'what': '原典の中国語固定文字を和訳する', 'value': FIXED_LABELS},
]


def adapt(svg):
    svg = svg.replace(f'font-family="{c.SANS}"', f'font-family="{JA_FONT}"')
    svg = svg.replace(f'font-family="{c.MONO}"', f'font-family="{JA_FONT}"')
    for old, new in FIXED_LABELS.items():
        svg = svg.replace(old, new)
    return svg


# ---- ページ（原典関数を phase-05 の引数で呼ぶだけ） ----
def content_page(n, body):
    p = PAGES[n]
    return themes.chrome('A', p['section_label'], p['action_title'], plan['footer_source'],
                         n, TOTAL, body, brand=BRAND)


def page1():
    a = PAGES[1]['render_args']
    return themes.cover('A', BRAND, a['title'], a['subtitle'], [tuple(ch) for ch in a['chapters']],
                        meta_left=a['meta_left'], meta_right=a['meta_right'])


def page2():
    a = PAGES[2]['render_args']
    mins = data['durations_minutes']
    if len(mins) != len(a['steps']):
        raise SystemExit(f'P2: data.json の所要時間 {len(mins)} 件と手順 {len(a["steps"])} 件が合わない')
    tasks, start = [], 0
    for i, (name, m) in enumerate(zip(a['steps'], mins)):
        tasks.append((name, start, m, i == a['highlight_index']))
        start += m
    if start != a['total_minutes']:
        raise SystemExit(f'P2: 合計 {start} 分が計画の {a["total_minutes"]} 分と合わない')
    body = charts.gantt(a['x'], a['y'], a['w'], a['h'], tasks, start, unit_label=a['unit_label'])
    nt = a['note']
    body += c.T(nt['x'], nt['y'], nt['text'], 13, c.INK2)
    return content_page(2, body)


def page3():
    a = PAGES[3]['render_args']
    return content_page(3, layouts.stage_flow([tuple(s) for s in a['stages']], dark_last=a['dark_last']))


def page4():
    a = PAGES[4]['render_args']
    lanes = [(ln['name'], [tuple(b) for b in ln['blocks']]) for ln in a['lanes']]
    body = structures.swimlane_roadmap(a['x'], a['y'], a['w'], a['h'], lanes, a['total_units'],
                                       unit_label=a['unit_label'], legend=a['legend'])
    return content_page(4, body)


BUILDERS = {1: page1, 2: page2, 3: page3, 4: page4}

# ---- 変換前の検査 ----
WHITELIST = {'#FFFFFF', '#F3F3F3', '#E7E7E7', '#262626', '#6E6E6E', '#9B9B9B', '#C4C4C4', '#C8C8C8',
             '#336EE8', '#5C8BED', '#164EC3', '#5A7A60', '#8A7333', '#9E3D31', '#2E2E2E', '#333333', '#F7F8FA'}
LAYOUT_ID = re.compile(r'\b(?:SPC(?:-\d{2})?|[SEGIVTA]-\d{2})\b')
LEFTOVER = ['来源', '目录', '本册导览', 'IN THIS DECK', '已承诺', '计划中', '拟议']
TEXT_RE = re.compile(r'<text ([^>]*)>([^<]*)</text>')


def svg_texts(svg):
    return [(attrs, html.unescape(t)) for attrs, t in TEXT_RE.findall(svg)]


def attr(attrs, name, default=None):
    m = re.search(rf'\b{name}="([^"]*)"', attrs)
    return m.group(1) if m else default


def lint_width(s, fs, ls=0.0):
    # 原典 lint_svg.py と同じ見積り（CJK=1em、それ以外0.55em）
    return sum(fs if ord(ch) > 0x2E00 else fs * 0.55 for ch in s) + ls * max(0, len(s) - 1)


def pre_checks(svgs, pages):
    res = []

    def add(name, ok, detail):
        res.append({'check': name, 'ok': bool(ok), 'detail': detail})

    colors = sorted({m.upper() for s in svgs.values() for m in re.findall(r'#[0-9A-Fa-f]{6}\b', s)})
    add('brand_color_whitelist', not [x for x in colors if x not in WHITELIST],
        {'used': colors, 'outside': [x for x in colors if x not in WHITELIST]})
    fams = sorted({m for s in svgs.values() for m in re.findall(r'font-family="([^"]*)"', s)})
    add('font_family_only_ja_font', fams == [JA_FONT], {'used': fams})
    ids, left, over = [], [], []
    for name, s in svgs.items():
        for attrs, t in svg_texts(s):
            if LAYOUT_ID.search(t):
                ids.append(f'{name}: {t}')
            left += [f'{name}: {t}' for w in LEFTOVER if w in t]
            x, fs = float(attr(attrs, 'x', 0)), float(attr(attrs, 'font-size', 16))
            wd = lint_width(t, fs, float(attr(attrs, 'letter-spacing', 0)))
            anc = attr(attrs, 'text-anchor', 'start')
            right = x + (wd if anc == 'start' else wd / 2 if anc == 'middle' else 0)
            if right > 1240:
                over.append(f'{name}: 右端{right:.1f} 「{t}」')
    add('no_layout_id_in_text', not ids, ids)
    add('no_untranslated_fixed_label', not left, left)
    add('text_right_edge_le_1240', not over, over)
    gap = 1240 - lint_width(BRAND, 12) - 1197
    add('chrome_brand_clear_of_dot', gap >= 6, {'brand': BRAND, 'gap_px': round(gap, 1)})
    bad_titles, subjects = [], []
    for n in pages:
        t = PAGES[n].get('action_title')
        if not t:
            continue
        if len(t) > 24 or re.search(r'[0-9０-９]', t) or t.count('、') + t.count(',') > 1:
            bad_titles.append(t)
        subjects.append(t.split('は')[0])
    add('action_title_rules', not bad_titles, {'violations': bad_titles})
    add('subject_alternates', all(a != b for a, b in zip(subjects, subjects[1:])), subjects)
    return res


def run_lint(svg_dir, out_dir, tag, expected):
    cmd = [sys.executable, str(UP / 'assets/tools/lint_svg.py'), str(svg_dir)]
    env = dict(os.environ, PYTHONIOENCODING='utf-8', PYTHONUTF8='1')
    r = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', errors='replace',
                       cwd=str(HERE), env=env)
    (out_dir / f'lint_svg.{tag}.txt').write_text(r.stdout + ('\n[stderr]\n' + r.stderr if r.stderr else ''),
                                                 encoding='utf-8')
    findings = [ln.strip() for ln in r.stdout.splitlines() if ln.strip().startswith('[!]') or '⚠' in ln]
    m = re.search(r'扫描 (\d+) 个 SVG', r.stdout)
    scanned = int(m.group(1)) if m else None
    ok = r.returncode == 0 and scanned == expected and not findings
    return {'check': f'upstream_lint_svg_{tag}', 'ok': ok,
            'detail': {'command': ['python', 'upstream/assets/tools/lint_svg.py', svg_dir.relative_to(HERE).as_posix()],
                       'returncode': r.returncode, 'scanned': scanned, 'expected': expected,
                       'findings': findings, 'stdout_file': f'lint_svg.{tag}.txt',
                       'note': 'lint_svg.py は常に終了コード0を返すため、[!] 行と ⚠ 行の有無で判定する'}}


# ---- 変換と変換後の検査 ----
def convert(paths, out_pptx, notes, out_dir):
    du.EA_FONTS.add('Zen Kaku Gothic New')
    buf = io.StringIO()
    err = None
    try:
        with contextlib.redirect_stdout(buf):
            ok = svg_to_native_pptx(paths, out_pptx, canvas_format='auto', notes=notes)
    except Exception as e:  # 変換器の例外も失敗として記録する
        ok, err = False, f'{type(e).__name__}: {e}'
    log = buf.getvalue()
    (out_dir / 'convert.log').write_text(log + (f'\n[exception] {err}\n' if err else ''), encoding='utf-8')
    print(log)
    warnings = [ln.strip() for ln in log.splitlines() if 'Warning:' in ln]
    errors = [ln.strip() for ln in log.splitlines() if 'Error:' in ln]
    m = re.search(r'Succeeded: (\d+), Failed: (\d+)', log)
    good = bool(ok) and not errors and m is not None and int(m.group(1)) == len(paths) and m.group(2) == '0'
    return {'check': 'upstream_svg_to_native_pptx', 'ok': good,
            'detail': {'returned': bool(ok), 'exception': err, 'succeeded_failed': m.groups() if m else None,
                       'warnings': warnings, 'errors': errors, 'log_file': 'convert.log'}}


def post_checks(pptx_path, svgs, notes):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from lxml import etree

    def walk(shapes):
        for sh in shapes:
            try:
                st = sh.shape_type
            except NotImplementedError:
                st = None
            if st == MSO_SHAPE_TYPE.GROUP:
                yield from walk(sh.shapes)
            else:
                yield sh, st

    res = []
    prs = Presentation(str(pptx_path))
    names = list(svgs)
    res.append({'check': 'slide_count', 'ok': len(prs.slides) == len(names),
                'detail': {'pptx': len(prs.slides), 'svg': len(names)}})
    for slide, name in zip(prs.slides, names):
        shapes = list(walk(slide.shapes))
        pics = sum(1 for _, st in shapes if st == MSO_SHAPE_TYPE.PICTURE)
        ptexts = sorted(sh.text_frame.text for sh, _ in shapes if sh.has_text_frame and sh.text_frame.text.strip())
        stexts = sorted(t for _, t in svg_texts(svgs[name]) if t.strip())
        xml = etree.tostring(slide._element, encoding='unicode')
        ea = re.findall(r'<a:ea typeface="([^"]*)"', xml)
        latin = re.findall(r'<a:latin typeface="([^"]*)"', xml)
        note_ok = True
        if notes.get(name):
            note_ok = slide.has_notes_slide and notes[name][:12] in slide.notes_slide.notes_text_frame.text
        detail = {'shapes': len(shapes), 'pictures': pics,
                  'text_missing_in_pptx': sorted(set(stexts) - set(ptexts)),
                  'text_extra_in_pptx': sorted(set(ptexts) - set(stexts)),
                  'ea_fonts': sorted(set(ea)), 'latin_fonts': sorted(set(latin)), 'notes_ok': note_ok}
        ok = (pics == 0 and stexts == ptexts and set(ea) == {'Zen Kaku Gothic New'}
              and set(latin) == {'Manrope'} and note_ok)
        res.append({'check': f'native_{name}', 'ok': ok, 'detail': detail})
    kb = round(pptx_path.stat().st_size / 1024, 1)
    res.append({'check': 'file_size_info', 'ok': True,
                'detail': {'kb': kb, 'upstream_normal_range_kb': [50, 300], 'in_range': 50 <= kb <= 300}})
    return res


def preview_html(svgs, out_dir):
    rel = os.path.relpath(HERE / 'fonts', out_dir).replace(os.sep, '/')
    faces = (f"@font-face{{font-family:'Zen Kaku Gothic New';src:url('{rel}/ZenKakuGothicNew-Regular.ttf');font-weight:400}}"
             f"@font-face{{font-family:'Zen Kaku Gothic New';src:url('{rel}/ZenKakuGothicNew-Bold.ttf');font-weight:700}}"
             f"@font-face{{font-family:'Manrope';src:url('{rel}/Manrope%5Bwght%5D.ttf');font-weight:200 800}}")
    body = ''.join(f'<section><h2>{n}</h2>{s}</section>' for n, s in svgs.items())
    (out_dir / 'preview.html').write_text(
        '<!doctype html><meta charset="utf-8"><title>preview</title>'
        f'<style>{faces}body{{background:#E7E7E7;margin:24px;font-family:sans-serif}}'
        'section{margin:0 0 32px}svg{width:100%;max-width:1280px;height:auto;background:#fff;display:block}</style>'
        + body, encoding='utf-8')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument('--sample', action='store_true', help='P6見本（表紙とP2）を source/sample/ へ出す')
    g.add_argument('--final', action='store_true', help='P6判断が approve のとき全4ページを source/final/ へ出す')
    args = ap.parse_args()

    if args.final:
        p6 = load('phase-06-decision.json')
        if p6.get('decision') != 'approve':
            raise SystemExit('P6未判断のため停止：統合担当が source/sample/ を見て input/phase-06-decision.json の '
                             f'decision を "approve" にしてから --final を実行する（現在: {p6.get("decision")!r}）')
        pages, out_dir, svg_dir, pptx = list(range(1, TOTAL + 1)), HERE / 'source/final', 'svg_final', 'deck.pptx'
    else:
        pages = load('phase-06-decision.json')['sample_pages']
        out_dir, svg_dir, pptx = HERE / 'source/sample', 'svg', 'sample.pptx'

    out_dir.mkdir(parents=True, exist_ok=True)
    sdir = out_dir / svg_dir
    sdir.mkdir(exist_ok=True)
    for old in sdir.glob('*.svg'):
        old.unlink()
    svgs = {f'slide-{n:02}': adapt(BUILDERS[n]()) for n in pages}
    paths = []
    for name, s in svgs.items():
        p = sdir / f'{name}.svg'
        p.write_text(s, encoding='utf-8')
        paths.append(p)
    preview_html(svgs, out_dir)

    checks = pre_checks(svgs, pages)
    checks.append(run_lint(sdir, out_dir, 'before_convert', len(paths)))
    notes = {f'slide-{n:02}': PAGES[n]['speaker_note'] for n in pages}
    out_pptx = out_dir / pptx
    if out_pptx.exists():
        out_pptx.unlink()
    conv = convert(paths, out_pptx, notes, out_dir)
    checks.append(conv)
    if conv['ok'] and out_pptx.exists():
        checks += post_checks(out_pptx, svgs, notes)
    else:
        checks.append({'check': 'native_post_checks', 'ok': False, 'detail': '変換に失敗したため未実施'})
    checks.append(run_lint(sdir, out_dir, 'after_convert', len(paths)))

    failed = [x['check'] for x in checks if not x['ok']]
    report = {
        'mode': 'final' if args.final else 'sample',
        'generated_at': datetime.datetime.now().astimezone().isoformat(timespec='seconds'),
        'pages': pages,
        'outputs': {'svg_dir': sdir.relative_to(HERE).as_posix(), 'pptx': out_pptx.relative_to(HERE).as_posix(),
                    'preview': (out_dir / 'preview.html').relative_to(HERE).as_posix()},
        'adapters': ADAPTERS,
        'status': 'pass' if not failed else 'fail',
        'failed_checks': failed,
        'checks': checks,
        'not_performed_here': ['PowerPoint / LibreOffice での描画と目視', '原典 Phase 7 の3ページごとの確認',
                               '原典 Phase 8 の五層QCと人による内容・視覚監査'],
    }
    (out_dir / 'checks.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f'[{report["status"]}] {report["mode"]}: {report["outputs"]["pptx"]}  失敗: {failed or "なし"}')
    if failed:
        raise SystemExit(1)


if __name__ == '__main__':
    main()
