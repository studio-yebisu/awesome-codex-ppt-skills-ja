"""Case-specific native authoring following the supplied PradaSlides contracts."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
ROOT=Path(__file__).resolve().parents[1]
NAVY='12324A'; PAPER='F5F3ED'; WHITE='FFFFFF'; ACCENT='B7472A'; MUTED='435B6A'
def main():
 data=json.loads((ROOT/'input/data.json').read_text('utf-8')); prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
 def text(s,x,y,w,h,value,size=22,color=NAVY,bold=False):
  box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=0;tf.margin_top=tf.margin_bottom=0
  for i,line in enumerate(value.split('\n')):
   p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Yu Gothic';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(8)
  return box
 def rect(s,x,y,w,h,color):
  sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(color);sh.line.fill.background();return sh
 for i,p in enumerate(data['pages']):
  s=prs.slides.add_slide(prs.slide_layouts[6]);dark=p['tone']=='dark';fg=WHITE if dark else NAVY;s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(NAVY if dark else PAPER)
  text(s,.6,.25,11,.3,f'町内会の備品貸出  /  {i+1:02d}',12,fg)
  text(s,.6,7.02,12,.25,data['notice']+'  ・  実際のルールは担当窓口へ',11,fg)
  s.notes_slide.notes_text_frame.text=p['notes']
  if i==0:
   text(s,.75,1.12,8.1,1.8,'町内会の備品貸出を、\n迷わない手順に',43,WHITE,True)
   rect(s,.75,3.5,.09,1.13,ACCENT);text(s,1.05,3.48,7.3,1.25,'借りる前に、\n返すところまで決める。',30,WHITE,True)
   text(s,.78,5.75,7.8,.65,'申込み → 受取り → 使用 → 返却',23,WHITE)
   rect(s,9.1,2.02,3.35,4.32,WHITE);text(s,9.36,2.3,2.85,.5,'返却までの記録',21,NAVY,True)
   for j,(label,copy) in enumerate([('貸出番号','________________'),('受取数 / 返却数','______ / ______'),('返す日時','____月____日____時'),('担当者の確認','□ 返却確認済み')]):
    y=3.0+j*.72;text(s,9.36,y,2.83,.3,label,15,ACCENT,True);text(s,9.36,y+.31,2.83,.35,copy,17,NAVY)

  elif i==1:
   text(s,.6,.93,12,1.25,p['title'],35,fg,True);rect(s,.75,2.24,11.85,3.72,WHITE)
   text(s,1,2.39,10,.45,'申込みの記入例',23,ACCENT,True)
   for j,(k,v) in enumerate(data['form_example'].items()):
    y=3+j*.66;text(s,1,y,2.9,.47,k,20,MUTED,True);text(s,4.04,y,8,.5,v,22);rect(s,1,y+.52,11.2,.012,'CBD4D5')
   text(s,.85,6.28,11.8,.43,'受付方法と貸出条件は、実際の窓口で確認する。',21)
  elif i==2:
   text(s,.6,.93,12,1.3,p['title'],35,fg,True)
   for j,(title,body) in enumerate([('品名と数を照合','申込書と現物の\n数量・付属品を比べる'),('貸出前の状態','傷・汚れ・欠品を\n担当者と記録する'),('受取記録を保存','受取数・状態・担当者を\n返却時まで残す')]):
    x=.75+j*4.08;rect(s,x,2.75,3.56,.045,ACCENT);text(s,x,2.08,3.2,.6,f'0{j+1}',30,ACCENT,True);text(s,x,3.18,3.8,.6,title,29,NAVY,True);text(s,x,4.2,3.8,1.2,body,23)
   text(s,.75,6.18,12,.45,'受取記録を、返却時の照合に使う。',25,NAVY,True)
  elif i==3:
   text(s,.6,.93,12,1.3,p['title'],34,fg,True)
   for j,(title,body) in enumerate([('受取りのとき','傷や欠品があった'),('使っているとき','ぐらつきや破損に気づいた')]):
    x=.8+j*6.35;text(s,x,2.26,5.8,.6,title,27,WHITE,True);text(s,x,2.96,5.8,.65,body,22,WHITE)
   rect(s,.8,3.84,11.7,.02,'7793A4');text(s,.8,4.05,11.8,.65,'使用を止め、窓口へ。扱いを確認する。',28,WHITE,True)
   rect(s,.8,5.02,11.7,1.59,WHITE);text(s,1.03,5.13,11.1,.34,'窓口へ伝える記録',18,ACCENT,True)
   text(s,1.03,5.65,5.3,.38,'品名・数量  ______________',20,NAVY);text(s,6.8,5.65,5.4,.38,'気づいた時刻  ____________',20,NAVY)
   text(s,1.03,6.13,11.1,.38,'状態・窓口の指示  ___________________________________',20,NAVY)
  elif i==4:
   text(s,.6,.93,12,1.3,p['title'],35,fg,True)
   for j,(a,b) in enumerate([('数量','借りた数と、返す数が合うか'),('付属品','袋・部品などの戻し忘れはないか'),('状態','使用中の変化を担当者へ伝えたか'),('確認','担当者の確認を記録したか')]):
    y=2.4+j*.88;rect(s,.8,y,2.05,.74,NAVY);label=text(s,1.02,y,1.7,.74,a,23,WHITE,True);label.text_frame.vertical_anchor=MSO_ANCHOR.MIDDLE;label.text_frame.paragraphs[0].space_after=Pt(0);text(s,3.17,y+.11,9.3,.5,b,24);rect(s,3.15,y+.73,9.2,.015,'9DAEB4')
   text(s,.85,6.32,12,.45,'無人で置いてよいかも、あらかじめ確認する。',21)
  else:
   text(s,.75,1.03,12,1.75,'借りる前に、\n返却方法まで聞こう。',43,NAVY,True)
   for j,(a,b) in enumerate([('誰に','申込みと返却の窓口'),('いつ・どこで','受取りと返却の日時・場所'),('変わったら','予定変更や不具合の連絡先')]):
    y=3.28+j*.87;text(s,.85,y,3.05,.6,a,24,ACCENT,True);text(s,4.02,y,8.4,.6,b,26)
   text(s,.8,6.24,12,.48,'実際の貸出ルールを確認してから、申し込む。',24,NAVY,True)
 (ROOT/'output').mkdir(exist_ok=True);prs.save(ROOT/'output/deck.pptx')
if __name__=='__main__':main()
