"""Self-authored 16:9 workshop template for the agent-slides showcase.

Authors two named layouts whose placeholders carry the deck's type styles, so
original set_semantic_text ops inherit font, size, colour and position with no
explicit overrides in ops. Background, side panel and the English label are
layout-owned, not shape ops. Depends only on python-pptx; no third-party
branding or photos. The two sample slides are design evidence, not the deck.

Usage from the main generator: build_template(Path) -> JSON-friendly inventory.
Reproduction: python build_styled_template.py <target.pptx>
"""
from pathlib import Path
from xml.sax.saxutils import escape
import argparse, json, sys
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.oxml import serialize_part_xml
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu

NAVY='102D45'; ORANGE='FF6A13'; WHITE='FFFFFF'; PALE='D9E7F1'
JP='Zen Kaku Gothic New'; EN='Manrope'
COVER='Workshop Cover'; SPLIT='Workshop Split'
LABEL='A DAY TO MAKE / CONCEPT PROPOSAL'
SLIDE_W=12192000; SLIDE_H=6858000  # true 16:9, 13.333 x 7.5 in
THEME_COLORS={'dk1':NAVY,'lt1':WHITE,'accent1':ORANGE}

def _emu(v):return int(round(v*914400))
def _box(x,y,w,h):return tuple(_emu(v) for v in (x,y,w,h))
def _inches(v):return round(Emu(int(v)).inches,3)

# Type specs. Japanese family on latin/ea/cs; English labels keep ea on Zen Kaku.
SPLIT_TITLE=dict(size=52,bold=True,color=WHITE,font=JP)
SPLIT_BODY=dict(size=23,bold=False,color=WHITE,font=JP)
COVER_TITLE=dict(size=48,bold=True,color=WHITE,font=JP)
COVER_SUB=dict(size=20,bold=False,color=PALE,font=JP)
LABEL_SPEC=dict(size=12,bold=False,color=PALE,font=EN)

# ---- XML fragments -------------------------------------------------------

def _rpr(tag,spec,extra=''):
    return (f'<a:{tag}{extra} sz="{spec["size"]*100}" b="{int(spec["bold"])}">'
            f'<a:solidFill><a:srgbClr val="{spec["color"]}"/></a:solidFill>'
            f'<a:latin typeface="{spec["font"]}"/><a:ea typeface="{JP}"/><a:cs typeface="{spec["font"]}"/></a:{tag}>')

def _xfrm(box):
    x,y,w,h=box
    return f'<a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>'

def _body_pr(anchor,autofit):
    return (f'<a:bodyPr vert="horz" wrap="square" lIns="91440" tIns="45720" rIns="91440" bIns="45720" rtlCol="0" anchor="{anchor}">'
            f'<a:{autofit}/></a:bodyPr>')

def _lvl(n,spec,algn):
    # No bullets, no indent, single spacing: matches the original add_text boxes.
    return (f'<a:lvl{n}pPr marL="{_emu(.3*(n-1))}" indent="0" algn="{algn}">'
            '<a:lnSpc><a:spcPct val="100000"/></a:lnSpc><a:spcBef><a:spcPts val="0"/></a:spcBef><a:buNone/>'
            f'{_rpr("defRPr",spec)}</a:lvl{n}pPr>')

def _tx_body(spec,levels,algn,anchor,autofit,prompt):
    lst=''.join(_lvl(n,spec,algn) for n in range(1,levels+1))
    # p.font on the prompt paragraph mirrors lvl1, so the layout itself previews correctly.
    paras=''.join(f'<a:p><a:pPr>{_rpr("defRPr",spec)}</a:pPr><a:r><a:rPr lang="ja-JP" altLang="en-US"/>'
                  f'<a:t>{escape(line)}</a:t></a:r></a:p>' for line in prompt)
    return f'<p:txBody {nsdecls("a","p")}>{_body_pr(anchor,autofit)}<a:lstStyle>{lst}</a:lstStyle>{paras}</p:txBody>'

def _rect(sid,name,box,color):
    return parse_xml(f'<p:sp {nsdecls("a","p")}><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>'
                     f'<p:spPr>{_xfrm(box)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
                     f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>')

def _label(sid,name,box,text,spec,algn):
    run_pr=_rpr('rPr',spec,' lang="en-US" dirty="0"')
    return parse_xml(f'<p:sp {nsdecls("a","p")}><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr>'
                     f'<p:spPr>{_xfrm(box)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
                     f'<p:txBody>{_body_pr("t","noAutofit")}<a:lstStyle/><a:p><a:pPr algn="{algn}"/>'
                     f'<a:r>{run_pr}<a:t>{escape(text)}</a:t></a:r></a:p></p:txBody></p:sp>')

# ---- template authoring --------------------------------------------------

def _widen(element,ratio):
    # Default python-pptx master/layouts are 4:3; stretch their horizontal geometry to 16:9.
    for xfrm in element.iter(qn('a:xfrm')):
        off,ext=xfrm.find(qn('a:off')),xfrm.find(qn('a:ext'))
        if off is not None:off.set('x',str(round(int(off.get('x'))*ratio)))
        if ext is not None:ext.set('cx',str(round(int(ext.get('cx'))*ratio)))

def _update_theme(prs):
    part=prs.slide_master.part.part_related_by(RT.THEME)
    xml_part=hasattr(part,'_element')
    root=part._element if xml_part else parse_xml(part.blob)
    root.set('name','Workshop Theme')
    scheme=root.find('.//'+qn('a:clrScheme'));scheme.set('name','Workshop Palette')
    for key,value in THEME_COLORS.items():
        slot=scheme.find(qn('a:'+key))
        for child in list(slot):slot.remove(child)
        slot.append(parse_xml(f'<a:srgbClr {nsdecls("a")} val="{value}"/>'))
    fonts=root.find('.//'+qn('a:fontScheme'));fonts.set('name','Workshop Type')
    for kind in ('a:majorFont','a:minorFont'):
        group=fonts.find(qn(kind))
        group.find(qn('a:latin')).set('typeface',JP);group.find(qn('a:ea')).set('typeface',JP)
        for font in group.findall(qn('a:font')):
            if font.get('script')=='Jpan':font.set('typeface',JP)
    if not xml_part:part._blob=serialize_part_xml(root)

def _prepare_layout(layout,name,background):
    layout._element.find(qn('p:cSld')).set('name',name)
    for ph in [p for p in layout.placeholders if p.placeholder_format.type in (PP_PLACEHOLDER.DATE,PP_PLACEHOLDER.FOOTER,PP_PLACEHOLDER.SLIDE_NUMBER)]:
        ph._element.getparent().remove(ph._element)
    c_sld=layout._element.find(qn('p:cSld'))
    old=c_sld.find(qn('p:bg'))
    if old is not None:c_sld.remove(old)
    c_sld.insert(0,parse_xml(f'<p:bg {nsdecls("a","p")}><p:bgPr><a:solidFill><a:srgbClr val="{background}"/></a:solidFill><a:effectLst/></p:bgPr></p:bg>'))
    return {p.placeholder_format.idx:p._element for p in layout.placeholders}

def _restyle(sp,name,box,spec,algn,anchor,autofit,prompt,levels=1,ph_type=None):
    sp.find(f'{qn("p:nvSpPr")}/{qn("p:cNvPr")}').set('name',name)
    ph=sp.find('.//'+qn('p:ph'))
    if ph_type:ph.set('type',ph_type)
    ph.set('hasCustomPrompt','1')
    sp_pr=sp.find(qn('p:spPr'))
    for old in sp_pr.findall(qn('a:xfrm')):sp_pr.remove(old)
    sp_pr.insert(0,parse_xml(_xfrm(box).replace('<a:xfrm>',f'<a:xfrm {nsdecls("a")}>',1)))
    new=parse_xml(_tx_body(spec,levels,algn,anchor,autofit,prompt))
    old=sp.find(qn('p:txBody'))
    if old is not None:old.addprevious(new);sp.remove(old)
    else:sp.append(new)

def _insert_fixed(layout,builders):
    # Fixed shapes go directly after grpSpPr so they sit behind every placeholder.
    tree=layout._element.find(qn('p:cSld')).find(qn('p:spTree'))
    pos=list(tree).index(tree.find(qn('p:grpSpPr')))+1
    sid=max(int(e.get('id')) for e in layout._element.iter(qn('p:cNvPr')))+1
    for offset,build in enumerate(builders):tree.insert(pos+offset,build(sid+offset))

def _style_split(layout):
    phs=_prepare_layout(layout,SPLIT,NAVY)
    _restyle(phs[0],'Workshop Title',_box(.58,.9,9.1,2.28),SPLIT_TITLE,'l','t','noAutofit',['見出しを入力'])
    # Standard OBJECT content placeholder becomes BODY, keeping idx 1.
    _restyle(phs[1],'Workshop Body',_box(.65,3.35,8.72,.95),SPLIT_BODY,'l','t','normAutofit',['本文を入力'],levels=5,ph_type='body')
    panel_x=_emu(10.02)
    _insert_fixed(layout,[lambda i:_rect(i,'Workshop Side Panel',(panel_x,0,SLIDE_W-panel_x,SLIDE_H),ORANGE),
                          lambda i:_label(i,'Workshop Label',_box(.62,.33,8.9,.34),LABEL,LABEL_SPEC,'l')])

def _style_cover(layout):
    phs=_prepare_layout(layout,COVER,NAVY)
    _restyle(phs[0],'Workshop Cover Title',_box(.9,2.2,11.533,2.1),COVER_TITLE,'ctr','b','noAutofit',['表紙の見出しを入力'])
    _restyle(phs[1],'Workshop Cover Subtitle',_box(1.667,4.75,10.0,1.0),COVER_SUB,'ctr','t','normAutofit',['補足の一文を入力'])
    _insert_fixed(layout,[lambda i:_rect(i,'Workshop Cover Band',(0,_emu(7.1),SLIDE_W,SLIDE_H-_emu(7.1)),ORANGE),
                          lambda i:_rect(i,'Workshop Cover Rule',_box(6.0665,4.47,1.2,.07),ORANGE),
                          lambda i:_label(i,'Workshop Cover Label',_box(2.2165,.62,8.9,.34),LABEL,LABEL_SPEC,'ctr')])

def build_template(target):
    target=Path(target);target.parent.mkdir(parents=True,exist_ok=True)
    prs=Presentation()
    ratio=SLIDE_W/prs.slide_width
    prs.slide_width=Emu(SLIDE_W);prs.slide_height=Emu(SLIDE_H)
    _widen(prs.slide_master._element,ratio)
    for layout in prs.slide_layouts:_widen(layout._element,ratio)
    _update_theme(prs)
    cover,split=prs.slide_layouts[0],prs.slide_layouts[1]
    _style_cover(cover);_style_split(split)
    # Sample slides: placeholder text only, every style inherited from the layout.
    for layout,title,body in [(cover,'ワークショップ企画書\nテンプレート見本','表紙レイアウトの書式見本。架空の企画用。'),
                              (split,'見出しは、\n二行で大きく。','書体・色・位置は\nレイアウトから継承する見本。')]:
        slide=prs.slides.add_slide(layout)
        slide.shapes.title.text_frame.text=title
        slide.placeholders[1].text_frame.text=body
    prs.save(str(target))
    return inventory(target)

# ---- inventory (read back from the saved file) ---------------------------

def _fill(sp):
    clr=sp.find(f'{qn("p:spPr")}/{qn("a:solidFill")}/{qn("a:srgbClr")}')
    return None if clr is None else clr.get('val')

def _geometry(shape):
    return dict(left=_inches(shape.left),top=_inches(shape.top),width=_inches(shape.width),height=_inches(shape.height))

def _style(rpr):
    if rpr is None:return None
    clr=rpr.find('.//'+qn('a:srgbClr'));latin=rpr.find(qn('a:latin'));ea=rpr.find(qn('a:ea'))
    return dict(font=None if latin is None else latin.get('typeface'),ea_font=None if ea is None else ea.get('typeface'),
                size_pt=int(rpr.get('sz'))/100 if rpr.get('sz') else None,bold=rpr.get('b')=='1',color=None if clr is None else clr.get('val'))

def _layout_record(index,layout):
    bg=layout._element.find(f'{qn("p:cSld")}/{qn("p:bg")}')
    bg_clr=None if bg is None else bg.find('.//'+qn('a:srgbClr'))
    placeholders=[];fixed=[]
    for shape in layout.shapes:
        sp=shape._element
        if shape.is_placeholder:
            lst=sp.find(f'{qn("p:txBody")}/{qn("a:lstStyle")}');lvl1=lst.find(qn('a:lvl1pPr'))
            body_pr=sp.find(f'{qn("p:txBody")}/{qn("a:bodyPr")}')
            placeholders.append(dict(idx=shape.placeholder_format.idx,type=str(shape.placeholder_format.type),name=shape.name,
                                     **_geometry(shape),style=_style(lvl1.find(qn('a:defRPr'))),align=lvl1.get('algn'),
                                     anchor=body_pr.get('anchor'),styled_levels=len(lst),
                                     prompt_paragraph_font=_style(sp.find('.//'+qn('a:pPr')+'/'+qn('a:defRPr')))))
        else:
            # Check txBody directly: shape.text_frame would add one to plain rectangles.
            text=''.join(t.text or '' for t in sp.iter(qn('a:t'))) if sp.find(qn('p:txBody')) is not None else None
            fixed.append(dict(name=shape.name,kind='text' if text else 'rectangle',**_geometry(shape),fill=_fill(sp),
                              text=text,style=_style(sp.find('.//'+qn('a:rPr')))))
    return dict(index=index,name=layout.name,background=None if bg_clr is None else bg_clr.get('val'),
                shape_order=[s.name for s in layout.shapes],fixed_shapes=fixed,placeholders=placeholders)

def _explicit_overrides(sp):
    if sp.find(f'{qn("p:spPr")}/{qn("a:xfrm")}') is not None:return True
    for rpr in sp.iter(qn('a:rPr'),qn('a:defRPr'),qn('a:endParaRPr')):
        if rpr.get('sz') or rpr.get('b') or len(rpr):return True
    return False

def inventory(target):
    prs=Presentation(str(target))
    root=parse_xml(prs.slide_master.part.part_related_by(RT.THEME).blob)
    palette={}
    for slot in root.find('.//'+qn('a:clrScheme')):
        clr=slot[0];palette[slot.tag.split('}')[1]]=clr.get('val') if clr.tag==qn('a:srgbClr') else clr.get('lastClr',clr.get('val'))
    font_scheme=root.find('.//'+qn('a:fontScheme'))
    theme_fonts={kind:dict(latin=font_scheme.find(f'{qn("a:"+kind)}/{qn("a:latin")}').get('typeface'),
                           ea=font_scheme.find(f'{qn("a:"+kind)}/{qn("a:ea")}').get('typeface')) for kind in ('majorFont','minorFont')}
    layouts=list(prs.slide_layouts)
    samples=[dict(slide_number=n,layout=s.slide_layout.name,
                  placeholders=[dict(idx=p.placeholder_format.idx,type=str(p.placeholder_format.type),text=p.text_frame.text,
                                     explicit_overrides=_explicit_overrides(p._element)) for p in s.placeholders])
             for n,s in enumerate(prs.slides,1)]
    return dict(path=str(Path(target)),self_authored=True,third_party_assets=False,dependency='python-pptx',
                slide_size_in=dict(width=_inches(prs.slide_width),height=_inches(prs.slide_height)),
                palette=dict(theme=palette,requested=dict(dk1=NAVY,lt1=WHITE,accent1=ORANGE),extra_text=PALE),
                theme_fonts=theme_fonts,layout_names=[l.name for l in layouts],
                styled_layouts=[_layout_record(i,l) for i,l in enumerate(layouts) if l.name in (COVER,SPLIT)],
                fixed_label=dict(layout=SPLIT,text=LABEL,note='Layout-owned; render ops may omit it.'),
                sample_slides=samples,
                sample_note='Template design evidence only; not the final event deck.')

def main(argv=None):
    p=argparse.ArgumentParser(description='Build the self-authored Workshop Cover / Workshop Split template.')
    p.add_argument('target',type=Path,help='output .pptx path')
    args=p.parse_args(argv)
    if hasattr(sys.stdout,'reconfigure'):sys.stdout.reconfigure(encoding='utf-8')
    print(json.dumps(build_template(args.target),ensure_ascii=False,indent=2))

if __name__=='__main__':main()
