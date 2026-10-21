"""Inspect actual native elements; this does not score slide aesthetics."""
import collections
import json
from pathlib import Path
from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
prs = Presentation(ROOT / "deck.pptx")
slides = []
for index, slide in enumerate(prs.slides):
    shapes = []
    for shape in slide.shapes:
        geometry = shape._element.find(".//" + qn("a:prstGeom"))
        shapes.append({"id": shape.shape_id, "name": shape.name,
                       "type": str(shape.shape_type),
                       "bounds_inches": {"left": round(shape.left/914400,4), "top": round(shape.top/914400,4),
                                          "width": round(shape.width/914400,4), "height": round(shape.height/914400,4)},
                       "preset": geometry.get("prst") if geometry is not None else None,
                       "text": shape.text if shape.has_text_frame else None,
                       "has_native_chart": shape.has_chart,
                       "has_native_table": shape.has_table})
    slides.append({"index": index, "shape_count": len(shapes),
                   "picture_count": sum(s["type"].startswith("PICTURE") for s in shapes),
                   "block_arc_count": sum(s["preset"] == "blockArc" for s in shapes),
                   "text_shape_count": sum(bool(s["text"]) for s in shapes),
                   "office_chart_count": sum(s["has_native_chart"] for s in shapes),
                   "office_table_count": sum(s["has_native_table"] for s in shapes),
                   "shapes": shapes})
font_data = json.loads((ROOT / "input/oss-font-kit.json").read_text(encoding="utf-8"))
counts = collections.Counter(x["language"] for x in font_data["families"])
result = {"scope": "PPTX native-object and data inspection; not visual approval",
          "slide_count": len(slides), "font_family_count": len(font_data["families"]),
          "language_counts": dict(counts), "slides": slides,
          "editability": "Donut arcs, legends, table text/rules and vertical-step text/nodes are native shapes. They are not data-linked Office chart/table objects."}
(ROOT / "native-elements.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print(json.dumps({"slide_count": len(slides), "font_family_count": len(font_data["families"]),
                  "block_arcs": [s["block_arc_count"] for s in slides],
                  "pictures": [s["picture_count"] for s in slides]}, ensure_ascii=False))
