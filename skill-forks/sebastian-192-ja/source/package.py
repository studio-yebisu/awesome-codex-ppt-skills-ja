import pathlib,json,shutil,hashlib,zipfile,html
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
ROOT=pathlib.Path(__file__).resolve().parent.parent; W=ROOT/'work'; D=ROOT/'delivery'; E=D/'examples/repositories/sebastian-192'
spec=json.loads((W/'page-spec.json').read_text()); pages=json.loads((W/'content.json').read_text())
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
checks=[]
for d,s in zip(pages,spec['slides']):
 f=W/s['image_file']; im=Image.open(f); width,height=im.size; assert abs(height-width*9/16)<=1
 slide=prs.slides.add_slide(prs.slide_layouts[6]);slide.shapes.add_picture(str(f),0,0,width=prs.slide_width,height=prs.slide_height);slide.notes_slide.notes_text_frame.text=d['speaker_notes'];shutil.copy2(f,E/'pages'/f.name)
 s['image_status']='approved';s['prompt']=(W/f'prompts/page-{d["number"]:02}.txt').read_text();s['image_sha256']=hashlib.sha256(f.read_bytes()).hexdigest()
 checks.append({'id':s['id'],'dimensions':[width,height],'sha256':s['image_sha256'],'visual_review':'初回画像を目視。見出し・主要本文・主要ラベル・概念/未実施注記・切れ/重なりを確認。','scope':'AI文字の完全OCR一致検査ではない','issues':[]})
prs.save(E/'deck.pptx');r=Presentation(E/'deck.pptx');assert len(r.slides)==8
for i,d in enumerate(pages):assert r.slides[i].notes_slide.notes_text_frame.text==d['speaker_notes'];assert len(r.slides[i].shapes)==1
with zipfile.ZipFile(E/'deck.pptx') as z:
 media=[x for x in z.namelist() if x.startswith('ppt/media/')];assert len(media)==8
 for f in (E/'pages').glob('*.png'):assert any(z.read(n)==f.read_bytes() for n in media)
(W/'page-spec.json').write_text(json.dumps(spec,ensure_ascii=False,indent=2));shutil.copy2(W/'page-spec.json',E/'input/page-spec.json');shutil.copy2(W/'deck-spec.md',E/'input/deck-spec.md');shutil.copy2(W/'content.json',E/'input/data.json');shutil.copy2('/workspace/shared/slide-next-inputs/204-summer-shade.md',E/'input/BRIEF.md');shutil.copy2(ROOT/'upstream/LICENSE',E/'LICENSE');shutil.copytree(W/'prompts',E/'source/prompts',dirs_exist_ok=True);shutil.copytree(W/'styles',E/'source/styles',dirs_exist_ok=True);shutil.copy2(W/'style-options.json',E/'source/style-options.json');shutil.copy2(__file__,E/'source/package.py')
# Preserve original textual workflow/documentation without executing its programs.
import subprocess
tracked=subprocess.check_output(['git','ls-files'],cwd=ROOT/'upstream',text=True).splitlines()
for name in tracked:
 src=ROOT/'upstream'/name
 if src.is_file() and src.suffix.lower() in ['.md','.json','.py','.txt','.yml','.yaml','.toml','.sh'] or name=='LICENSE':
  dst=E/'source/upstream'/name;dst.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(src,dst)
size=(E/'deck.pptx').stat().st_size
qa={'status':'independent_minimal_check_complete','upstream_full_qa':'not_run','page_count':8,'slide_ratio':'16:9 with PNG rounding within 1 pixel','image_only_slide_objects':True,'speaker_notes_exact_match':True,'embedded_media_byte_match':True,'pptx_bytes':size,'image_backend':'host-native-imagegen-model-not-exposed','image_calls':{'style_overviews':4,'final_pages':8,'repairs':0},'review':checks,'limitations':['overview内部縮略頁の16:9要件は未達。比較用の4案を保持。','PPTX実レンダラーは未使用。プレビューは完成PPTX内とバイト一致した生成PNG。','原典validator/packer/renderおよび原典全QAは未実行。','PPTX本体は画像主体。文字・図形はラスタで、話者ノートのみ編集可能。','03/08など図内の補助注釈はAI生成、仮説・想定/未実施境界内の説明。完全な文字OCR照合は未実施。']}
(E/'QA.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2)); prov={'repository':'helloo1568/slidemuse','upstream_commit':'7034412b4c19915c17f8d7206e3723e9e53cf8d2','license':'MIT','source_url':'https://github.com/helloo1568/slidemuse/tree/7034412b4c19915c17f8d7206e3723e9e53cf8d2','backend':'host native imagegen; model not exposed','selected_style':3,'selection':'親がユーザー委任に基づき選択','method':'原典の内容稿→同内容4独立overview→選択→Page Spec→逐頁AI PNGの手順。独立python-pptx包装と最小構造チェックへ変更。','rendering':'No completed-PPTX raster renderer used; original PNG previews exactly match embedded bytes.','image_processing':'No generated image edits or resizing. Images embedded directly, placement scaling only.','qa_limitations':qa['limitations'],'sample_notice':'架空・未実施デザイン研究。測定温度・効果・実績なし。'};(E/'PROVENANCE.json').write_text(json.dumps(prov,ensure_ascii=False,indent=2))
body=''.join(f'<section><h2>{d["number"]:02}｜{html.escape(d["title"])}</h2><img src="{d["number"]:02}-slide.png" alt="{html.escape(d["title"])}"><p>{html.escape(d["speaker_notes"])}</p></section>' for d in pages)
(E/'pages/index.html').write_text('<!doctype html><html lang="ja"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>夏の余白｜8頁</title><style>body{margin:24px auto;max-width:1400px;background:#faf8f2;color:#1e2f22;font-family:sans-serif}section{margin:40px 0}img{display:block;width:100%}p{max-width:70em;line-height:1.7}</style><h1>夏の余白</h1><p>架空・未実施のデザイン研究。画像主体PPTX、話者ノート編集可。プレビューはPPTX内の画像とバイト一致した生成PNGで、PPTXレンダリングではありません。</p><a href="../deck.pptx">PowerPoint</a>'+body+'</html>')
print('PPTX bytes',size,'MiB',size/1024**2)
