import importlib.metadata
import importlib.util
import json
import os
from pathlib import Path, PurePosixPath
import subprocess
import sys
from types import SimpleNamespace
import uuid
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches
import pytest


ROOT = Path(__file__).resolve().parents[1]
INSPECTOR = ROOT / "scripts" / "inspect_pptx.py"
QUALITY_GATES = ROOT / "references" / "quality-gates.md"


def run_inspector(
    deck: Path,
    *extra: str,
    cwd: Path = ROOT,
) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(INSPECTOR), str(deck), "--json", *extra],
        cwd=cwd,
        capture_output=True,
        text=True,
        check=False,
    )


def result_json(result: subprocess.CompletedProcess[str]) -> dict:
    assert result.stderr == ""
    return json.loads(result.stdout)


def load_inspector_module():
    name = f"paper2ppt_inspector_{uuid.uuid4().hex}"
    spec = importlib.util.spec_from_file_location(name, INSPECTOR)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def rewrite_package(
    path: Path,
    *,
    remove: set[str] | None = None,
    replacements: dict[str, bytes] | None = None,
    additions: dict[str, bytes] | None = None,
) -> None:
    removed = remove or set()
    replaced = replacements or {}
    added = additions or {}
    temporary = path.with_name(f".{path.name}.rewrite")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(temporary, "w") as target:
        for entry in source.infolist():
            if entry.filename in removed or entry.filename in replaced:
                continue
            target.writestr(entry, source.read(entry))
        for name, data in replaced.items():
            target.writestr(name, data)
        for name, data in added.items():
            target.writestr(name, data)
    os.replace(temporary, path)


def package_part(path: Path, prefix: str) -> str:
    with zipfile.ZipFile(path) as package:
        return next(name for name in package.namelist() if name.startswith(prefix))


def replace_package_text(path: Path, part: str, transform) -> None:
    with zipfile.ZipFile(path) as package:
        original = package.read(part)
    replacement = transform(original)
    assert replacement != original
    rewrite_package(path, replacements={part: replacement})


def corrupt_stored_member_crc(path: Path, member: str) -> None:
    temporary = path.with_name(f".{path.name}.stored")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_STORED
    ) as target:
        for entry in source.infolist():
            target.writestr(entry.filename, source.read(entry))
    os.replace(temporary, path)

    with zipfile.ZipFile(path) as package:
        entry = package.getinfo(member)
    archive = bytearray(path.read_bytes())
    header = entry.header_offset
    filename_length = int.from_bytes(archive[header + 26 : header + 28], "little")
    extra_length = int.from_bytes(archive[header + 28 : header + 30], "little")
    data_offset = header + 30 + filename_length + extra_length
    archive[data_offset] ^= 0x01
    path.write_bytes(archive)


def add_meaningful_notes(slide, text: str = "Evidence, limitation, and transition.") -> None:
    slide.notes_slide.notes_text_frame.text = text


def make_deck(
    path: Path,
    *,
    slides: int = 1,
    notes: bool = True,
    image: Path | None = None,
) -> Presentation:
    presentation = Presentation()
    for index in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(
            Inches(0.75), Inches(0.75), Inches(3), Inches(0.75)
        ).text = f"Slide {index + 1}"
        if notes:
            add_meaningful_notes(slide, f"Evidence for slide {index + 1}.")
        if image is not None and index == 0:
            slide.shapes.add_picture(
                str(image), Inches(1), Inches(2), width=Inches(1)
            )
    presentation.save(path)
    return presentation


def test_valid_pptx_reports_slides_media_notes_and_bounds(tmp_path: Path) -> None:
    image = tmp_path / "figure.png"
    Image.new("RGB", (32, 24), "white").save(image)
    deck = tmp_path / "deck.pptx"
    make_deck(deck, slides=2, image=image)

    result = run_inspector(deck, "--require-notes")

    assert result.returncode == 0
    payload = result_json(result)
    assert list(payload) == [
        "ok",
        "file",
        "slides",
        "media",
        "notes_slides",
        "out_of_bounds",
        "warnings",
        "errors",
    ]
    assert payload == {
        "ok": True,
        "file": str(deck.resolve()),
        "slides": 2,
        "media": 1,
        "notes_slides": 2,
        "out_of_bounds": [],
        "warnings": [],
        "errors": [],
    }


def test_missing_referenced_media_is_a_structural_error(tmp_path: Path) -> None:
    image = tmp_path / "figure.png"
    Image.new("RGB", (32, 24), "white").save(image)
    deck = tmp_path / "missing-media.pptx"
    make_deck(deck, image=image)
    media = package_part(deck, "ppt/media/")
    rewrite_package(deck, remove={media})

    result = run_inspector(deck)

    assert result.returncode == 1
    payload = result_json(result)
    assert any(
        "missing" in error.lower() and media in error
        for error in payload["errors"]
    )


def test_duplicate_package_names_are_rejected(tmp_path: Path) -> None:
    deck = tmp_path / "duplicate.pptx"
    make_deck(deck)
    with zipfile.ZipFile(deck) as package:
        duplicate = package.read("[Content_Types].xml")
    with pytest.warns(UserWarning, match="Duplicate name"):
        with zipfile.ZipFile(deck, "a") as package:
            package.writestr("[Content_Types].xml", duplicate)

    result = run_inspector(deck)

    assert result.returncode == 1
    assert any(
        "duplicate" in error.lower() and "[Content_Types].xml" in error
        for error in result_json(result)["errors"]
    )


def test_archive_metadata_limits_are_enforced_before_reads(monkeypatch) -> None:
    module = load_inspector_module()
    monkeypatch.setattr(module, "MAX_ZIP_ENTRIES", 2)
    monkeypatch.setattr(module, "MAX_ENTRY_UNCOMPRESSED_BYTES", 100)
    monkeypatch.setattr(module, "MAX_TOTAL_UNCOMPRESSED_BYTES", 150)
    monkeypatch.setattr(module, "MAX_COMPRESSION_RATIO", 10)
    monkeypatch.setattr(module, "MIN_RATIO_CHECK_BYTES", 20)

    entries = [
        SimpleNamespace(filename="one", file_size=80, compress_size=80),
        SimpleNamespace(filename="two", file_size=80, compress_size=80),
        SimpleNamespace(filename="three", file_size=30, compress_size=1),
    ]
    errors = module._archive_limit_errors(entries)

    assert any("entry count" in error.lower() for error in errors)
    assert any("total uncompressed" in error.lower() for error in errors)
    assert any("compression ratio" in error.lower() for error in errors)


def test_content_type_must_cover_slide_parts(tmp_path: Path) -> None:
    deck = tmp_path / "missing-content-type.pptx"
    make_deck(deck)

    def remove_slide_override(xml: bytes) -> bytes:
        text = xml.decode("utf-8")
        import re

        return re.sub(
            r'<Override PartName="/ppt/slides/slide1.xml"[^>]*/>',
            "",
            text,
        ).encode("utf-8")

    replace_package_text(deck, "[Content_Types].xml", remove_slide_override)

    result = run_inspector(deck)

    assert result.returncode == 1
    assert any(
        "content type" in error.lower() and "ppt/slides/slide1.xml" in error
        for error in result_json(result)["errors"]
    )


def test_orphaned_media_part_is_rejected(tmp_path: Path) -> None:
    deck = tmp_path / "orphan-media.pptx"
    make_deck(deck)
    image = tmp_path / "orphan.png"
    Image.new("RGB", (8, 8), "white").save(image)
    rewrite_package(
        deck,
        additions={"ppt/media/orphan.png": image.read_bytes()},
    )

    result = run_inspector(deck)

    assert result.returncode == 1
    assert any(
        "orphan" in error.lower() and "ppt/media/orphan.png" in error
        for error in result_json(result)["errors"]
    )


def test_orphaned_notes_part_is_rejected(tmp_path: Path) -> None:
    deck = tmp_path / "orphan-notes.pptx"
    make_deck(deck)
    rels = "ppt/slides/_rels/slide1.xml.rels"

    def remove_notes_relationship(xml: bytes) -> bytes:
        import re

        return re.sub(
            rb'<Relationship\b[^>]*Type="[^"]*/notesSlide"[^>]*/>',
            b"",
            xml,
            count=1,
        )

    replace_package_text(deck, rels, remove_notes_relationship)

    result = run_inspector(deck)

    assert result.returncode == 1
    assert any(
        "orphan" in error.lower() and "notesSlide" in error
        for error in result_json(result)["errors"]
    )


def test_orphaned_slide_part_is_rejected(tmp_path: Path) -> None:
    deck = tmp_path / "orphan-slide.pptx"
    make_deck(deck)
    rels = "ppt/_rels/presentation.xml.rels"

    def remove_slide_relationship(xml: bytes) -> bytes:
        import re

        return re.sub(
            rb'<Relationship\b[^>]*Type="[^"]*/slide"[^>]*/>',
            b"",
            xml,
            count=1,
        )

    replace_package_text(deck, rels, remove_slide_relationship)

    result = run_inspector(deck)

    assert result.returncode == 1
    assert any(
        "orphan" in error.lower() and "ppt/slides/slide1.xml" in error
        for error in result_json(result)["errors"]
    )


def test_missing_notes_relationship_target_is_rejected(tmp_path: Path) -> None:
    deck = tmp_path / "missing-notes-part.pptx"
    make_deck(deck)
    notes = package_part(deck, "ppt/notesSlides/notesSlide")
    rewrite_package(deck, remove={notes})

    result = run_inspector(deck)

    assert result.returncode == 1
    assert any(
        "missing" in error.lower() and notes in error
        for error in result_json(result)["errors"]
    )


def test_external_asset_relationship_is_rejected(tmp_path: Path) -> None:
    deck = tmp_path / "external-asset.pptx"
    make_deck(deck)
    rels = "ppt/slides/_rels/slide1.xml.rels"

    def add_external_image(xml: bytes) -> bytes:
        marker = b"</Relationships>"
        relationship = (
            b'<Relationship Id="rIdExternalImage" '
            b'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            b'relationships/image" Target="https://example.org/image.png" '
            b'TargetMode="External"/>'
        )
        assert marker in xml
        return xml.replace(marker, relationship + marker)

    replace_package_text(deck, rels, add_external_image)

    result = run_inspector(deck)

    assert result.returncode == 1
    assert any(
        "external asset" in error.lower() and "image.png" in error
        for error in result_json(result)["errors"]
    )


def test_ordinary_external_hyperlinks_are_allowed(tmp_path: Path) -> None:
    deck = tmp_path / "hyperlink.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(3), Inches(1)
    )
    run = text_box.text_frame.paragraphs[0].add_run()
    run.text = "Source"
    run.hyperlink.address = "https://example.org/paper"
    add_meaningful_notes(slide)
    presentation.save(deck)

    result = run_inspector(deck)

    assert result.returncode == 0
    assert result_json(result)["errors"] == []


def test_corrupt_referenced_raster_is_rejected(tmp_path: Path) -> None:
    image = tmp_path / "figure.png"
    Image.new("RGB", (32, 24), "white").save(image)
    deck = tmp_path / "corrupt-image.pptx"
    make_deck(deck, image=image)
    media = package_part(deck, "ppt/media/")
    rewrite_package(deck, replacements={media: b"not a valid raster image"})

    result = run_inspector(deck)

    assert result.returncode == 1
    assert any(
        "image" in error.lower() and media in error
        for error in result_json(result)["errors"]
    )


def test_corrupt_png_declared_octet_stream_is_still_decoded_and_rejected(
    tmp_path: Path,
) -> None:
    image = tmp_path / "figure.png"
    Image.new("RGB", (32, 24), "white").save(image)
    deck = tmp_path / "disguised-corrupt-image.pptx"
    make_deck(deck, image=image)
    media = package_part(deck, "ppt/media/")

    def disguise_png_content_type(xml: bytes) -> bytes:
        return xml.replace(
            b'Extension="png" ContentType="image/png"',
            b'Extension="png" ContentType="application/octet-stream"',
        )

    with zipfile.ZipFile(deck) as package:
        content_types = package.read("[Content_Types].xml")
    disguised_types = disguise_png_content_type(content_types)
    assert disguised_types != content_types
    rewrite_package(
        deck,
        replacements={
            "[Content_Types].xml": disguised_types,
            media: b"not a valid PNG",
        },
    )

    result = run_inspector(deck)

    assert result.returncode == 1
    errors = result_json(result)["errors"]
    assert any(
        "image relationship" in error.lower()
        and "application/octet-stream" in error
        for error in errors
    )
    assert any(
        media in error and "not a valid raster image" in error.lower()
        for error in errors
    )


def test_image_relationship_cannot_target_an_existing_notes_part(
    tmp_path: Path,
) -> None:
    image = tmp_path / "figure.png"
    Image.new("RGB", (32, 24), "white").save(image)
    deck = tmp_path / "image-to-notes.pptx"
    make_deck(deck, image=image)
    media = package_part(deck, "ppt/media/")
    notes = package_part(deck, "ppt/notesSlides/notesSlide")
    rels = "ppt/slides/_rels/slide1.xml.rels"
    with zipfile.ZipFile(deck) as package:
        relationships = package.read(rels)
    retargeted = relationships.replace(
        f"../media/{PurePosixPath(media).name}".encode("utf-8"),
        f"../notesSlides/{PurePosixPath(notes).name}".encode("utf-8"),
    )
    assert retargeted != relationships
    rewrite_package(
        deck,
        remove={media},
        replacements={rels: retargeted},
    )

    result = run_inspector(deck)

    assert result.returncode == 1
    assert any(
        "image relationship" in error.lower()
        and notes in error
        and "incompatible" in error.lower()
        for error in result_json(result)["errors"]
    )


def test_crc_corruption_is_reported_before_xml_reads(tmp_path: Path) -> None:
    deck = tmp_path / "bad-crc.pptx"
    make_deck(deck)
    corrupt_stored_member_crc(deck, "[Content_Types].xml")

    result = run_inspector(deck)

    assert result.returncode == 1
    assert any("CRC" in error for error in result_json(result)["errors"])


def test_image_format_not_decoded_by_pillow_has_controlled_warning(
    tmp_path: Path,
) -> None:
    module = load_inspector_module()
    package_path = tmp_path / "image-format.zip"
    part = "ppt/media/vector.svg"
    with zipfile.ZipFile(package_path, "w") as package:
        package.writestr(part, b"<svg xmlns='http://www.w3.org/2000/svg'/>")

    with zipfile.ZipFile(package_path) as package:
        errors, warnings = module._validate_images(
            package,
            [part],
            {part: "image/svg+xml"},
        )

    assert errors == []
    assert any(part in warning and "not decoded" in warning for warning in warnings)


def test_shape_outside_slide_is_reported_actionably(tmp_path: Path) -> None:
    deck = tmp_path / "outside.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(
        -int(presentation.slide_width * 0.02),
        Inches(1),
        Inches(1),
        Inches(0.5),
    )
    shape.name = "Outside evidence label"
    shape.text = "Source"
    add_meaningful_notes(slide)
    presentation.save(deck)

    result = run_inspector(deck)

    assert result.returncode == 1
    payload = result_json(result)
    assert payload["ok"] is False
    assert len(payload["out_of_bounds"]) == 1
    finding = payload["out_of_bounds"][0]
    assert finding["slide"] == 1
    assert finding["shape"] == "Outside evidence label"
    assert finding["severity"] == "error"
    assert finding["sides"] == ["left"]
    assert finding["overflow_percent"]["left"] > 1
    assert "move or resize" in finding["action"].lower()
    assert any(
        "Slide 1" in error and "Outside evidence label" in error
        for error in payload["errors"]
    )


def test_one_percent_or_smaller_bleed_is_not_a_hard_failure(tmp_path: Path) -> None:
    deck = tmp_path / "bleed.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(
        -int(presentation.slide_width * 0.01),
        Inches(1),
        Inches(1),
        Inches(0.5),
    )
    shape.name = "Intentional bleed"
    shape.text = "Edge"
    add_meaningful_notes(slide)
    presentation.save(deck)

    result = run_inspector(deck)

    assert result.returncode == 0
    payload = result_json(result)
    assert payload["ok"] is True
    assert payload["out_of_bounds"][0]["severity"] == "warning"
    assert payload["out_of_bounds"][0]["sides"] == ["left"]
    assert any("bleed" in warning.lower() for warning in payload["warnings"])
    assert payload["errors"] == []


def test_more_than_one_percent_bleed_is_a_hard_failure(tmp_path: Path) -> None:
    deck = tmp_path / "material-bleed.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(
        Inches(1),
        Inches(1),
        presentation.slide_width,
        Inches(0.5),
    )
    shape.name = "Wide background"
    shape.text = "Background"
    add_meaningful_notes(slide)
    presentation.save(deck)

    payload = result_json(run_inspector(deck))

    assert payload["ok"] is False
    finding = payload["out_of_bounds"][0]
    assert finding["severity"] == "error"
    assert finding["sides"] == ["right"]
    assert finding["overflow_percent"]["right"] > 1


def test_one_emu_above_the_bleed_limit_is_a_hard_failure(tmp_path: Path) -> None:
    deck = tmp_path / "just-over-limit.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(
        -int(presentation.slide_width * 0.01) - 1,
        Inches(1),
        Inches(1),
        Inches(0.5),
    )
    shape.name = "Barely outside"
    add_meaningful_notes(slide)
    presentation.save(deck)

    result = run_inspector(deck)

    assert result.returncode == 1
    finding = result_json(result)["out_of_bounds"][0]
    assert finding["severity"] == "error"
    assert finding["overflow_percent"]["left"] > 1


def test_zero_area_non_connector_is_a_structural_error(tmp_path: Path) -> None:
    deck = tmp_path / "zero-area.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(
        Inches(1), Inches(1), 0, Inches(1)
    )
    shape.name = "Collapsed evidence box"
    add_meaningful_notes(slide)
    presentation.save(deck)

    result = run_inspector(deck)

    assert result.returncode == 1
    assert any(
        "zero-area" in error.lower() and "Collapsed evidence box" in error
        for error in result_json(result)["errors"]
    )


def test_one_axis_zero_connector_is_allowed(tmp_path: Path) -> None:
    deck = tmp_path / "vertical-connector.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(1),
        Inches(1),
        Inches(1),
        Inches(2),
    )
    add_meaningful_notes(slide)
    presentation.save(deck)

    result = run_inspector(deck)

    assert result.returncode == 0
    assert result_json(result)["errors"] == []


def test_tiny_shape_wholly_outside_is_always_an_error(tmp_path: Path) -> None:
    deck = tmp_path / "tiny-outside.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    width = int(presentation.slide_width * 0.005)
    shape = slide.shapes.add_textbox(
        -width, Inches(1), width, Inches(0.25)
    )
    shape.name = "Tiny off-canvas label"
    add_meaningful_notes(slide)
    presentation.save(deck)

    result = run_inspector(deck)

    assert result.returncode == 1
    finding = result_json(result)["out_of_bounds"][0]
    assert finding["severity"] == "error"
    assert "wholly outside" in finding["action"].lower()


def test_rotated_visual_aabb_outside_canvas_is_reported(tmp_path: Path) -> None:
    deck = tmp_path / "rotated-outside.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(
        0, Inches(1), Inches(2), Inches(2)
    )
    shape.name = "Rotated method box"
    shape.rotation = 45
    add_meaningful_notes(slide)
    presentation.save(deck)

    result = run_inspector(deck)

    assert result.returncode == 1
    finding = result_json(result)["out_of_bounds"][0]
    assert finding["shape"] == "Rotated method box"
    assert finding["overflow_percent"]["left"] > 1


def test_group_children_are_inspected_recursively(tmp_path: Path) -> None:
    deck = tmp_path / "group-child.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    child = group.shapes.add_textbox(
        -int(presentation.slide_width * 0.02),
        Inches(1),
        Inches(1),
        Inches(0.5),
    )
    child.name = "Grouped evidence label"
    add_meaningful_notes(slide)
    presentation.save(deck)

    result = run_inspector(deck)

    assert result.returncode == 1
    findings = result_json(result)["out_of_bounds"]
    assert any(finding["shape"] == "Grouped evidence label" for finding in findings)


def test_group_coordinate_adapter_matches_python_pptx_1_0_2() -> None:
    module = load_inspector_module()
    assert importlib.metadata.version("python-pptx") == "1.0.2"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(
        Inches(2), Inches(3), Inches(1), Inches(0.5)
    )

    offset_x, offset_y, extent_width, extent_height = (
        module._group_child_coordinates(group)
    )

    assert (offset_x, offset_y) == (Inches(2), Inches(3))
    assert (extent_width, extent_height) == (Inches(1), Inches(0.5))


def test_missing_notes_warn_by_default_and_error_when_required(tmp_path: Path) -> None:
    deck = tmp_path / "missing-notes.pptx"
    make_deck(deck, slides=2, notes=False)

    optional = run_inspector(deck)
    required = run_inspector(deck, "--require-notes")

    assert optional.returncode == 0
    optional_payload = result_json(optional)
    assert optional_payload["notes_slides"] == 0
    assert any("slides 1, 2" in warning.lower() for warning in optional_payload["warnings"])
    assert optional_payload["errors"] == []

    assert required.returncode == 1
    required_payload = result_json(required)
    assert any("slides 1, 2" in error.lower() for error in required_payload["errors"])


def test_empty_or_whitespace_only_notes_are_not_meaningful(tmp_path: Path) -> None:
    deck = tmp_path / "empty-notes.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = "  \n\t"
    presentation.save(deck)

    result = run_inspector(deck, "--require-notes")

    assert result.returncode == 1
    payload = result_json(result)
    assert payload["notes_slides"] == 1
    assert any("slide 1" in error.lower() for error in payload["errors"])


def test_invalid_zip_returns_package_error(tmp_path: Path) -> None:
    deck = tmp_path / "not-a-package.pptx"
    deck.write_text("not a zip package", encoding="utf-8")

    result = run_inspector(deck)

    assert result.returncode == 1
    payload = result_json(result)
    assert payload["slides"] == 0
    assert any("ZIP" in error for error in payload["errors"])


def test_unreadable_presentation_inside_zip_returns_reopen_error(tmp_path: Path) -> None:
    deck = tmp_path / "corrupt.pptx"
    with zipfile.ZipFile(deck, "w") as package:
        package.writestr("[Content_Types].xml", "<Types/>")

    result = run_inspector(deck)

    assert result.returncode == 1
    payload = result_json(result)
    assert any("reopen" in error.lower() for error in payload["errors"])


@pytest.mark.parametrize("mode", ["missing", "zero"])
def test_invalid_slide_dimensions_return_stable_json(
    tmp_path: Path, mode: str
) -> None:
    deck = tmp_path / f"{mode}-dimensions.pptx"
    make_deck(deck)

    def invalidate_dimensions(xml: bytes) -> bytes:
        import re

        if mode == "missing":
            return re.sub(rb"<p:sldSz\b[^>]*/>", b"", xml, count=1)

        def zero_width(match) -> bytes:
            return re.sub(rb'cx="\d+"', b'cx="0"', match.group(0), count=1)

        return re.sub(rb"<p:sldSz\b[^>]*/>", zero_width, xml, count=1)

    replace_package_text(deck, "ppt/presentation.xml", invalidate_dimensions)

    result = run_inspector(deck)

    assert result.returncode == 1
    payload = result_json(result)
    assert list(payload) == [
        "ok",
        "file",
        "slides",
        "media",
        "notes_slides",
        "out_of_bounds",
        "warnings",
        "errors",
    ]
    assert any("slide dimensions" in error.lower() for error in payload["errors"])


def test_presentation_traversal_exception_returns_stable_result(
    tmp_path: Path, monkeypatch
) -> None:
    module = load_inspector_module()
    deck = tmp_path / "traversal.pptx"
    make_deck(deck)
    from pptx.slide import Slides

    def fail_iteration(_self):
        raise RuntimeError("injected traversal failure")

    monkeypatch.setattr(Slides, "__iter__", fail_iteration)

    result, dimensions = module.inspect_pptx(deck)

    assert result["ok"] is False
    assert dimensions is not None
    assert any("injected traversal failure" in error for error in result["errors"])


def test_missing_input_is_a_clean_qa_error(tmp_path: Path) -> None:
    deck = tmp_path / "missing.pptx"

    result = run_inspector(deck)

    assert result.returncode == 1
    payload = result_json(result)
    assert payload["file"] == str(deck.resolve())
    assert any("does not exist" in error.lower() for error in payload["errors"])


def test_markdown_report_contains_actionable_findings(tmp_path: Path) -> None:
    deck = tmp_path / "report-deck.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(
        -int(presentation.slide_width * 0.02),
        Inches(1),
        Inches(1),
        Inches(0.5),
    )
    shape.name = "Escaping source"
    shape.text = "Source"
    presentation.save(deck)
    report = tmp_path / "qa_report.md"

    first = run_inspector(deck, "--report", str(report))
    first_text = report.read_text(encoding="utf-8")
    second = run_inspector(deck, "--report", str(report))

    assert first.returncode == second.returncode == 1
    assert first_text == report.read_text(encoding="utf-8")
    assert "# Paper2PPT structural QA report" in first_text
    assert "Status: fail" in first_text
    assert "Slides: 1" in first_text
    assert "Slide dimensions:" in first_text
    assert "Slide 1" in first_text
    assert "Escaping source" in first_text
    assert "Move or resize" in first_text
    assert "structural checks only" in first_text


def test_report_parent_is_not_created_and_write_failure_is_reported(
    tmp_path: Path,
) -> None:
    deck = tmp_path / "deck.pptx"
    make_deck(deck)
    report = tmp_path / "missing-parent" / "qa_report.md"

    result = run_inspector(deck, "--report", str(report))

    assert result.returncode == 1
    assert not report.parent.exists()
    payload = result_json(result)
    assert payload["ok"] is False
    assert any(
        "report" in error.lower() and "parent" in error.lower()
        for error in payload["errors"]
    )


def test_report_cannot_replace_the_input_deck(tmp_path: Path) -> None:
    deck = tmp_path / "same-path.pptx"
    make_deck(deck)
    original = deck.read_bytes()

    result = run_inspector(deck, "--report", str(deck))

    assert result.returncode == 1
    assert deck.read_bytes() == original
    assert any(
        "same file" in error.lower() for error in result_json(result)["errors"]
    )


def test_report_symlink_to_input_is_rejected(tmp_path: Path) -> None:
    deck = tmp_path / "symlink-source.pptx"
    make_deck(deck)
    original = deck.read_bytes()
    report = tmp_path / "report-link.md"
    try:
        report.symlink_to(deck)
    except (NotImplementedError, OSError) as error:
        pytest.skip(f"symlinks unavailable: {error}")

    result = run_inspector(deck, "--report", str(report))

    assert result.returncode == 1
    assert deck.read_bytes() == original
    assert report.is_symlink()
    assert any(
        "same file" in error.lower() for error in result_json(result)["errors"]
    )


def test_report_hardlink_to_input_is_rejected(tmp_path: Path) -> None:
    deck = tmp_path / "hardlink-source.pptx"
    make_deck(deck)
    original = deck.read_bytes()
    report = tmp_path / "report-hardlink.md"
    try:
        os.link(deck, report)
    except (NotImplementedError, OSError) as error:
        pytest.skip(f"hardlinks unavailable: {error}")

    result = run_inspector(deck, "--report", str(report))

    assert result.returncode == 1
    assert deck.read_bytes() == original
    assert any(
        "same file" in error.lower() for error in result_json(result)["errors"]
    )


def test_existing_report_is_atomically_replaced(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    make_deck(deck)
    report = tmp_path / "qa_report.md"
    report.write_text("old report\n", encoding="utf-8")

    result = run_inspector(deck, "--report", str(report))

    assert result.returncode == 0
    assert report.read_text(encoding="utf-8").startswith(
        "# Paper2PPT structural QA report\n"
    )
    assert list(tmp_path.glob(f".{report.name}.*.tmp")) == []


def test_injected_temp_write_failure_preserves_existing_report(
    tmp_path: Path, monkeypatch, capsys
) -> None:
    module = load_inspector_module()
    deck = tmp_path / "deck.pptx"
    make_deck(deck)
    report = tmp_path / "qa_report.md"
    report.write_text("old report\n", encoding="utf-8")

    def fail_write(*_args, **_kwargs):
        raise OSError("injected write failure")

    monkeypatch.setattr(module, "_write_temp_file", fail_write)

    returncode = module.main(
        [str(deck), "--json", "--report", str(report)]
    )
    captured = capsys.readouterr()
    payload = json.loads(captured.out)

    assert returncode == 1
    assert captured.err == ""
    assert list(payload) == [
        "ok",
        "file",
        "slides",
        "media",
        "notes_slides",
        "out_of_bounds",
        "warnings",
        "errors",
    ]
    assert any("injected write failure" in error for error in payload["errors"])
    assert report.read_text(encoding="utf-8") == "old report\n"
    assert list(tmp_path.glob(f".{report.name}.*.tmp")) == []


def test_injected_replace_failure_preserves_existing_report(
    tmp_path: Path, monkeypatch, capsys
) -> None:
    module = load_inspector_module()
    deck = tmp_path / "deck.pptx"
    make_deck(deck)
    report = tmp_path / "qa_report.md"
    report.write_text("old report\n", encoding="utf-8")

    def fail_replace(*_args, **_kwargs):
        raise OSError("injected replace failure")

    monkeypatch.setattr(module.os, "replace", fail_replace)

    returncode = module.main(
        [str(deck), "--json", "--report", str(report)]
    )
    captured = capsys.readouterr()
    payload = json.loads(captured.out)

    assert returncode == 1
    assert captured.err == ""
    assert any("injected replace failure" in error for error in payload["errors"])
    assert report.read_text(encoding="utf-8") == "old report\n"
    assert list(tmp_path.glob(f".{report.name}.*.tmp")) == []


def test_markdown_escapes_document_derived_shape_names_and_paths(
    tmp_path: Path,
) -> None:
    deck = tmp_path / "deck\n## Injected file heading.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(
        -int(presentation.slide_width * 0.02),
        Inches(1),
        Inches(1),
        Inches(0.5),
    )
    shape.name = "bad`\n## Injected shape heading\n![remote](https://evil.invalid/x.png)"
    add_meaningful_notes(slide)
    presentation.save(deck)
    report = tmp_path / "safe-report.md"

    result = run_inspector(deck, "--report", str(report))

    assert result.returncode == 1
    payload = result_json(result)
    assert "\n" not in payload["out_of_bounds"][0]["shape"]
    markdown = report.read_text(encoding="utf-8")
    assert markdown.count("# Paper2PPT structural QA report") == 1
    assert "\n## Injected" not in markdown
    assert "![remote](" not in markdown
    assert "\\#\\# Injected" in markdown


def test_cli_usage_errors_exit_two() -> None:
    result = subprocess.run(
        [sys.executable, str(INSPECTOR)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 2
    assert "usage:" in result.stderr.lower()


def test_quality_gate_documents_the_structural_inspector_contract() -> None:
    guidance = QUALITY_GATES.read_text(encoding="utf-8")

    assert (
        'python "<SKILL_DIR>/scripts/inspect_pptx.py" deck.pptx '
        "[--require-notes] [--json] [--report qa_report.md]"
    ) in guidance
    assert "package-level count" in guidance
    assert "meaningful speaker-note text" in guidance
    assert "does not replace content QA or rendered visual QA" in guidance
    assert "CRC integrity" in guidance
    assert "internal relationships" in guidance
    assert "external asset" in guidance
    assert "ordinary hyperlinks" in guidance
    assert "atomic replacement" in guidance
