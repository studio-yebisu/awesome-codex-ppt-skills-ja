import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys

from PIL import Image, ImageStat
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "examples" / "source"
OUTPUT = ROOT / "examples" / "output"
SHOWCASE = ROOT / "examples" / "showcase"
DECK = OUTPUT / "paper2ppt-example-zh.pptx"
INSPECTOR = ROOT / "scripts" / "inspect_pptx.py"


def _slide_text(slide) -> str:
    return "\n".join(
        shape.text for shape in slide.shapes if hasattr(shape, "text_frame")
    )


def _meaningful_notes(slide) -> str:
    return slide.notes_slide.notes_text_frame.text.strip()


def _manifest_assets(manifest: str) -> list[dict[str, str]]:
    records = []
    for section in re.split(r"(?m)^## Asset \d+\s*$", manifest)[1:]:
        fields = dict(re.findall(r"(?m)^- ([^:]+): (.+)$", section))
        records.append(fields)
    return records


def _embedded_media() -> dict[str, set[int]]:
    media: dict[str, set[int]] = {}
    presentation = Presentation(DECK)
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for relationship in slide.part.rels.values():
            if relationship.is_external or relationship.reltype != RT.IMAGE:
                continue
            digest = hashlib.sha256(relationship.target_part.blob).hexdigest()
            media.setdefault(digest, set()).add(slide_number)
    return media


def test_example_source_has_explicit_open_license_metadata() -> None:
    metadata = json.loads((SOURCE / "paper-metadata.json").read_text(encoding="utf-8"))

    assert metadata["doi"] == "10.1371/journal.pone.0276250"
    assert metadata["publisher"] == "Public Library of Science"
    assert metadata["license"] == "CC BY 4.0"
    assert metadata["license_url"] == "https://creativecommons.org/licenses/by/4.0/"
    assert metadata["retrieved"] == "2026-07-15"
    assert metadata["article_url"].startswith("https://")
    assert metadata["full_text_url"].startswith("https://")
    assert metadata["figures"]
    for figure in metadata["figures"]:
        assert figure["figure"]
        assert figure["url"].startswith("https://")
        assert figure["attribution"]

    source_readme = (SOURCE / "README.md").read_text(encoding="utf-8")
    assert metadata["doi"] in source_readme
    assert metadata["license_url"] in source_readme


def test_example_deck_is_editable_chinese_and_fully_noted() -> None:
    presentation = Presentation(DECK)

    assert 8 <= len(presentation.slides) <= 10
    assert presentation.slide_width / presentation.slide_height == 16 / 9
    all_text = "\n".join(_slide_text(slide) for slide in presentation.slides)
    assert re.search(r"[\u4e00-\u9fff]", all_text)
    assert "nature-paper2ppt" not in all_text.lower()
    assert "Nature 风格" not in all_text
    for slide in presentation.slides:
        assert len(_meaningful_notes(slide)) >= 30

    result_slides = [
        slide for slide in presentation.slides if "结果页" in _meaningful_notes(slide)
    ]
    assert len(result_slides) >= 3
    for slide in result_slides:
        assert "来源：" in _slide_text(slide)


def test_example_deck_passes_packaged_inspector_with_required_notes() -> None:
    result = subprocess.run(
        [
            sys.executable,
            str(INSPECTOR),
            str(DECK),
            "--require-notes",
            "--json",
        ],
        cwd=ROOT,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stdout + result.stderr
    payload = json.loads(result.stdout)
    assert payload["ok"] is True
    assert payload["notes_slides"] == payload["slides"]
    assert payload["errors"] == []


def test_asset_manifest_traces_every_used_asset() -> None:
    manifest = (OUTPUT / "asset_manifest.md").read_text(encoding="utf-8")
    assets = _manifest_assets(manifest)
    embedded_media = _embedded_media()

    assert assets
    assert len(assets) == len(embedded_media)
    assert {asset["SHA-256"] for asset in assets} == set(embedded_media)
    for asset in assets:
        for field in (
            "Filename",
            "Source URL",
            "Source anchor",
            "Figure / panel",
            "Crop",
            "Destination slide",
            "License",
            "SHA-256",
        ):
            assert asset.get(field, "").strip(), field
        assert asset["Source URL"].startswith("https://")
        assert "CC BY 4.0" in asset["License"]
        destination_slides = {
            int(number) for number in re.findall(r"\d+", asset["Destination slide"])
        }
        assert destination_slides == embedded_media[asset["SHA-256"]]


def test_qa_report_records_all_three_gates_and_every_slide_reviewed() -> None:
    report = (OUTPUT / "qa_report.md").read_text(encoding="utf-8")
    presentation = Presentation(DECK)

    assert "- Source sufficiency: full-paper" in report
    assert "## content QA" in report
    assert "## structural QA" in report
    assert "## visual QA" in report
    assert report.count("- Status: pass") >= 3
    assert "LibreOffice" in report
    assert "Poppler" in report
    assert f"- Slides inspected individually: 1-{len(presentation.slides)}" in report
    assert "- Contact sheet inspected: pass" in report
    assert "## Completion state" in report
    assert re.search(r"(?ms)^## Completion state\s+.*^- Status: pass$", report)


def test_showcase_contains_nonblank_overview_and_three_slide_previews() -> None:
    images = [SHOWCASE / "overview.png", *sorted(SHOWCASE.glob("slide-*.png"))]

    assert len(images) == 4
    for path in images:
        with Image.open(path) as image:
            image.load()
            assert image.width >= 1000
            assert image.height >= 560
            assert ImageStat.Stat(image.convert("L")).var[0] > 50


def test_example_documents_do_not_reintroduce_legacy_product_branding() -> None:
    text_files = [
        SOURCE / "README.md",
        SOURCE / "paper-metadata.json",
        OUTPUT / "qa_report.md",
        OUTPUT / "asset_manifest.md",
    ]
    combined = "\n".join(path.read_text(encoding="utf-8") for path in text_files)

    assert "nature-paper2ppt" not in combined.lower()
    assert "Nature 风格" not in combined
