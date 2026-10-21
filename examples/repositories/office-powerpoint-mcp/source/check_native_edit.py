"""Bounded real MCP edit/save proof on a separate output; canonical deck is preserved."""
from pathlib import Path
import asyncio,hashlib,json,os,sys
from mcp import ClientSession,StdioServerParameters
from mcp.client.stdio import stdio_client
from pptx import Presentation
CASE=Path(__file__).resolve().parents[1];UP=CASE/'source/upstream'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
async def main():
 original=CASE/'deck.pptx';before=sha(original);out=CASE/'source/native-edit-check.pptx';calls=[]
 params=StdioServerParameters(command=sys.executable,args=[str(UP/'ppt_mcp_server.py'),'--transport','stdio'],cwd=str(UP),env={**os.environ,'PYTHONIOENCODING':'utf-8'})
 with (CASE/'source/native-edit-server.log').open('w',encoding='utf-8') as err:
  async with stdio_client(params,errlog=err) as (r,w):
   async with ClientSession(r,w) as session:
    await session.initialize()
    async def call(name,**args):
     result=await session.call_tool(name,arguments=args);obj=result.model_dump(mode='json');calls.append({'tool':name,'arguments':args,'result':obj})
     for block in result.content:
      if hasattr(block,'text'):
       try:value=json.loads(block.text)
       except json.JSONDecodeError:continue
       assert not (isinstance(value,dict) and value.get('error')),value
     assert not result.isError
    await call('open_presentation',file_path=str(original),id='native-edit-proof')
    await call('manage_text',slide_index=0,operation='format_runs',shape_index=0,text_runs=[{'text':'編集検証用：本文を書き換え','font_name':'BIZ UDPGothic','font_size':28,'bold':True}],presentation_id='native-edit-proof')
    source=Presentation(original);shape_index=next(i for i,s in enumerate(source.slides[2].shapes) if s.has_chart);chart=source.slides[2].shapes[shape_index].chart
    values=[list(s.values) for s in chart.series];values[0][0]=221
    await call('update_chart_data',slide_index=2,shape_index=shape_index,categories=[v.label for v in chart.plots[0].categories],series_data=[{'name':s.name,'values':v} for s,v in zip(chart.series,values)],presentation_id='native-edit-proof')
    await call('save_presentation',file_path=str(out),presentation_id='native-edit-proof')
    await call('extract_slide_text',slide_index=0,presentation_id='native-edit-proof')
 proof=Presentation(out);assert any('編集検証用' in s.text for s in proof.slides[0].shapes if s.has_text_frame)
 assert list(next(s.chart for s in proof.slides[2].shapes if s.has_chart).series[0].values)[0]==221
 assert sha(original)==before
 (CASE/'source/native-edit-calls.json').write_text(json.dumps(calls,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
 result={'scope':'Actual original MCP tools on a separate native PPTX copy; no GUI test or canonical mutation','canonical_sha256':before,'canonical_unchanged':True,'output':'source/native-edit-check.pptx','output_sha256':sha(out),'text_edit_saved':True,'native_chart_point_edit_saved':{'slide_index':2,'series':0,'category':0,'before':220,'after':221},'slide_count':len(proof.slides),'tool_calls':len(calls)}
 (CASE/'source/native-edit-check.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(json.dumps(result,ensure_ascii=False))
if __name__=='__main__':asyncio.run(main())
