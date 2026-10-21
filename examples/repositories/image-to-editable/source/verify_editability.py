"""Change one native label in a separate copy and optionally render it in PowerPoint."""
from pathlib import Path
import argparse, hashlib, json, subprocess
from pptx import Presentation
CASE=Path(__file__).resolve().parents[1]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    ap=argparse.ArgumentParser();ap.add_argument('--render',action='store_true');args=ap.parse_args()
    src=CASE/'deck.pptx';original_hash=sha(src);out=CASE/'qa/editability/native-label-edited.pptx';out.parent.mkdir(parents=True,exist_ok=True)
    prs=Presentation(src);changed=[]
    for s in prs.slides[1].shapes:
        if not s.has_text_frame:continue
        for p in s.text_frame.paragraphs:
            for r in p.runs:
                if '次に記録する内容' in r.text:
                    r.text=r.text.replace('次に記録する内容','次回の記録内容');changed.append(s.name)
    assert len(changed)==1,changed
    prs.save(out)
    assert sha(src)==original_hash
    check=Presentation(out);assert any('次回の記録内容' in s.text for s in check.slides[1].shapes if s.has_text_frame)
    receipt={'canonical_sha256':original_hash,'canonical_unchanged':True,'output':'qa/editability/native-label-edited.pptx','output_sha256':sha(out),'operation':'Replace text in one native run, retaining its formatting and geometry','slide_index':1,'shapes':changed,'before':'次に記録する内容','after':'次回の記録内容','gui_edit_test':False,'render_requested':args.render}
    if args.render:
        argv=['pwsh','-NoProfile','-File',str(CASE/'source/render_powerpoint.ps1'),'-PresentationPath',str(out),'-OutputDirectory',str(CASE/'qa/editability/render')]
        run=subprocess.run(argv,capture_output=True,text=True,encoding='utf-8',errors='replace');(CASE/'qa/editability/powerpoint.log').write_text(run.stdout+run.stderr,encoding='utf-8');receipt['powerpoint_exit_code']=run.returncode
        if run.returncode:raise SystemExit(run.stdout+run.stderr)
    (CASE/'qa/editability/native-edit.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(receipt,ensure_ascii=False))
if __name__=='__main__':main()
