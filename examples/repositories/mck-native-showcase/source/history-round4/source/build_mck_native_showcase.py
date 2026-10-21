"""『フォント構成と制作手順』— likaku/Mck-ppt-design-skill の MckEngine 標準版式で作る4ページ。

原典SKILLの五段階に沿って S1 brief → S2 outline → S3 content を先に書き、
S3 は原典 gate_check_s3.py、S4 は原典 gate_check.py を改変せずに subprocess で実行し、
書き出された JSON の passed だけで次へ進む。描画は MckEngine の cover / donut /
data_table / timeline をそのまま呼ぶ。上書きや自作 renderer は使わない。

原典constantsの書体名をプロセス内で日本語書体へ指定する。donutの凡例は、原典が
整数切捨てする百分率の文字だけを小数1桁へ直す。形状・座標・サイズと原典ファイルは不変。
"""
from collections import Counter
from pathlib import Path
import argparse, hashlib, importlib.util, inspect, json, os, re, shutil, subprocess, sys

ROOT = Path(__file__).resolve().parents[1]
ZIP_MODE = (ROOT / 'upstream/mck_ppt').exists()
UPSTREAM_CANDIDATES = [
    ROOT / 'upstream',                                            # 配布ZIPの同梱原典
    ROOT / '.build/actual-upstreams/likaku__Mck-ppt-design-skill',  # repoの固定取得物
    ROOT / 'skill-forks/sol-editorial-ja/upstream',               # repoに保存済みの同一commit
]
UPSTREAM_REQUIRED = ['mck_ppt/__init__.py', 'mck_ppt/constants.py', 'mck_ppt/core.py', 'mck_ppt/engine.py',
                     'mck_ppt/qa.py', 'references/scripts/gate_check_s3.py', 'references/scripts/gate_check.py',
                     'references/layout-matrix.yaml']
LAYOUTS = ('cover', 'donut', 'data_table', 'timeline')
PROJECT_DIR = 'ppt-project-font-kit'  # 原典SKILLの ppt-project-{slug}/ 規約


def fail(stage, message):
    print(f'[{stage}] 失敗: {message}', file=sys.stderr)
    sys.exit(1)


def load(p): return json.loads(Path(p).read_text(encoding='utf-8-sig'))
def write(p, s): p.parent.mkdir(parents=True, exist_ok=True); p.write_text(s, encoding='utf-8')
def dump(p, v): write(p, json.dumps(v, ensure_ascii=False, indent=2) + '\n')
def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def rel(p): return Path(p).resolve().relative_to(ROOT).as_posix()
def first_existing(paths): return next((p for p in paths if p.exists()), None)


def find_upstream():
    for up in UPSTREAM_CANDIDATES:
        if all((up / f).exists() for f in UPSTREAM_REQUIRED):
            return up
    fail('setup', '原典 likaku/Mck-ppt-design-skill が見つかりません: ' + ', '.join(rel(p) for p in UPSTREAM_CANDIDATES))


def upstream_identity(up):
    for meta in (up / 'ACQUISITION.json', up.parent / 'SOURCE.json', up / 'SOURCE.json'):
        if meta.exists():
            d = load(meta)
            if d.get('repository') and d.get('commit'):
                return {'repository': d['repository'], 'commit': d['commit']}  # root等の絶対パスは持ち出さない
    return {'repository': 'likaku/Mck-ppt-design-skill', 'commit': None}


def load_engine(up, fonts):
    """原典constantsを読み、実在する名前だけ差し替えてから原典パッケージを読み込む。"""
    if 'mck_ppt' in sys.modules:
        fail('setup', 'mck_ppt が既に読み込まれています。新しいプロセスで実行してください')
    spec = importlib.util.spec_from_file_location('mck_ppt.constants', up / 'mck_ppt/constants.py')
    const = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(const)
    for name, value in fonts.items():
        if not isinstance(getattr(const, name, None), str):
            fail('setup', f'原典constantsに書体名 {name} がありません')
        setattr(const, name, value)
    sys.modules['mck_ppt.constants'] = const
    sys.path.insert(0, str(up))
    import mck_ppt, mck_ppt.core as core, mck_ppt.engine as engine
    if Path(mck_ppt.__file__).resolve().parent != (up / 'mck_ppt').resolve():
        fail('setup', '固定原典以外の mck_ppt が読み込まれました')
    defaults = inspect.signature(core.add_text).parameters
    if (engine.FONT_HEADER, defaults['font_name'].default, defaults['ea_font'].default) != \
            (fonts['FONT_HEADER'], fonts['FONT_BODY'], fonts['FONT_EA']):
        fail('setup', '書体名の局所指定が原典エンジンに反映されていません')
    return engine.MckEngine, const


def read_layout_matrix(path):
    """layout-matrix.yaml の Layout Capacity Matrix 表を読む。"""
    out = {}
    row = re.compile(r'^\|\s*`(\w+)`\s*\|\s*([^|]*?)\s*\|\s*([^|]*?)\s*\|\s*([^|]*?)\s*\|\s*([^|]*?)\s*\|')
    for line in path.read_text(encoding='utf-8').splitlines():
        m = row.match(line)
        if not m:
            continue
        name, _, items, title, body = m.groups()
        n = re.search(r'(\d+)', items)
        out[name] = {'max_items': int(n.group(1)) if n else None,
                     'title': int(title) if title.isdigit() else None,
                     'body': {k.strip(): int(v) for k, v in re.findall(r'([^,:]+?):\s*(\d+)', body)}}
    return out


# ─── S1 ──────────────────────────────────────────────────────────────

def font_stats(data, plan):
    langs, cats = list(plan['languages']), list(plan['categories'])
    fams = data['families']
    ids = [f['id'] for f in fams]
    if len(ids) != len(set(ids)):
        fail('S1', 'フォントデータの id が重複しています')
    unknown = sorted({f['language'] for f in fams} - set(langs)) + sorted({f['category'] for f in fams} - set(cats))
    if unknown:
        fail('S1', f'plan.json に表示名がない language/category: {unknown}')
    by_lang = Counter(f['language'] for f in fams)
    by_cat = Counter(f['category'] for f in fams)
    cross = Counter((f['category'], f['language']) for f in fams)
    order = sorted((c for c in cats if by_cat[c]), key=lambda c: (-by_cat[c], cats.index(c)))
    first = {}
    for f in fams:
        first.setdefault((f['category'], f['language']), f['family'])
    exclusive = [(c, l) for c in order for l in langs if cross[(c, l)] == by_cat[c]]
    family_lang = {f['family']: f['language'] for f in fams}
    presets = data.get('presets', [])
    split = bool(presets) and all(family_lang.get(p.get('heading')) == 'ja' and family_lang.get(p.get('body')) == 'ja'
                                  and family_lang.get(p.get('latin')) == 'latin' for p in presets)
    return {'total': len(fams), 'updated_on': data.get('updated_on', ''), 'langs': langs, 'order': order,
            'by_lang': by_lang, 'by_cat': by_cat, 'cross': cross, 'first': first, 'exclusive': exclusive,
            'presets': len(presets), 'presets_split': split}


def summary_json(st):
    return {'family_count': st['total'], 'updated_on': st['updated_on'],
            'by_language': {l: st['by_lang'][l] for l in st['langs']},
            'by_category': {c: st['by_cat'][c] for c in st['order']},
            'by_category_language': {c: {l: st['cross'][(c, l)] for l in st['langs']} for c in st['order']},
            'preset_count': st['presets'], 'presets_all_ja_with_latin': st['presets_split']}


def write_brief(path, plan, st):
    L, K = plan['languages'], plan['categories']
    top = st['order'][0]
    messages = [
        f"収録は{st['total']}字体（" + '・'.join(f"{L[l]}{st['by_lang'][l]}" for l in st['langs']) + '）',
        f"分類別の最多は{K[top]}の{st['by_cat'][top]}字体",
        '、'.join(f'{K[c]}は{L[l]}のみ' for c, l in st['exclusive']) or '各分類に複数の言語がある',
        '制作は要件→構成→原稿→描画・検査の順。所要時間は計測していない',
    ]
    b = plan['brief']
    if not (b.get('audience') and b.get('goal') and messages):
        fail('S1', 'audience / goal / key_messages のいずれかが空です')
    fonts = ' / '.join(f'{k}={v}' for k, v in plan['fonts'].items())
    write(path, f"# {plan['title']} — brief\n\n"
                f"- 受け手（audience）: {b['audience']}\n"
                f"- 目的（goal）: {b['goal']}\n"
                f"- 時間: {b['duration_minutes']}分 / 4ページ\n"
                f"- 主要メッセージ（key_messages）:\n" + ''.join(f'  {i}. {m}\n' for i, m in enumerate(messages, 1)) +
                f"- データ: {b['data_source']}（{st['updated_on']}更新）\n"
                f"- 扱わないもの: {b['exclusions']}\n"
                f"- 書体: 原典constantsの {fonts} をこの生成プロセス内でだけ指定\n")
    return messages


# ─── S2 / S3 ─────────────────────────────────────────────────────────

def build_content(plan, st):
    L, K = plan['languages'], plan['categories']
    fmt = {'total': st['total'], 'updated_on': st['updated_on']}
    fmt.update({f'{l}_count': st['by_lang'][l] for l in st['langs']})
    top = st['order'][0]
    tie = len(st['order']) > 1 and st['by_cat'][st['order'][1]] == st['by_cat'][top]
    d, t, p = plan['donut'], plan['table'], plan['process']
    if len(st['order']) > len(d['colors']):
        fail('S3', f"分類が{len(st['order'])}種あり、plan.json の donut.colors が足りません")

    segments = [[st['by_cat'][c] / st['total'], d['colors'][i], f"{K[c]} {st['by_cat'][c]}字体"]
                for i, c in enumerate(st['order'])]
    summary = '・'.join(f"{L[l]}{st['by_lang'][l]}字体" for l in st['langs']) + '。' + \
        f'{K[top]}は' + '・'.join(f"{L[l]}{st['cross'][(top, l)]}" for l in st['langs'])
    if d.get('group_by') == 'language':
        segments = [[st['by_lang'][l] / st['total'], d['colors'][i], f"{L[l]} {st['by_lang'][l]}字体"] for i, l in enumerate(st['langs'])]
        summary = '／'.join(f"{L[l]}{st['by_lang'][l]}字体" for l in st['langs']) + '。\n' + f"プリセット{st['presets']}件は、和文の見出し・本文と欧文の英数字を組み合わせる。"
    summary = d.get('summary', summary)

    widths = t['col_widths_in']
    if len(widths) != len(st['langs']) + 3:
        fail('S3', 'table.col_widths_in の列数が表と一致しません')
    headers = [t['headers']['category'], *(L[l] for l in st['langs']), t['headers']['total'], t['headers']['sample']]
    rows = [[K[c], *(str(st['cross'][(c, l)]) for l in st['langs']), str(st['by_cat'][c]),
             ' / '.join(st['first'][(c, l)] for l in st['langs'] if (c, l) in st['first'])] for c in st['order']]
    rows.append([t['total_label'], *(str(st['by_lang'][l]) for l in st['langs']), str(st['total']), '—'])
    table_title = (t['title_prefix'] + '、'.join(f'{K[c]}は{L[l]}のみ' for c, l in st['exclusive'])
                   if st['exclusive'] else t['title_fallback'])
    table_bar = [t['bottom_bar_label'], t['presets_split'].format(n=st['presets'])] if st['presets_split'] else None

    slides = [
        {'idx': 1, 'layout': 'cover', 'title': plan['title'],
         'key_point': '',
         'subtitle': plan['cover']['subtitle'].format(**fmt), 'author': plan['cover']['author'],
         'date': plan['cover']['date'].format(**fmt)},
        {'idx': 2, 'layout': 'donut',
         'title': (d['title_tie'] if tie else d['title']).format(**fmt, top_label=K[top], top_count=st['by_cat'][top]),
         'key_point': f"収録{st['total']}字体を{'言語' if d.get('group_by') == 'language' else '分類'}で数え、構成比を一つの円で示す",
         'segments': segments, 'center_label': str(st['total']), 'center_sub': d['center_sub'],
         'summary': summary if d.get('show_summary', True) else None,
         'source': d['source'].format(**fmt)},
        {'idx': 3, 'layout': 'data_table', 'title': table_title,
         'key_point': '分類ごとに日本語と欧文の字体数を分け、代表書体を添える',
         'headers': headers, 'rows': rows, 'col_widths': widths, 'bottom_bar': table_bar,
         'source': t['source'].format(**fmt)},
        {'idx': 4, 'layout': 'timeline', 'title': p['title'],
         'key_point': '原典SKILLのS1〜S4を順に通し、S3とS4は原典の検査スクリプトで判定する',
         'milestones': p['milestones'], 'bottom_bar': p['bottom_bar'], 'source': p['source']},
    ]
    return slides


def check_outline(outline, matrix):
    """原典SKILLのS2門禁（AI自評項目）をコードで確認する。動詞の有無は判定しない。"""
    slides, minutes = outline['slides'], outline['brief']['duration_minutes']
    issues = []
    if not any(s['layout'] == 'cover' for s in slides):
        issues.append('cover がありません')
    if len(slides) > minutes * 1.2:
        issues.append(f'{len(slides)}ページは {minutes}分×1.2 を超えます')
    issues += [f"{s['idx']}: {s['layout']} は layout-matrix.yaml にありません" for s in slides if s['layout'] not in matrix]
    issues += [f"{s['idx']}: Action Title が10字以下" for s in slides if s['layout'] != 'cover' and len(s['title']) <= 10]
    if sum(s['layout'] == 'two_column_text' for s in slides) > 1:
        issues.append('two_column_text が2枚以上あります')
    return issues


def budget_issues(slide, matrix):
    """layout-matrix.yaml の char_budget / Max Items による事前確認（原典S3 gateの前段）。"""
    m, lay, out = matrix[slide['layout']], slide['layout'], []

    def over(field, text, limit):
        n = len(str(text).replace('\n', ''))
        if limit is not None and n > limit:
            out.append(f"slide {slide['idx']} {lay} {field}: {n} > {limit}字")
    over('title', slide['title'], m['title'])
    b = m['body']
    if lay == 'cover':
        over('subtitle', slide['subtitle'], b.get('subtitle'))
        over('author', slide['author'], b.get('author'))
    elif lay == 'donut':
        for seg in slide['segments']:
            over('segment label', seg[2], b.get('segment label'))
    elif lay == 'data_table':
        for h in slide['headers']:
            over('header', h, b.get('header'))
        for row in slide['rows']:
            for cell in row:
                over('cell', cell, b.get('cell'))
    elif lay == 'timeline':
        for label, desc in slide['milestones']:
            over('milestone label', label, b.get('milestone label'))
            over('desc', desc, b.get('desc'))
    items = {'donut': 'segments', 'data_table': 'rows', 'timeline': 'milestones'}.get(lay)
    if items and m['max_items'] and len(slide[items]) > m['max_items']:
        out.append(f"slide {slide['idx']} {lay} {items}: {len(slide[items])} > {m['max_items']}")
    return out


def method_kwargs(slide, const):
    """content.json の1枚を原典メソッドの引数へ変換する（色名→原典constantsのRGB、列幅→Inches）。"""
    from pptx.util import Inches
    kw = {k: v for k, v in slide.items() if k not in ('idx', 'layout', 'key_point')}
    if 'segments' in kw:
        for _, color, _ in kw['segments']:
            if not hasattr(const, color):
                fail('S3', f'原典constantsに色 {color} がありません')
        kw['segments'] = [(float(p), getattr(const, c), label) for p, c, label in kw['segments']]
    if 'milestones' in kw:
        kw['milestones'] = [tuple(m) for m in kw['milestones']]
    if 'col_widths' in kw:
        kw['col_widths'] = [Inches(w) for w in kw['col_widths']]
    if kw.get('bottom_bar'):
        kw['bottom_bar'] = tuple(kw['bottom_bar'])
    return kw


def decimal_donut_labels(slide, segments):
    """Original donut geometry is unchanged; replace truncated legend text only."""
    replacements = {f'{label}  {int(pct * 100)}%': f'{label}  {pct * 100:.1f}%'
                    for pct, _, label in segments}
    changed = []
    for shape in slide.shapes:
        if not shape.has_text_frame or shape.text not in replacements:
            continue
        old, new = shape.text, replacements[shape.text]
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.text = run.text.replace(old, new)
        changed.append({'shape_id': shape.shape_id, 'before': old, 'after': new})
    if len(changed) != len(segments):
        fail('S4', '原典donut凡例の文字を全て特定できません。百分率adapterを見直してください')
    return changed


# ─── 原典 gate ───────────────────────────────────────────────────────

def run_upstream_script(up, script, args, cwd):
    env = dict(os.environ, PYTHONIOENCODING='utf-8')
    env['PYTHONPATH'] = os.pathsep.join(filter(None, [str(up), os.environ.get('PYTHONPATH')]))
    proc = subprocess.run([sys.executable, str(up / 'references/scripts' / script), *args], cwd=cwd, env=env,
                          capture_output=True, text=True, encoding='utf-8', errors='replace')
    print(proc.stdout, end='')
    if proc.stderr:
        print(proc.stderr, end='', file=sys.stderr)
    return proc.returncode


def upstream_s3_checkers(up, slides):
    """原典 gate_check_s3.py の LAYOUT_CHECKERS を読み、各ページに掛かる検査名を記録する。"""
    spec = importlib.util.spec_from_file_location('_upstream_gate_s3', up / 'references/scripts/gate_check_s3.py')
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return {str(s['idx']): [f.__name__ for f in mod.LAYOUT_CHECKERS.get(s['layout'], [mod.check_source])]
            for s in slides}


def assert_public(paths):
    private = set()
    for base in (ROOT, Path.home()):
        s = str(base)
        private |= {s, s.replace('\\', '/'), json.dumps(s)[1:-1]}
    for p in paths:
        text = p.read_text(encoding='utf-8-sig')
        if any(x in text for x in private):
            fail('S5', f'{rel(p)} に端末固有の絶対パスが含まれています')


# ─── main ────────────────────────────────────────────────────────────

def main():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8', errors='replace')  # 原典 save() の絵文字出力対策
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('--render', action='store_true', help='Microsoft PowerPointで4ページのPNG/PDFを書き出す')
    ap.add_argument('--font-data', type=Path, help='oss-font-kit.json の場所')
    ap.add_argument('--out', type=Path, help='出力先（既定: repoでは作例フォルダー、配布ZIPでは rebuild/）')
    args = ap.parse_args()

    case_dir = ROOT if ZIP_MODE else ROOT / 'examples/repositories/mck-native-showcase'
    out = (args.out or (ROOT / 'rebuild' if ZIP_MODE else case_dir)).resolve()
    if not out.is_relative_to(ROOT.resolve()):
        fail('setup', '--out は repository / ZIP の中を指定してください')
    plan_path = case_dir / 'input/plan.json'
    data_path = args.font_data or first_existing([ROOT / 'data/oss-font-kit.json', ROOT / 'input/oss-font-kit.json'])
    if not plan_path.exists():
        fail('setup', f'{rel(plan_path)} がありません')
    if not data_path or not Path(data_path).exists():
        fail('setup', 'oss-font-kit.json が見つかりません（--font-data で指定）')
    data_path = Path(data_path).resolve()
    data_label = rel(data_path) if data_path.is_relative_to(ROOT.resolve()) else data_path.name
    plan, data = load(plan_path), load(data_path)
    up = find_upstream()
    MckEngine, const = load_engine(up, plan['fonts'])
    matrix = read_layout_matrix(up / 'references/layout-matrix.yaml')

    project = out / PROJECT_DIR
    deck = out / 'deck.pptx'
    record_path = out / 'build-record.json'
    for stale in (record_path, deck, project / 'gate_s3.json', project / 'gate_result.json'):
        stale.unlink(missing_ok=True)  # 前回の合格結果を残さない

    # S1 brief
    st = font_stats(data, plan)
    messages = write_brief(project / 'brief.md', plan, st)
    print(f"[S1] {st['total']}字体 / " + ', '.join(f"{l}={st['by_lang'][l]}" for l in st['langs']))

    # S2 outline
    slides = build_content(plan, st)
    outline = {'brief': {'audience': plan['brief']['audience'], 'goal': plan['brief']['goal'],
                         'duration_minutes': plan['brief']['duration_minutes'], 'key_messages': messages},
               'slides': [{k: s[k] for k in ('idx', 'layout', 'title', 'key_point')} for s in slides]}
    dump(project / 'outline.json', outline)
    if issues := check_outline(outline, matrix):
        fail('S2', ' / '.join(issues))
    print('[S2] outline: ' + ' → '.join(s['layout'] for s in slides))

    # S3 content + 事前確認 + 原典 gate_check_s3.py
    for s in slides:
        if s['layout'] not in LAYOUTS:
            fail('S3', f"{s['layout']} は今回使う原典メソッドに含めていません")
        try:
            inspect.signature(getattr(MckEngine, s['layout'])).bind(None, **method_kwargs(s, const))
        except TypeError as e:
            fail('S3', f"slide {s['idx']} {s['layout']} の引数が原典メソッドと一致しません: {e}")
    if issues := [i for s in slides for i in budget_issues(s, matrix)]:
        fail('S3', 'layout-matrix.yaml の文字数・項目数を超えています: ' + ' / '.join(issues))
    dump(project / 'content.json', {'title': plan['title'], 'units': {'col_widths': 'inch'},
                                    'colors': 'mck_ppt.constants の名前', 'slides': slides})
    code = run_upstream_script(up, 'gate_check_s3.py', ['content.json', '.'], project)
    if not (project / 'gate_s3.json').exists():
        fail('S3', f'原典 gate_check_s3.py が gate_s3.json を書かずに終了しました（exit {code}）')
    gate_s3 = load(project / 'gate_s3.json')
    if gate_s3.get('passed') is not True:
        fail('S3', 'gate_s3.json passed=false: ' + ' / '.join(
            f"slide {i.get('slide_idx')} [{i.get('check')}] {i.get('message')}" for i in gate_s3.get('fail_items', [])))

    # S4 原典メソッドで描画 + 原典 gate_check.py
    content = load(project / 'content.json')  # 門禁を通ったファイルそのものから描く
    eng = MckEngine(total_slides=len(content['slides']))
    label_changes = []
    for s in content['slides']:
        generated_slide = getattr(eng, s['layout'])(**method_kwargs(s, const))
        if s['layout'] == 'donut':
            label_changes += decimal_donut_labels(generated_slide, s['segments'])
    eng.save(str(deck))
    from pptx import Presentation
    if len(Presentation(str(deck)).slides) != len(content['slides']):
        fail('S4', 'PPTXのページ数が content.json と一致しません')
    shadow = Path(os.path.expanduser('~/.workbuddy/skills/mck-ppt-design/mck_ppt'))
    if shadow.exists():
        fail('S4', '原典gate_check.pyが優先読込する ~/.workbuddy/skills/mck-ppt-design があり、固定原典のQAを保証できません')
    code = run_upstream_script(up, 'gate_check.py', ['../deck.pptx', '.'], project)
    if not (project / 'gate_result.json').exists():
        fail('S4', f'原典 gate_check.py が gate_result.json を書かずに終了しました（exit {code}）')
    gate_s4 = load(project / 'gate_result.json')
    if gate_s4.get('passed') is not True:
        detail = ' / '.join(f"slide {e.get('slide')} [{e.get('category')}] {e.get('message')}"
                            for e in gate_s4.get('user_code_error_detail', []))
        fail('S4', 'gate_result.json passed=false: ' + (detail or gate_s4.get('error', '詳細なし')))

    # PowerPoint 描画（任意）
    render = {'run': False}
    public = [project / f for f in ('outline.json', 'content.json', 'gate_s3.json', 'gate_result.json')]
    if args.render:
        ps1 = first_existing([ROOT / 'tools/render_powerpoint.ps1', ROOT / 'source/render_powerpoint.ps1'])
        shell = shutil.which('pwsh') or shutil.which('powershell')
        if not ps1 or not shell:
            fail('render', 'render_powerpoint.ps1 または PowerShell が見つかりません')
        rdir = out / 'render'
        for f in [*rdir.glob('slide-*.png'), rdir / 'deck.pdf', rdir / 'powerpoint-render.json']:
            f.unlink(missing_ok=True)
        code = subprocess.run([shell, '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(ps1),
                               '-PresentationPath', str(deck), '-OutputDirectory', str(rdir)]).returncode
        if code != 0:
            fail('render', f'{rel(ps1)} が exit {code} で終了しました')
        report_path = rdir / 'powerpoint-render.json'
        report = load(report_path)
        report['source'] = '../deck.pptx'  # 描画scriptが書く絶対パスを公開用に相対化
        dump(report_path, report)
        pngs = sorted(p.name for p in rdir.glob('slide-*.png'))
        if report.get('slide_count') != len(content['slides']) or len(pngs) != len(content['slides']) \
                or not (rdir / 'deck.pdf').exists():
            fail('render', f"PowerPoint出力が不足しています（slides={report.get('slide_count')}, png={len(pngs)}）")
        render = {'run': True, 'application': report.get('application'), 'version': report.get('version'),
                  'pages': [f'render/{n}' for n in pngs], 'pdf': 'render/deck.pdf',
                  'reported_fonts': sorted({t.get('font') for s in report.get('slides', [])
                                            for t in s.get('text_shapes', []) if t.get('font')}),
                  'source_path_relativized': True}
        public.append(report_path)

    # S5 記録（相対パスのみ。モデル実行ログは含めない）
    record = {
        'id': plan['id'], 'title': plan['title'],
        'upstream': {**upstream_identity(up), 'location': rel(up),
                     'sha256': {f: sha(up / f) for f in UPSTREAM_REQUIRED if f != 'mck_ppt/__init__.py'}},
        'engine_methods': [s['layout'] for s in content['slides']] + ['save'],
        'engine_overrides': [], 'custom_renderer': False,
        'font_overrides': {'module': 'mck_ppt.constants', 'scope': 'この生成プロセス内のみ', 'values': plan['fonts']},
        'label_format_adapter': {'reason': '原典donutのint切捨て凡例を小数1桁へ表記。図形・位置・サイズは不変。', 'changes': label_changes},
        'data': {'file': data_label, 'sha256': sha(data_path), **summary_json(st)},
        'stages': {
            'S1': {'file': f'{PROJECT_DIR}/brief.md', 'check': 'audience / goal / key_messages 非空（コード確認）'},
            'S2': {'file': f'{PROJECT_DIR}/outline.json', 'check': 'SKILL S2門禁の項目をコード確認（動詞判定は除く）'},
            'S3': {'file': f'{PROJECT_DIR}/content.json', 'precheck': 'layout-matrix.yaml char_budget / Max Items',
                   'gate': '原典 references/scripts/gate_check_s3.py（未改変・subprocess）',
                   'result': f'{PROJECT_DIR}/gate_s3.json', 'passed': gate_s3['passed'],
                   'checkers': upstream_s3_checkers(up, content['slides'])},
            'S4': {'file': 'deck.pptx', 'gate': '原典 references/scripts/gate_check.py（未改変・subprocess・PYTHONPATH=固定原典）',
                   'result': f'{PROJECT_DIR}/gate_result.json', 'passed': gate_s4['passed'],
                   'overall_score': gate_s4.get('overall_score'), 'checklist': gate_s4.get('checklist')},
        },
        'pptx_sha256': sha(deck), 'render': render,
        'deviations': ['原典SKILLのエンジン固定パス(~/.workbuddy)ではなく固定commitの原典を読み込む',
                       '原典brand-guideのKaiTi指定を日本語書体へ局所変更（色・位置・字号は原典のまま）',
                       'TaskCreate による進捗管理は行わず、各段階のファイルとJSONで記録する'],
        'not_done_by_generator': ['S5 Self-Refinement（experiences/への記録）',
                     *([] if args.render else ['PowerPointによるPNG/PDF描画（--render 未指定）']),
                     '全ページの目視確認', '独立レビュー', 'フォーク梱包とプロベナンス'],
    }
    dump(record_path, record)
    assert_public([project / 'brief.md', record_path, *public])
    print(f"[S4] gate passed / score {gate_s4.get('overall_score')} → {rel(deck)}")


if __name__ == '__main__':
    main()
