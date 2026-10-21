"""Original Japanese-corporate native template editing; bounded Codex repair."""
from pathlib import Path
import copy,json
from pptx.util import Inches,Pt
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
CASE=Path(__file__).resolve().parents[1]
def dump(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def write(p,v):p.write_text(v,encoding='utf-8')
EC=json.loads((CASE/'input/data.json').read_text(encoding='utf-8'))
PLAN=json.loads((CASE/'source/template-map.json').read_text(encoding='utf-8'))
GTITLES=[p['title'] for p in PLAN];TITLE=EC['title']
NOTES=['架空のEC月次報告。今回の判断は商品説明修正と粗利分析への着手。広告増額は別判断である。', '売上は20%増えた一方、粗利率は2ポイント低下し返品率は2ポイント上昇。粗利額は380万円から432万円へ52万円増えた。率と額を分けて読む。', '増収200万円のうち自然検索90万円で45%、メール50万円で25%。合わせて70%。売上寄与は利益改善の証拠ではない。', '返品理由を顧客対応が分類し、商品担当が説明を修正する。説明と実物の差は仮説で、分類・修正記録・返品率を照合して判断する。', '上の四つの帯は段階名。実際の受け渡しは担当をまたぐ線で示す。商品・販売が対象一覧を顧客対応へ渡し、顧客対応は返品理由を分類する。分類表を受けた事業責任者は、理由分類が完了したかを判断する。不足なら再分類へ戻す。完了なら修正対象を商品・販売へ渡す。改稿記録を顧客対応が照合し、新たな返品記録を商品・販売へ渡す。効果と商品別粗利がそろってから事業責任者が広告増額を判断する。分類率100%は10月9日の目標で実績ではない。', '提案・要確認・継続審議を色と文字で区別する。担当と期限を確認し、次回は分類表・改稿記録・粗利差分を持ち寄る。着手や承認の実績は示していない。']

def font(shape,size,bold=False,color='111111'):
    from pptx.dml.color import RGBColor
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.name='BIZ UDPゴシック';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)

def repair_metrics(slide):
    from pptx.dml.color import RGBColor
    sh=list(slide.shapes)
    sh[10].left=Inches(.573);sh[10].top=Inches(1.45);sh[10].height=Inches(.42)
    font(sh[10],21,True)
    table=slide.shapes.add_table(4,4,Inches(.573),Inches(2.02),Inches(12.188),Inches(2.02)).table
    widths=[3.7,2.85,2.85,2.788]
    values=[['指標','2025年9月','2026年9月','変化'],['売上','1,000万円','1,200万円','20%増'],['粗利率','38%','36%','2ポイント低下'],['返品率','4%','6%','2ポイント上昇']]
    for col,w in zip(table.columns,widths):col.width=Inches(w)
    for i,row in enumerate(table.rows):
        for j,cell in enumerate(row.cells):
            set_text(cell,values[i][j]);cell.margin_left=Inches(.13);cell.margin_right=Inches(.10);cell.margin_top=Inches(.07);cell.margin_bottom=Inches(.05)
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string('2474B5' if i==0 else 'FFFFFF' if i%2 else 'EEF4F9')
            font(cell,16 if i==0 else 17,i==0,'FFFFFF' if i==0 else 'C00000' if j==3 and i>=2 else '111111')
    slide.shapes[-1].name='page-02-native-metric-table'
    sh[11].top=Inches(4.16);sh[11].height=Inches(.40);font(sh[11],12)
    sh[12].top=Inches(4.79);sh[12].height=Inches(1.26)
    sh[13].top=Inches(4.92);sh[13].height=Inches(.34);sh[13].width=Inches(11.8);font(sh[13],17,True)
    sh[14].top=Inches(5.37);sh[14].height=Inches(.56);font(sh[14],14)

def repair_process(slide):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    sh=list(slide.shapes);cols=[2.2,5.0,7.8,10.6];ys=[1.95,3.20,4.45];w=2.1;h=.86
    for r,y in enumerate(ys):
        role=sh[[19,29,39][r]];role.top=Inches(y);role.height=Inches(h)
        label=sh[[20,30,40][r]];label.top=Inches(y+.23);label.height=Inches(.4)
        for col,x in enumerate(cols):
            idx=21+r*10+col*2
            for item in [sh[idx],sh[idx+1]]:item.left=Inches(x);item.top=Inches(y);item.width=Inches(w);item.height=Inches(h)
            tf=sh[idx+1].text_frame;tf.margin_left=Inches(.06);tf.margin_right=Inches(.06);tf.margin_top=Inches(.04);tf.margin_bottom=Inches(.04);tf.vertical_anchor=MSO_ANCHOR.MIDDLE
            font(sh[idx+1],12.5)
    # The original loose diamond becomes a real gate in the responsible person's lane.
    gate=sh[49];gate.left=Inches(5.15);gate.top=Inches(1.95);gate.width=Inches(1.8);gate.height=Inches(.86);gate.name='page-05-decision-classification-complete'
    label=sh[50];label.left=Inches(5.39);label.top=Inches(2.24);label.width=Inches(1.32);label.height=Inches(.28);font(label,13,True)
    # Unused stage/role slots are deliberately removed to leave room for the actual route.
    for idx in [23,24,25,26,33,34]:
        el=sh[idx]._element;el.getparent().remove(el)
    for idx in [51,52]:sh[idx].left=Inches(2.2);sh[idx].top=Inches(5.68);sh[idx].width=Inches(5.9);sh[idx].height=Inches(.36)
    for idx in [53,54]:sh[idx].left=Inches(8.3);sh[idx].top=Inches(5.68);sh[idx].width=Inches(4.64);sh[idx].height=Inches(.36)
    font(sh[52],11.5,True,'092B4C');font(sh[54],11.5,True,'C00000')
    for idx in [52,54]:sh[idx].text_frame.margin_top=Inches(.05);sh[idx].text_frame.margin_bottom=0
    plan=json.loads((CASE/'input/flow-plan.json').read_text('utf-8'))
    for edge in plan['edges']:
        pts=edge['points'];color=edge.get('color','2474B5')
        for i,(a,b) in enumerate(zip(pts,pts[1:])):
            line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]),Inches(a[1]),Inches(b[0]),Inches(b[1]))
            line.line.color.rgb=RGBColor.from_string(color);line.line.width=Pt(1.5);line.name='page-05-handoff-'+edge['id']+'-'+str(i+1)
            if i==len(pts)-2:
                end=OxmlElement('a:tailEnd');end.set('type','triangle');line._element.spPr.get_or_add_ln().append(end)
            # Connectors stay behind the node/text objects, visible in the reserved gutters.
            el=line._element;el.getparent().remove(el);slide.shapes._spTree.insert(12,el)
        if edge.get('label'):
            x,y,lw,lh=edge['label_box'];label=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(lw),Inches(lh));set_text(label,edge['label']);label.name='page-05-handoff-label-'+edge['id'];label.text_frame.margin_left=0;label.text_frame.margin_right=0;label.text_frame.margin_top=0;label.text_frame.margin_bottom=0;font(label,11,False,color)
            for p in label.text_frame.paragraphs:p.alignment=PP_ALIGN.CENTER
    dump(CASE/'source/process-relationships.json',plan)
def prepare(_):return CASE,CASE/'source/upstream'
def set_text(shape,text):
    """Replace wording while retaining the template's first paragraph/run formatting."""
    tf=shape.text_frame
    p0=tf.paragraphs[0]
    ppr=copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    rpr=copy.deepcopy(p0.runs[0]._r.rPr) if p0.runs and p0.runs[0]._r.rPr is not None else None
    tf.clear()
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if ppr is not None:
            if p._p.pPr is not None:p._p.remove(p._p.pPr)
            p._p.insert(0,copy.deepcopy(ppr))
        r=p.add_run();r.text=line
        if rpr is not None:r._r.insert(0,copy.deepcopy(rpr))

def build_gonta():
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_LEGEND_POSITION
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    d,u=prepare('japanese-corporate')
    template=u/'skills/japanese-corporate-pptx/assets/japanese-corporate-ponchi-template.pptx'
    prs=Presentation(template)
    selected=[1,2,11,9,8,14]
    slides=[prs.slides[i-1] for i in selected]
    # Reorder existing slide ids, then remove unselected relationships; native chart and table stay intact.
    ids=list(prs.slides._sldIdLst)
    for el in ids:prs.slides._sldIdLst.remove(el)
    for n in selected:prs.slides._sldIdLst.append(ids[n-1])
    for i,el in enumerate(ids,1):
        if i not in selected:prs.part.drop_rel(el.rId)
    common={7:'くらし便 EC事業部　架空の制作例'}
    maps=[{int(k):v for k,v in p['text_replacements'].items()} for p in PLAN]
    for i,(s,m) in enumerate(zip(slides,maps)):
        for idx,text in {**common,8:str(i+1),**m}.items():set_text(s.shapes[idx],text)
        s.notes_slide.notes_text_frame.text=NOTES[i]
        for j,shape in enumerate(s.shapes):shape.name=f'page-{i+1:02d}-object-{j:03d}'
    chart=slides[2].shapes[11].chart
    data=CategoryChartData();data.categories=[c['name'] for c in EC['channels']]
    data.add_series('2025年9月',[c['previous_man_yen'] for c in EC['channels']]);data.add_series('2026年9月',[c['sales_man_yen'] for c in EC['channels']]);chart.replace_data(data)
    chart.has_legend=True;chart.legend.position=XL_LEGEND_POSITION.BOTTOM;chart.legend.include_in_layout=False;chart.legend.font.size=Pt(10)
    chart.value_axis.minimum_scale=0;chart.value_axis.maximum_scale=500;chart.value_axis.major_unit=100
    chart.chart_title.text_frame.text='売上／万円'
    for series,color in zip(chart.series,['777777','2474B5']):series.format.fill.solid();series.format.fill.fore_color.rgb=RGBColor.from_string(color)
    rows=[['No.','論点','把握している状況','対応方針・判断条件','担当','期限','状態'],['1','返品理由','返品率 4% → 6%','上位20商品の理由を分類','森','10/09','提案'],['2','商品説明','説明と実物の差は仮説','返品上位5商品の説明を改稿','青木','10/16','提案'],['3','粗利率','38% → 36%','値引きと商品構成を分解','山田','10/16','要確認'],['4','施策の効果','改善後の結果は未測定','対象商品の返品率を週次比較','青木','週次','要確認'],['5','広告増額','売上20%増、粗利率2ポイント低下','施策の完了と分析後に判断','責任者','次回','継続審議']]
    for row,vals in zip(slides[-1].shapes[14].table.rows,rows):
        for cell,text in zip(row.cells,vals):set_text(cell,text)
    status_colors={'提案':'2474B5','要確認':'C00000','継続審議':'666666'}
    for row in slides[-1].shapes[14].table.rows:
        cell=row.cells[6]
        if cell.text in status_colors:
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(status_colors[cell.text])
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb=RGBColor.from_string('FFFFFF');run.font.bold=True
    for s in slides:
        for sh in s.shapes:
            if sh.has_text_frame and '{{' in sh.text:raise ValueError(sh.text)
            if sh.has_table:
                assert not any('{{' in c.text for row in sh.table.rows for c in row.cells)
    # Explicit local visual repairs: no change to the native chart/table engine.
    for shape in slides[1].shapes:
        if shape.has_text_frame:
            for par in shape.text_frame.paragraphs:
                for run in par.runs:run.font.italic=False
    slides[1].shapes[11].height=Inches(.57)
    from pptx.enum.text import PP_ALIGN
    for par in slides[1].shapes[9].text_frame.paragraphs:
        par.alignment=PP_ALIGN.LEFT
        for run in par.runs:run.font.color.rgb=RGBColor.from_string('111111')
    for slide in slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for par in shape.text_frame.paragraphs:
                    for run in par.runs:
                        run.font.underline=False
                        if run.font.name and ('Noto Sans' in run.font.name):run.font.name='Noto Sans JP'
    repair_metrics(slides[1])
    repair_process(slides[4])
    write(CASE/'NOTES.md','\n\n'.join('## '+str(i+1)+'. '+t+'\n\n'+n for i,(t,n) in enumerate(zip(GTITLES,NOTES))))
    prs.core_properties.title=TITLE;prs.core_properties.subject='架空のデータによる原典テンプレート実行例';prs.core_properties.language='ja-JP';prs.core_properties.author='';prs.core_properties.last_modified_by=''
    prs.save(d/'deck.pptx')
    dump(d/'source/template-map.json',[{'page':i+1,'source_slide':n,'title':t,'text_replacements':m} for i,(n,t,m) in enumerate(zip(selected,GTITLES,maps))])
    write(d/'source/build.log','Loaded official 15-slide PPTX. Selected source slides '+str(selected)+'. Replaced text preserving styles; replaced native chart data; edited native issue table; added native metric table and ten labelled handoff relations with a connected decision; removed other slide relationships. Saved and reopened with python-pptx.\n')
    assert len(Presentation(d/'deck.pptx').slides)==6
    print('Gonta native template saved, six slides')
if __name__=='__main__':build_gonta()
