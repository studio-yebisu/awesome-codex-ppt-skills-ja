"""Two upstream-native Japanese examples: Jinba HTML parts and Gonta PPTX template.

Run with the catalog as cwd and upstream snapshots in .build/actual-upstreams/.
No upstream files are modified. Each path uses its repository's own entry point.
"""
from __future__ import annotations
import argparse, copy, hashlib, importlib.metadata, json, os, re, shutil, subprocess, sys, zipfile
from pathlib import Path
from datetime import datetime, timezone
import xml.etree.ElementTree as ET

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'examples/repositories'
UP=ROOT/'.build/actual-upstreams'
REPOS={'jinba-consulting':'carnot-tech/jinba-consulting-pptx-skill','japanese-corporate':'gonta223/japanese-corporate-pptx-skill'}
EC=json.loads((ROOT/'data/case-inputs.json').read_text(encoding='utf-8-sig'))['ec']
TITLE='くらし便 EC事業 9月報告'
PARTS=['b01','b09','m05','b06','b05','m35','b10']
JTITLES=[TITLE,'売上は20%増え、粗利率と返品率が悪化','自然検索が増収200万円の45%を占める','粗利額は52万円増え、率の低下要因は未確認','返品理由を分類してから商品説明を改稿','商品説明の修正と粗利分析を先行する','くらし便']
GTITLES=[TITLE,'売上増と利益・返品の変化を分けて確認','自然検索が増収200万円の45%を占める','返品理由の分類を、説明修正と効果確認につなぐ','販売・商品・顧客対応で確認結果を引き継ぐ','商品説明の修正と粗利分析を先行する']

def write(p,s):
    p=Path(p);p.parent.mkdir(parents=True,exist_ok=True);p.write_text(s,encoding='utf-8')
def dump(p,obj):write(p,json.dumps(obj,ensure_ascii=False,indent=2)+'\n')
def read(p):return Path(p).read_text(encoding='utf-8-sig')
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def run(argv,log,*,check=True,env=None):
    r=subprocess.run([str(x) for x in argv],cwd=ROOT,env={**os.environ,'PYTHONIOENCODING':'utf-8',**(env or {})},capture_output=True,text=True,encoding='utf-8',errors='replace')
    write(log,r.stdout+r.stderr)
    if check and r.returncode:raise RuntimeError(read(log))
    return r
def prepare(cid):
    d=OUT/cid;(d/'source').mkdir(parents=True,exist_ok=True)
    dump(d/'input/data.json',EC)
    write(d/'input/BRIEF.md','# '+TITLE+'\n\n架空のEC事業の月次会議。事業責任者・販売・商品・顧客対応の担当者が、商品説明修正と粗利分析の着手を判断する。組織・計数・担当者は制作例用。広告増額は確認後に判断する。\n')
    return d,UP/REPOS[cid].replace('/','__')

def build_jinba_html():
    d,u=prepare('jinba-consulting')
    run([sys.executable,u/'scripts/new_deck.py','--parts',','.join(PARTS),'--title',TITLE,'-o',d/'source/seed.html'],d/'source/new-deck.log')
    h=read(d/'source/seed.html')
    original=re.findall(r'<section\b.*?</section>',h,re.S)
    def bar():return '<div class="bar"><div class="logo">'+TITLE+'</div><div class="date"></div></div>'
    def foot(n):return '<div class="foot"><b>'+TITLE+'</b><span>'+str(n)+'</span></div>'
    def base(n,title,body):return '<section class="s">'+bar()+'<h1>'+title+'</h1><div class="c">'+body+'</div>'+foot(n)+'</section>'
    def more(n,title,body):return '<section class="slide slide--no-title-rule"><div class="sinner"><div class="hdr"><div class="logo">'+TITLE+'</div><div class="date"></div></div><h1 class="title">'+title+'</h1><div class="content">'+body+'</div><footer class="footer"><span class="source"></span><span>'+str(n)+'</span></footer></div></section>'
    cover=original[0].replace('Text 1','くらし便<br>EC事業 9月報告').replace('ラベル 1：Text 2に基づく整理','2026年9月／前年同月比較<br>架空の組織・数値による制作例').replace('パーツ01｜表紙','月次会議')
    metrics=[['売上','1,000万円','1,200万円','20%増'],['注文','3,500件','4,000件','500件増'],['粗利率','38%','36%','2ポイント低下'],['返品率','4%','6%','2ポイント上昇']]
    table='<table class="eq"><tr>'+''.join('<th'+(' class="ax"' if i==0 else '')+'>'+x+'</th>' for i,x in enumerate(['指標','2025年9月','2026年9月','前年差']))+'</tr>'+''.join('<tr>'+''.join('<td'+(' class="ax"' if i==0 else '')+'>'+x+'</td>' for i,x in enumerate(row))+'</tr>' for row in metrics)+'</table>'
    # m05's bar-chart/bar-wrap/insight-panel vocabulary is retained; all four bars share the same scale.
    chart='<div class="two-col"><div><div class="section-label">チャネル別の増収額／万円</div><div class="bar-chart" style="grid-template-columns:repeat(4,minmax(0,1fr))">'
    for name,value in [('広告',20),('自然検索',90),('メール',50),('リピート直訪',40)]:
        chart+='<div class="bar-wrap"><div class="bar-value">'+str(value)+'</div><div class="bar" style="height:'+str(value)+'%"></div><div class="bar-label">'+name+'</div></div>'
    chart+='</div></div><div class="insight-panel"><div class="section-label">増収の内訳</div><ul class="bullets"><li>自然検索は90万円増<br>全体の45%を占める</li><li>メールは50万円増<br>自然検索と合わせて70%</li><li>広告は20万円増<br>増収への寄与は10%</li></ul></div></div>'
    gross='<div class="two"><div><div class="colh">売上と粗利率から計算</div><table><tr><td class="ax">前年同月</td><td>1,000万円 × 38%<br><b>粗利額 380万円</b></td></tr><tr><td class="ax">当月</td><td>1,200万円 × 36%<br><b>粗利額 432万円</b></td></tr><tr><td class="ax">差分</td><td>432万円 − 380万円<br><b>52万円の増加</b></td></tr></table></div><div class="tri"></div><div><div class="colh">商品別の分析が必要</div><ul><li>売上増が、粗利率低下の影響を上回った</li><li>値引きと商品構成の影響は仮説である</li><li>販売担当が差分を分解し、原因を確認する</li></ul></div></div>'
    flow='<div class="chev"><div class="cv"><div class="n">10月9日まで</div><div class="t">返品理由を分類</div><div class="x">顧客対応・森<br>上位20商品を対象に<br>分類率100%を確認</div></div><div class="cv"><div class="n">10月16日まで</div><div class="t">商品説明を改稿</div><div class="x">商品・青木<br>返品上位5商品の<br>説明と実物の差を修正</div></div><div class="cv"><div class="n">改稿後は週次</div><div class="t">返品率を比較</div><div class="x">商品・青木<br>対象商品の返品率を比較<br>効果は測定後に判断</div></div></div>'
    flow+='<div class="rows"><div class="r"><div class="l">確認する仮説</div><div>商品説明と実物の差が、返品率上昇の一因か</div></div><div class="r"><div class="l">判断に使う記録</div><div>返品理由の分類表、改稿した説明、週次の返品率</div></div></div>'
    decision='<div class="decision-layout"><div class="decision-ask"><div class="decision-ask-title">今回の提案</div><div class="decision-ask-copy">商品説明の修正と<br>粗利分析を先行</div></div><div class="decision-list"><div class="decision-item"><div>顧客対応・商品担当が返品理由を確認<br>10月16日までに返品上位5商品の説明を改稿</div></div><div class="decision-item"><div>販売担当が10月16日までに粗利差分を分析<br>値引きと商品構成の影響を報告</div></div><div class="decision-item"><div>事業責任者は次回、完了記録と分析結果を確認<br>返品率の週次比較も踏まえて広告増額を判断</div></div></div></div>'
    back=original[-1].replace('会社名','くらし便').replace('連絡先','EC事業部／架空の制作例').replace('パーツ10｜裏表紙','')
    sections=[cover,base(1,JTITLES[1],table),more(2,JTITLES[2],chart),base(3,JTITLES[3],gross),base(4,JTITLES[4],flow),more(5,JTITLES[5],decision),back]
    for old,new in zip(original,sections):h=h.replace(old,new,1)
    # Local adjustments apply only to this deck. Upstream source and parts remain untouched.
    css='''
/* EC example: preserve original warm skin, serif titles, parts and 0.8 scaling. */
:root{--font-sans:"Yu Gothic";--font-serif:"Yu Mincho"}
.s .chev .t{font-size:15pt}.s .chev .x{font-size:12pt;line-height:1.7}
.s .chev .cv{padding-top:9mm;padding-bottom:9mm}
.s .two{grid-template-columns:1fr auto 1fr}.s .two ul{font-size:14pt;line-height:1.8}
.s .two .colh{height:13mm;box-sizing:border-box;display:flex;align-items:flex-end}
.slide .bar-chart{height:450px}.slide .bar{background:var(--accent)}
.slide .bar-wrap{position:relative}.slide .bar-label{white-space:nowrap}
.slide .insight-panel{background:transparent;border:0;padding:0 0 0 30px}
.slide .insight-panel li{font-size:24px;line-height:1.75;margin-bottom:24px}
.slide .decision-ask{background:transparent;border:1.8px solid var(--navy);color:var(--navy);padding:24px;box-sizing:border-box}
.slide .decision-ask-title,.slide .decision-ask-copy{color:var(--navy)}
.slide .decision-item{grid-template-columns:1fr;font-size:24px;line-height:1.65}
'''
    h=h.replace('</style>',css+'</style>',1)
    write(d/'deck.html',h)
    dump(d/'source/part-selection.json',[{'page':i+1,'part':p,'title':t} for i,(p,t) in enumerate(zip(PARTS,JTITLES))])
    run([sys.executable,u/'scripts/check_deck.py',d/'deck.html'],d/'source/check-html.log')
    print('Jinba HTML written and check_deck passed')

def convert_jinba():
    d,u=prepare('jinba-consulting')
    run([sys.executable,u/'scripts/html_to_pptx.py',d/'deck.html',d/'deck.pptx'],d/'source/html-to-pptx.log')
    # html_dump returns font sizes before CSS zoom, but geometry after zoom. Restore 0.8 on the two more-parts slides.
    from pptx import Presentation
    from pptx.util import Pt
    from pptx.oxml.ns import qn
    prs=Presentation(d/'deck.pptx')
    corrections=[]
    for i,slide in enumerate(prs.slides):
        slide.notes_slide.notes_text_frame.text=JTITLES[i]+'。'+EC['decision']
        if PARTS[i].startswith('m'):
            count=0
            for node in slide._element.iter():
                if node.tag in (qn('a:rPr'),qn('a:defRPr'),qn('a:endParaRPr')) and node.get('sz'):
                    node.set('sz',str(round(int(node.get('sz'))*.8)));count+=1
                if node.tag==qn('a:spcPts') and node.get('val'):
                    node.set('val',str(round(int(node.get('val'))*.8)))
            corrections.append({'page':i+1,'part':PARTS[i],'font_runs_scaled':count,'factor':.8,'reason':'upstream CSS zoom is included in box coordinates but absent from computed font sizes'})
    prs.core_properties.language='ja-JP';prs.save(d/'deck.pptx')
    dump(d/'source/conversion-corrections.json',corrections)
    run([sys.executable,u/'scripts/check_deck.py',d/'deck.pptx'],d/'source/check-pptx.log')

def set_text(shape,text):
    """Replace wording while retaining the template's first paragraph/run formatting."""
    tf=shape.text_frame
    p0=tf.paragraphs[0]
    ppr=copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    rpr=copy.deepcopy(p0.runs[0]._r.rPr) if p0.runs and p0.runs[0]._r.rPr is not None else None
    tf.clear()
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        if ppr is not None:
            if p._p.pPr is not None:p._p.remove(p._p.pPr)
            p._p.insert(0,copy.deepcopy(ppr))
        r=p.add_run();r.text=line
        if rpr is not None:r._r.insert(0,copy.deepcopy(rpr))

def build_gonta():
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_LEGEND_POSITION
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    d,u=prepare('japanese-corporate')
    template=u/'skills/japanese-corporate-pptx/assets/japanese-corporate-ponchi-template.pptx'
    prs=Presentation(template)
    selected=[1,2,11,9,8,14]
    slides=[prs.slides[i-1] for i in selected]
    # Reorder existing slide ids, then remove unselected relationships; native chart and table stay intact.
    ids=list(prs.slides._sldIdLst)
    for el in ids:prs.slides._sldIdLst.remove(el)
    for n in selected:prs.slides._sldIdLst.append(ids[n-1])
    for i,el in enumerate(ids,1):
        if i not in selected:prs.part.drop_rel(el.rId)
    common={7:'くらし便 EC事業部　架空の制作例'}
    maps=[
     {1:'EC事業 月次会議',3:'9月報告',9:TITLE,10:'商品説明の修正と粗利分析を先行する',12:'2026年9月／前年同月比較\n読者：事業責任者・販売・商品・顧客対応\n組織・数値・担当者は架空',16:'1分',17:'主要指標と判断の論点',18:'販売担当',20:'2分',21:'チャネル別の増収と利益',22:'販売担当',24:'2分',25:'商品説明修正の進め方',26:'商品担当',28:'2分',29:'着手する施策と次回の確認',30:'参加者',31:'注：実在の組織・業績ではなく、制作手順を比較するためのサンプル。'},
     {3:'主要指標',9:GTITLES[1],10:'● 売上は1,000万円から1,200万円へ20%増加\n● 粗利率は38%から36%へ低下し、返品率は4%から6%へ上昇',11:'対象：2026年9月と2025年9月の同月比較\n率の差は「ポイント」で表記し、売上の増加率と区別する',13:'※ 粗利額と返品率を同時に確認',14:'・前年の粗利額：1,000万円 × 38% = 380万円\n・当月の粗利額：1,200万円 × 36% = 432万円\n・粗利額は52万円増加したが、率の低下要因は未確認',18:'商品説明の修正と粗利分析を先行し、広告増額は次回の確認後に判断する',19:'出所：架空データ「くらし便 2026年9月・前年同月比較」'},
     {1:'チャネル別売上・増収の内訳',3:'定量根拠',9:GTITLES[2],10:'単位：万円／2025年9月と2026年9月を比較',13:'自然検索 +90万円\n増収全体の45%',18:'全体：1,000 → 1,200万円',20:'② 内訳',21:'自然検索90、メール50万円増',23:'③ 留意点',24:'売上だけでは利益改善を判断しない',27:'商品別の粗利差分を確認',29:'判断｜広告増額は分析結果を見てから',30:'注：前年・当月とも架空の集計。チャネル合計は全体売上に一致。',34:'自然検索とメールが増収の70%を占めるが、粗利と返品の確認を先行する',35:'出所：架空データ「くらし便 2026年9月・前年同月比較」'},
     {1:'返品理由から商品説明を見直す',3:'施策のつながり',9:GTITLES[3],10:'商品説明と実物の差は仮説。返品理由の分類で確かめる',17:'返品率が4%から6%へ上昇',19:'理由の内訳は未分類',21:'説明と実物の差は未確認',26:'顧客対応が返品理由を分類',28:'商品担当が上位5商品を改稿',30:'商品担当が返品率を週次比較',35:'返品理由を分類表で確認',37:'説明の修正箇所を記録',39:'実測で改善の有無を判断',41:'対象｜上位20商品の返品理由、返品上位5商品の説明\n前提｜改善効果は未確認　指標｜対象商品の返品率を週次比較',45:'10月9日に分類を確認し、10月16日までに商品説明を改稿する',46:'出所：架空データ「くらし便 9月報告・対応案」。改善後の成果は未測定。'},
     {3:'受け渡しの設計案',9:GTITLES[4],10:'承認前の作業案。記録を引き継ぎ、推測と確認済みの結果を分ける',12:'① 対象を決める',14:'② 理由を確かめる',16:'③ 説明を直す',18:'④ 結果を判断',20:'事業責任者',22:'着手方針を確認\n入力：9月の実績',24:'分類結果を確認\n出力：修正対象',26:'完了状況を確認\n入力：改稿記録',28:'広告増額を判断\n入力：効果・粗利',30:'商品・販売',32:'上位商品を抽出\n出力：対象一覧',34:'商品別粗利を分解\n出力：差分の要因',36:'商品説明を改稿\n出力：修正記録',38:'返品率を週次比較\n出力：効果の記録',40:'顧客対応',42:'返品記録を用意\n出力：返品の一覧',44:'返品理由を分類\n10/09・分類率100%',46:'説明と理由を照合\n入力：修正記録',48:'新たな返品を記録\n出力：理由の推移',50:'確認',52:'引き継ぐ記録｜返品理由・説明の修正・商品別粗利',54:'要確認｜説明と実物の差、値引きと商品構成の影響',58:'担当者が記録と期限を共有し、事業責任者は確認した結果を基に判断する',59:'注：架空の業務設計案。既存のシステム構成や実際の組織を示すものではない。'},
     {3:'着手方針の確認',9:GTITLES[5],11:'対応案\n2026年9月報告',13:'対象範囲｜EC事業・9月実績　　比較｜前年同月　　前提｜架空データ　　未決定｜広告増額',16:'【今回の判断】\n商品説明の修正と粗利分析に着手する案\n判断者：事業責任者／結果は会議後に記録',18:'【継続確認】\n返品率上昇と粗利率低下の要因は未確認\n次回提示物：分類表・改稿記録・粗利差分',22:'次回、各担当が施策の完了と分析結果を示し、事業責任者が広告増額を判断する',23:'出所：架空データ「くらし便 9月報告・対応案」。提案段階であり、実施済みの結果ではない。'}
    ]
    for i,(s,m) in enumerate(zip(slides,maps)):
        for idx,text in {**common,8:str(i+1),**m}.items():set_text(s.shapes[idx],text)
        s.notes_slide.notes_text_frame.text=GTITLES[i]+'。架空のEC月次報告。'+EC['decision']
        for j,shape in enumerate(s.shapes):shape.name=f'page-{i+1:02d}-object-{j:03d}'
    chart=slides[2].shapes[11].chart
    data=CategoryChartData();data.categories=[c['name'] for c in EC['channels']]
    data.add_series('2025年9月',[c['previous_man_yen'] for c in EC['channels']]);data.add_series('2026年9月',[c['sales_man_yen'] for c in EC['channels']]);chart.replace_data(data)
    chart.has_legend=True;chart.legend.position=XL_LEGEND_POSITION.BOTTOM;chart.legend.include_in_layout=False;chart.legend.font.size=Pt(10)
    chart.value_axis.minimum_scale=0;chart.value_axis.maximum_scale=500;chart.value_axis.major_unit=100
    chart.chart_title.text_frame.text='売上／万円'
    for series,color in zip(chart.series,['777777','2474B5']):series.format.fill.solid();series.format.fill.fore_color.rgb=RGBColor.from_string(color)
    rows=[['No.','論点','把握している状況','対応方針・判断条件','担当','期限','状態'],['1','返品理由','返品率 4% → 6%','上位20商品の理由を分類','森','10/09','提案'],['2','商品説明','説明と実物の差は仮説','返品上位5商品の説明を改稿','青木','10/16','提案'],['3','粗利率','38% → 36%','値引きと商品構成を分解','山田','10/16','要確認'],['4','施策の効果','改善後の結果は未測定','対象商品の返品率を週次比較','青木','週次','要確認'],['5','広告増額','売上20%増、粗利率2ポイント低下','施策の完了と分析後に判断','責任者','次回','継続審議']]
    for row,vals in zip(slides[-1].shapes[14].table.rows,rows):
        for cell,text in zip(row.cells,vals):set_text(cell,text)
    status_colors={'提案':'2474B5','要確認':'C00000','継続審議':'666666'}
    for row in slides[-1].shapes[14].table.rows:
        cell=row.cells[6]
        if cell.text in status_colors:
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(status_colors[cell.text])
    for s in slides:
        for sh in s.shapes:
            if sh.has_text_frame and '{{' in sh.text:raise ValueError(sh.text)
            if sh.has_table:
                assert not any('{{' in c.text for row in sh.table.rows for c in row.cells)
    prs.core_properties.title=TITLE;prs.core_properties.subject='架空のデータによる原典テンプレート実行例';prs.core_properties.language='ja-JP';prs.core_properties.author='';prs.core_properties.last_modified_by=''
    prs.save(d/'deck.pptx')
    dump(d/'source/template-map.json',[{'page':i+1,'source_slide':n,'title':t,'text_replacements':m} for i,(n,t,m) in enumerate(zip(selected,GTITLES,maps))])
    write(d/'source/build.log','Loaded official 15-slide PPTX. Selected source slides '+str(selected)+'. Replaced text preserving styles; replaced native chart data; edited native issue table; removed other slide relationships. Saved and reopened with python-pptx.\n')
    assert len(Presentation(d/'deck.pptx').slides)==6
    print('Gonta native template saved, six slides')

def render(cid):
    d=OUT/cid
    run(['pwsh','-NoProfile','-File',ROOT/'tools/render_powerpoint.ps1','-PresentationPath',d/'deck.pptx','-OutputDirectory',d/'render'],d/'render/render-command.log')
    print(cid,'rendered with PowerPoint')

def finalize():
    records=[];guides=[]
    for cid,repo in REPOS.items():
        d=OUT/cid
        if not (d/'deck.pptx').exists():continue
        u=UP/repo.replace('/','__');a=json.loads(read(u/'ACQUISITION.json'))
        audit_path=ROOT/'.build/repository-audit'/repo.replace('/','__')/'record.json'
        audit=json.loads(read(audit_path)) if audit_path.exists() else a
        commit=audit.get('commit',a.get('commit'))
        jinba=cid=='jinba-consulting';titles=JTITLES if jinba else GTITLES
        files=['SKILL.md','references/slide-rules.md','references/ai-smell-lexicon.md','references/content-review-prompt.md','templates/freeform_parts_16x9.html','templates/freeform_parts_more_16x9.html','scripts/new_deck.py','scripts/check_deck.py','scripts/html_to_pptx.py','scripts/html_dump.mjs','scripts/lib_cdp.mjs'] if jinba else ['skills/japanese-corporate-pptx/SKILL.md']+['skills/japanese-corporate-pptx/references/'+f+'.md' for f in ['research-backed-patterns','jtc-authenticity','visual-style','writing-and-density','layout-catalog','qa-checklist']]+['skills/japanese-corporate-pptx/assets/japanese-corporate-ponchi-template.pptx','skills/japanese-corporate-pptx/assets/japanese-corporate-ponchi-layouts.png']
        refs=[dict(path=f,sha256=sha(u/f),url=f'https://github.com/{repo}/blob/{commit}/{f}') for f in files]
        shutil.copy2(u/'LICENSE',d/'LICENSE');shutil.copy2(__file__,d/'source/build_japanese_repository_examples.py')
        mode='html_parts_cli' if jinba else 'instruction_skill_native_template'
        method='公式new_deck.pyで'+','.join(PARTS)+'を選択し、HTMLを組み直して公式html_to_pptx.pyで変換。' if jinba else '公式15型PPTXから1・2・11・9・8・14枚目を選択し、SKILLの規約に従ってネイティブ図形・チャート・表を編集。専用生成APIは同梱されていない。'
        dump(d/'PROVENANCE.json',dict(repository=repo,upstream_commit=commit,execution_type=mode,method=method,source_refs=refs,upstream_files=refs,input_sha256=sha(d/'input/data.json'),artifact_sha256=sha(d/'deck.pptx'),local_changes=['架空の日本語EC月次データ','作例の本文・図表・講稿']+(['原典HTMLパーツの局所CSS調整','追加パーツ2枚のPPTX文字サイズ・行間をCSS zoom 0.8へ合わせて補正','和文ゴシックと明朝のフォント指定をYu GothicとYu Minchoへ限定'] if jinba else ['原典の文字書式・ヘッダー・フッター・図形構造を保持','不要なテンプレートスライドを削除']),license='MIT',render_report='render/powerpoint-render.json',environment='environment.json'))
        dump(d/'environment.json',dict(python=sys.version,platform=sys.platform,packages={p:importlib.metadata.version(p) for p in ['python-pptx','lxml','Pillow']},external_api=False,renderer='Microsoft PowerPoint 16.0',converter='official HTML converter; Node 24 + Chrome headless' if jinba else 'python-pptx native template editing'))
        prompt='添付の架空EC月次データを使い、'+repo+'の固有の制作方法で完結した日本語資料を作成してください。売上増と粗利率低下を区別し、返品理由の分類、商品説明の修正、担当・期限、広告増額の判断条件を含めます。編集可能なPPTXと全ページPDFを出力してください。'
        write(d/'PROMPT.txt',prompt)
        notes=[t+'。'+EC['decision'] for t in titles]
        write(d/'NOTES.md','\n\n'.join('## '+str(i+1)+'. '+t+'\n\n'+n for i,(t,n) in enumerate(zip(titles,notes))))
        p='examples/repositories/'+cid+'/'
        viewer=(OUT/'office-powerpoint-mcp/index.html').read_text(encoding='utf-8')
        viewer=viewer.replace('Math.min(5,i)',f'Math.min({len(titles)-1},i)').replace("+' / 6'",f"+' / {len(titles)}'")
        write(d/'index.html',viewer)
        ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
        with zipfile.ZipFile(d/'deck.pptx') as z:
            assert z.testzip() is None
            slides=[x for x in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',x)]
            metrics=dict(slides=len(slides),native_shapes=sum(len(ET.fromstring(z.read(x)).findall('.//p:sp',ns)) for x in slides),tables=sum(len(ET.fromstring(z.read(x)).findall('.//a:tbl',ns)) for x in slides),charts=len([x for x in z.namelist() if re.fullmatch(r'ppt/charts/chart\d+\.xml',x)]))
        qa=json.loads(read(d/'QA.json')) if (d/'QA.json').exists() else {'status':'pending_visual_review'}
        if qa.get('artifact_sha256') != sha(d/'deck.pptx'):
            qa={'status':'pending_visual_review'}
        qa.update(structure=metrics,artifact_sha256=sha(d/'deck.pptx'),powerpoint_render='render/powerpoint-render.json',package_integrity='pass')
        dump(d/'QA.json',qa)
        req=['Python 3.13 / python-pptx / lxml / Pillow','Microsoft PowerPoint 16.0 と PowerShell 7（PNG/PDF描画）']+(['Node 22以降 / Google Chrome（原典HTML→PPTX変換）'] if jinba else ['同梱の japanese-corporate-ponchi-template.pptx'])
        target='.build/actual-upstreams/'+repo.replace('/','__')
        commands=['git clone https://github.com/'+repo+' '+target,'git -C '+target+' checkout '+str(commit),'python -m venv .venv',r'.venv\Scripts\python -m pip install python-pptx lxml Pillow',r'.venv\Scripts\python tools/build_japanese_repository_examples.py --only '+('jinba' if jinba else 'gonta')+' --render']
        limits=['HTML図形は編集可能なPPTX要素へ変換。SVG/画像は原典仕様で画像になる。作例の棒グラフはCSS矩形。','原典converterのCSS zoom未反映を、追加パーツ2枚のPPTX文字サイズ・行間0.8倍で補正。source/conversion-corrections.jsonに記録。','元check_layout.mjsのブラウザ検査は実行せず、PowerPoint全ページ視認と座標検査で確認。本文の充填率49%・裏表紙15%のWARNを受容。'] if jinba else ['独立した生成プログラムではなく指示SkillとPPTXテンプレートを用いた制作。原典テンプレートの再配布はMIT LICENSEを保持。']
        guides.append(dict(repository=repo,execution_type=mode,summary_ja=method,requirements=req,steps=['原典を固定commitで取得','入力と依頼文を用意','固有パーツまたはテンプレートを選択','本文と図表を作成','PowerPoint本体で全ページ描画'],commands=commands,source_refs=refs,license='MIT',example_ids=[cid],limitations=limits,reproduction=p+'README.md'))
        write(d/'README.md','# '+repo+' の日本語実行例\n\n'+method+'\n\n組織・数値・担当者は架空。input/data.json、PROMPT.txt、NOTES.mdに入力・依頼・講稿を保存。\n\n## 再現手順\n\nカタログ全体を取得し、そのrootで次を実行します。\n\n```powershell\n'+'\n'.join(commands)+'\n```\n\n'+'\n'.join('- '+s for s in req)+'\n\n'+('元CLIのnew_deck → deck.html編集 → check_deck → html_to_pptx → check_deckを実行。source/seed.htmlは選択直後、deck.htmlは完成HTML。source/part-selection.jsonに対応表。' if jinba else '元PPTXをpython-pptxで開き、source/template-map.jsonのページと図形を編集。source/build.logに実施操作。上流の15型テンプレートは保持したまま、選択6ページだけを書き出す。')+'\n\n## 根拠と制約\n\nPROVENANCE.jsonにcommitと原典ファイルSHA、environment.jsonに依存、QA.jsonとrender/powerpoint-render.jsonに検査範囲を保存。LICENSEは原文を保持。\n\n'+'\n'.join('- '+s for s in limits)+'\n')
        records.append(dict(id=cid,title=TITLE,use_case='月次報告',repository=repo,skill_repository=repo,upstream_commit=commit,fork_id=None,page_count=len(titles),format='pptx',preview_kind='powerpoint-render',style='Jinba warm consulting parts' if jinba else 'Japanese corporate template',cover=p+'render/slide-01.png',interior=p+'render/slide-03.png',preview=p+'index.html',artifact_download=p+'deck.pptx',source_download=p+'source/build_japanese_repository_examples.py',input=p+'input/BRIEF.md',input_data=p+'input/data.json',prompt=p+'PROMPT.txt',prompt_text=prompt,provenance=p+'PROVENANCE.json',qa=p+'QA.json',guide=p+'README.md',titles=titles,notes=notes,method=method,sample_notice='組織・数値は架空の制作サンプルです。',pages=[p+f'render/slide-{i:02d}.png' for i in range(1,len(titles)+1)],pdf_download=p+'render/deck.pdf'))
    dump(ROOT/'data/japanese-repository-examples.json',records);dump(ROOT/'data/japanese-repository-guides.json',guides)

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--only',choices=['jinba','gonta','both'],default='both');ap.add_argument('--render',action='store_true');ap.add_argument('--html-only',action='store_true');ap.add_argument('--convert-only',action='store_true');ap.add_argument('--finalize-only',action='store_true');a=ap.parse_args()
    if not a.finalize_only:
        if a.only in ('jinba','both'):
            if not a.convert_only:build_jinba_html()
            if not a.html_only:
                convert_jinba()
                if a.render:render('jinba-consulting')
        if a.only in ('gonta','both'):
            build_gonta()
            if a.render:render('japanese-corporate')
    finalize()
if __name__=='__main__':main()
