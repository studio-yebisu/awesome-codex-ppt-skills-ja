"""独立した最小画像PPTX包装。原典の生成器ではありません。"""
from pathlib import Path
import json,sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
root=Path(__file__).resolve().parents[1]
spec=json.loads((root/'input/page-spec.json').read_text())
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
for item in spec['slides']:
 p=root/'pages'/Path(item['image_file']).name
 w,h=Image.open(p).size
 assert abs(h-w*9/16)<=1
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.shapes.add_picture(str(p),0,0,width=prs.slide_width,height=prs.slide_height)
 s.notes_slide.notes_text_frame.text=item['speaker_notes']
output=Path(sys.argv[1]) if len(sys.argv)>1 else root/'reproduced.pptx'
prs.save(output)
assert len(Presentation(output).slides)==len(spec['slides'])
print(output)
