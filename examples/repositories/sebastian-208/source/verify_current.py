"""Mechanical package/source checks, not a visual or academic peer review."""
from pathlib import Path
import json,hashlib,re
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image
C=Path(__file__).resolve().parents[1]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
 d=json.loads((C/'input/data.json').read_text(encoding='utf-8'));p=Presentation(C/'deck.pptx');rows=[]
 for i,s in enumerate(p.slides,1):
  pics=[x for x in s.shapes if x.shape_type==MSO_SHAPE_TYPE.PICTURE];txt=[x.text for x in s.shapes if x.has_text_frame and x.text];png=C/f'image_pages/slide_{i:02}.png';note=s.notes_slide.notes_text_frame.text.strip();want=d['slides'][i-1]['notes'].strip();blobhash=hashlib.sha256(pics[0].image.blob).hexdigest() if len(pics)==1 else None
  row={'page':i,'full_slide_images':len(pics),'native_body_texts':len(txt),'image_size':list(Image.open(png).size),'embedded_image_sha256':blobhash,'source_image_sha256':sha(png),'image_bytes_match':blobhash==sha(png),'notes_match_input':note==want,'image_covers_slide':len(pics)==1 and pics[0].left==0 and pics[0].top==0 and abs(pics[0].width-p.slide_width)<1000 and abs(pics[0].height-p.slide_height)<1000};rows.append(row)
 report={'status':'pass' if len(rows)==len(d['slides']) and all(r['full_slide_images']==1 and r['native_body_texts']==0 and r['image_bytes_match'] and r['notes_match_input'] and r['image_covers_slide'] and r['image_size']==d.get('image_size',[1280,720]) for r in rows) else 'fail','slide_count':len(rows),'deck_sha256':sha(C/'deck.pptx'),'source_manuscript_sha256':sha(C/'input/manuscript-repair.md'),'pages':rows,'scope':'Mechanical image/notes consistency only; manual source/text/visual review is recorded separately.'}
 out=C/'qa/package-current.json';out.parent.mkdir(exist_ok=True);out.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(report['status'],len(rows));assert report['status']=='pass'
if __name__=='__main__':main()
