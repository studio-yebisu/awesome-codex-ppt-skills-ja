"""Execute MultiAgentPPT's local export function, localize output, and record evidence.

Run from this catalog with .build/ppt-master-venv/Scripts/python.exe.
--render uses the existing PowerPoint renderer. --finalize-only never regenerates.
--reviewed records only a review already performed on these exact rendered pages.
No outline/research/LLM/image API, server, or frontend is invoked.
"""
from pathlib import Path
from copy import deepcopy
from datetime import datetime, timezone
import argparse, hashlib, importlib.metadata, io, json, shutil, subprocess, sys
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'examples/repositories/multiagent-ppt'
UP=ROOT/'.build/actual-upstreams/johnson7788__MultiAgentPPT'
REPO='johnson7788/MultiAgentPPT'
COMMIT='ce8185cee83092290bdb913a528c6e3a72ee879e'
TITLES=['青葉デザイン はじめの一週間','一週間で取り組むこと','困った時は、早めに相談','まずは相談から']
NOTES=[
 '架空の青葉デザインに入社した方へ、最初の一週間の進め方を説明します。全てを一人でできるようになることではなく、仕事の流れと相談先を知ることが目標です。',
 '最初は端末とアカウントを整え、先輩の仕事を見てから練習用の仕事を試します。入れない場所や不明な役割は、解決を待たずに担当へ伝えましょう。',
 '相談するときは、目的と止まった場所、すでに試したことを分けて伝えます。画面共有前の情報確認と、解決後の手順記録も日常の習慣にします。',
 '困ったら、まず相談してください。今日の説明後に、直属の先輩と次の予定と相談先を確認しましょう。これは実在企業の規則ではなく、制作手順を試す架空サンプルです。']
REFS=['LICENSE','frontend/LICENSE','README.md','backend/save_ppt/README.md','backend/save_ppt/ppt_generator.py','backend/save_ppt/ppt_template_0717.pptx','backend/save_ppt/requirements.txt']
PROMPT='johnson7788/MultiAgentPPTの固定原典にあるAPI不要のJSON→PPTX export関数を使い、架空の青葉デザインの新入社員向けに日本語4枚を制作してください。入力を正本とし、固有テンプレートを維持してください。LLM・調査・画像APIは呼ばず、本文と図形を編集可能にし、講稿、原典と局所修正の記録、PowerPointでの全ページ描画を残してください。\n'

def dump(path,data):
    path.parent.mkdir(parents=True,exist_ok=True)
    path.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def rel(path):return path.relative_to(ROOT).as_posix()

def localize(raw):
    p=Presentation(raw)
    assert len(p.slides)==4
    # The original end page inherits a non-placeholder layout shape; make that
    # same shape page-local so its Japanese text is directly editable.
    end=p.slides[3]
    for sh in end.slide_layout.shapes:
        if sh.has_text_frame and sh.text:
            clone=deepcopy(sh._element)
            end.shapes._spTree.insert_element_before(clone,'p:extLst')
            sh.text=''
    for sh in end.shapes:
        if sh.has_text_frame and sh.text:
            first=next((r for par in sh.text_frame.paragraphs for r in par.runs),None)
            if first is not None:
                first.text=TITLES[3]
                for par in sh.text_frame.paragraphs:
                    for r in par.runs:
                        if r._r is not first._r:r.text=''
    for sh in p.slides[0].shapes:
        if sh.shape_id==3:sh.text='架空企業の研修サンプル｜2026年10月'
    ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    for part in list(p.slides)+list(p.slide_layouts)+list(p.slide_masters):
        for t in part._element.findall('.//a:t',ns):
            t.text=(t.text or '').replace('基本信息','青葉デザイン').replace('TRANSPORT','')
        for run in part._element.findall('.//a:r',ns):
            rpr=run.find('a:rPr',ns)
            if rpr is None:rpr=OxmlElement('a:rPr');run.insert(0,rpr)
            rpr.set('lang','ja-JP')
            for tag in ('latin','ea','cs'):
                el=rpr.find('a:'+tag,ns)
                if el is None:el=OxmlElement('a:'+tag);rpr.append(el)
                el.set('typeface','Noto Sans JP')
    for slide,note in zip(p.slides,NOTES):slide.notes_slide.notes_text_frame.text=note
    p.save(OUT/'deck.pptx')

def build():
    for name in REFS:
        target=OUT/'source/upstream'/name;target.parent.mkdir(parents=True,exist_ok=True)
        shutil.copy2(UP/name,target)
    cmd=[sys.executable,'-X','utf8',str(OUT/'source/run_export.py')]
    result=subprocess.run(cmd,cwd=ROOT,capture_output=True,text=True,encoding='utf8')
    log=OUT/'source/export-corrected.log';log.write_text(result.stdout+'\n'+result.stderr,encoding='utf8')
    dump(OUT/'source/commands.json',[{'argv':cmd,'exit_code':result.returncode,'log':rel(log),'log_sha256':digest(log)}])
    if result.returncode:raise RuntimeError('Original exporter failed; see '+str(log))
    data=json.loads((OUT/'input/data.json').read_text(encoding='utf8'))
    raw=OUT/'source/upstream/backend/save_ppt/output_ppts'/(data['title']+'.pptx')
    shutil.copy2(raw,OUT/'source/raw-export.pptx')
    localize(raw)

def finalize(reviewed=False):
    OUT.mkdir(parents=True,exist_ok=True)
    artifact=OUT/'deck.pptx';p=Presentation(artifact)
    report=json.loads((OUT/'render/powerpoint-render.json').read_text(encoding='utf8'))
    assert report['source_sha256'].lower()==digest(artifact)
    assert report['slide_count']==len(p.slides)==4
    pages=[OUT/f'render/slide-{i:02}.png' for i in range(1,5)]
    refs=[{'path':f,'sha256':digest(UP/f),'url':f'https://github.com/{REPO}/blob/{COMMIT}/{f}'} for f in REFS]
    images=[]
    for page,path in enumerate(pages,1):
        with Image.open(path) as im:im.verify()
        images.append({'page':page,'file':rel(path),'sha256':digest(path),'dimensions':[1600,900]})
    (OUT/'PROMPT.txt').write_text(PROMPT,encoding='utf8')
    (OUT/'NOTES.md').write_text('\n\n'.join(f'## {i+1}\n{n}' for i,n in enumerate(NOTES))+'\n',encoding='utf8')
    (OUT/'input/BRIEF.md').write_text('# 青葉デザイン はじめの一週間\n\n架空企業の新入社員向け4枚・3分の説明。目的は仕事の流れと相談先を知ること。input/data.jsonのtitle/sectionsを原典exportへ入力する。企業名・手順は制作用サンプルで、実在企業の規則や外部調査ではない。\n',encoding='utf8')
    shutil.copy2(UP/'LICENSE',OUT/'LICENSE.upstream')
    shutil.copy2(UP/'frontend/LICENSE',OUT/'LICENSE.frontend')
    changes=[
      '原典コードとPPTX templateはbyte同一でコピー。SlideConfigのデータだけを実templateのplaceholder IDへ補正: TITLE_PAGE_DATE 9→3、SUBCHAPTER_TITLE 5→2、3項目 [2,3,4]→[3,5,4]。',
      '生成後のPPTX文字をNoto Sans JP/ja-JPへ設定。内容に無関係で本文に隠れるmasterのTRANSPORTラベルを空にし、基本信息を青葉デザインへ置換。元のレイアウトと装飾は維持。',
      '原典終了pageは文字がlayout上にしかないため同じshapeをpageへ移し、Thanksをまずは相談からへ翻訳。表紙の日付欄を架空サンプル表示へ変更。ページ別日本語ノートを追加。',
      '原典generate_presentationはlistを受け付けずdict(title/sections)を要求。最初のsectionは表紙扱い、内容処理はindex1/2のみ。一般的な多頁exportとして紹介しない。',
      '最初のrequests依存不足と題名改行によるWindows filename失敗は初期ログに保持。requestsをproject venvへ追加し、題名から改行を除いて再実行。']
    dump(OUT/'source/local-changes.json',changes)
    commands=json.loads((OUT/'source/commands.json').read_text(encoding='utf8'))
    prov={'repository':REPO,'upstream_commit':COMMIT,'references':refs,'execution_type':'original_export_function_only','method':'原典start_generate_presentation(dict)と同梱PPTX templateのlocal export。LLM/研究agent/画像API/通常Web frontendは未実行。','commands_executed':commands,'artifact_sha256':digest(artifact),'input_sha256':digest(OUT/'input/data.json'),'generator':'tools/build_multiagent_repository_example.py','render_report':'render/powerpoint-render.json','local_changes':changes,'upstream_source_modified':False,'external_api_calls':0,'network_guard':'requests.Session.request is blocked in the runner; input has no image URLs','checked_at':datetime.now(timezone.utc).isoformat()}
    dump(OUT/'PROVENANCE.json',prov)
    structure={'slides':4,'native_text_shapes':sum(s.has_text_frame and bool(s.text) for slide in p.slides for s in slide.shapes),'picture_shapes':sum(s.shape_type==13 for slide in p.slides for s in slide.shapes),'speaker_notes':4}
    # Round-trip a disposable in-memory copy, leaving the reviewed artifact intact.
    editable=Presentation(artifact)
    shape=next(s for s in editable.slides[0].shapes if s.has_text_frame and s.text)
    shape.text='編集性の確認用テキスト'
    buffer=io.BytesIO();editable.save(buffer);buffer.seek(0)
    assert any(s.has_text_frame and s.text=='編集性の確認用テキスト' for s in Presentation(buffer).slides[0].shapes)
    dump(OUT/'QA.json',{'status':'qa_complete' if reviewed else 'pending_visual_review','page_count':4,'artifact_sha256':digest(artifact),'structure':structure,'images':images,'editability_check':{'status':'passed','method':'一時的なメモリーコピーでnative文字を変更、PPTX保存・再読込して変更の保持を確認。完成原本は変更していない。','powerpoint_interactive_edit_test':'not_run'},'visual_review':{'pages':[1,2,3,4] if reviewed else [],'method':'完成PPTXをMicrosoft PowerPoint16.0で1600×900PNGへ描画し、全4枚を実画像で視認。','result':'日本語本文・3列の順序・末尾の相談案内・文字の収まりを確認。編集可能なnativeテキストを保持。' if reviewed else 'pending'},'not_claimed':['通常のmulti-agent生成','LLM/画像API生成','人による正式採用','外部公開']})
    dump(OUT/'environment.json',{'python':sys.version,'packages':{name:importlib.metadata.version(name) for name in ['python-pptx','Pillow','requests']},'powerpoint':report['version'],'execution':'API-free local exporter only'})
    guide='''# MultiAgentPPT：日本語のローカルexport作例

架空の青葉デザインの新入社員向け4枚。原典の`backend/save_ppt/ppt_generator.py`の`start_generate_presentation`と同梱templateを実行した。JSONの原稿は制作担当が準備しており、multi-agent・調査・LLM・画像APIが生成した原稿ではない。

原典: johnson7788/MultiAgentPPT、固定commit ce8185cee83092290bdb913a528c6e3a72ee879e。READMEは現版の非メンテナンスを明記する。

## 自分の内容へ置き換える

1. Python、python-pptx、Pillow、requestsを用意する。確認バージョンはenvironment.json。キー・API・Web server・DBは不要。
2. input/data.jsonのtitleとsectionsを変更する。各sectionのcontentはh1/p/bulletsのblock。実装はdictを要求するため、原典末尾のlist形式exampleをそのまま渡さない。
3. `python source/run_export.py`で同梱した原典exportを実行する。出力はsource/upstream/backend/save_ppt/output_ppts/内。題名は出力ファイル名になるためWindowsで使えない文字や改行を避ける。
4. 日本語の書体・ノート・固定ラベルの局所化を行い、PowerPointで全ページ確認する。今回の補助コードのlocalize関数が実際の変更内容。新規原稿は別途目視が必要。

原典のcontent処理は最初の3sectionに限定される。今回の4頁は表紙、3項目、本文、終了頁。4番目以降のsectionを追加しても内容が出力されるとは限らない。この作例を通常のmulti-agent全経路・任意枚数の汎用生成の実証として扱わない。

原典のshape ID指定が同梱templateと一致しなかったため、run_export.pyがconfigデータを実placeholder対応へ補正する。元コードファイルの改変なし。ネットワーク要求はrunnerで拒否し、画像URLのない入力を使う。

## カタログ内での再制作

`python tools/build_multiagent_repository_example.py --render`

上記はカタログroot、.build/actual-upstreamsの固定原典、依存導入済みPython、Windows版PowerPoint、tools/render_powerpoint.ps1に依存する。画像を全頁確認してから`--finalize-only --reviewed`で確認記録を更新する。スクリプトの--reviewedは実視認を代行しない。

初期失敗・成功ログとraw-export.pptx、変更台帳はsource/。出力の本文はnative textで編集可能。背景・装飾は原典templateの要素で、画像化したページを貼ったものではない。Noto Sans JPは受け手にも必要。

root LICENSEはMIT Copyright 2025 johnson7788。frontend LICENSEはMIT Copyright 2024 ALLWEONE Team。両原文表示を保持する。原典templateも同梱リポ由来で、別の第三者素材の権利まで保証するものではない。外部画像は取得・使用していない。
'''
    (OUT/'README.md').write_text(guide,encoding='utf8')
    viewer='''<!doctype html><html lang="ja"><meta charset="utf-8"><title>青葉デザイン はじめの一週間</title><style>html,body{margin:0;background:#142032;height:100%}img{display:block;width:100%;height:100%;object-fit:contain}</style><img id="page" alt="PowerPoint描画の日本語スライド"><script>let n=1;function go(i){n=Math.max(1,Math.min(4,i));document.querySelector('img').src='render/slide-'+String(n).padStart(2,'0')+'.png';parent.postMessage({type:'yebisu-page',page:n,total:4},location.origin)}addEventListener('message',e=>{if(e.source!==parent||e.origin!==location.origin)return;if(e.data.type==='yebisu-slide')go(e.data.page)});addEventListener('keydown',e=>{if(e.key==='ArrowRight')go(n+1);if(e.key==='ArrowLeft')go(n-1)});go(1)</script></html>'''
    (OUT/'index.html').write_text(viewer,encoding='utf8')
    shutil.copy2(Path(__file__),OUT/'source/build_multiagent_repository_example.py')
    row={'id':'multiagent-ppt','title':TITLES[0],'use_case':'研修・共有','repository':REPO,'skill_repository':REPO,'upstream_commit':COMMIT,'fork_id':None,'page_count':4,'format':'pptx','preview_kind':'powerpoint-render','style':'原典export/template・編集可能PPTX','cover':rel(pages[0]),'interior':rel(pages[1]),'preview':rel(OUT/'index.html'),'artifact_download':rel(artifact),'source_download':rel(OUT/'source/build_multiagent_repository_example.py'),'input':rel(OUT/'input/BRIEF.md'),'input_data':rel(OUT/'input/data.json'),'prompt':rel(OUT/'PROMPT.txt'),'prompt_text':PROMPT,'provenance':rel(OUT/'PROVENANCE.json'),'qa':rel(OUT/'QA.json'),'guide':rel(OUT/'README.md'),'titles':TITLES,'notes':NOTES,'pages':[rel(v) for v in pages],'pdf_download':rel(OUT/'render/deck.pdf'),'input_sha256':digest(OUT/'input/data.json'),'method':prov['method'],'sample_notice':'架空企業・架空の研修手順。外部調査や実在企業の規則ではない。'}
    dump(ROOT/'data/multiagent-repository-examples.json',[row])
    dump(ROOT/'data/multiagent-repository-guides.json',[{'repository':REPO,'example_ids':['multiagent-ppt'],'production_status':'qa_complete' if reviewed else 'generated','requirements':['Python + python-pptx/Pillow/requests','Noto Sans JP','描画検証にはWindows版PowerPoint'],'steps':['原典のdict入力に合わせてinput/data.jsonを用意。','source/run_export.pyで原典export関数を実行。','日本語固定ラベル・書体・講稿を局所化。','PowerPointで全頁描画し、本人の原稿へ変更した場合も目視する。'],'source_refs':refs,'limitations':['API不要のexport関数のみ実行。multi-agent生成/画像APIの証明ではない。','原典の処理対象sectionに制約があり、任意枚数を保証しない。','同梱templateへのshape ID対応はrunnerで補正。原典は非メンテナンス。'],'guide_markdown':guide}])
    print(json.dumps({'status':'qa_complete' if reviewed else 'pending_visual_review','pages':4,'artifact':rel(artifact),'sha256':digest(artifact)},ensure_ascii=False))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--render',action='store_true');parser.add_argument('--finalize-only',action='store_true');parser.add_argument('--reviewed',action='store_true');a=parser.parse_args()
    if not a.finalize_only:build()
    if a.render:subprocess.run(['pwsh','-NoProfile','-File',str(ROOT/'tools/render_powerpoint.ps1'),'-PresentationPath',str(OUT/'deck.pptx'),'-OutputDirectory',str(OUT/'render')],check=True)
    finalize(a.reviewed)
