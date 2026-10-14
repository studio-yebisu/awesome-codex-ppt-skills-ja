#!/usr/bin/env python3
"""
Purpose: Assemble images/slide_XX.png into a 16:9 PPTX.
Input: Image directory and output PPTX path.
Output: PPTX containing full-slide images, plus optional report.
Failure behavior: If python-pptx/Pillow are unavailable, write a report and exit 0; never install packages.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--images", type=Path, required=True)
    parser.add_argument("--output", "-o", type=Path, required=True)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    report = {"images_dir": str(args.images), "output": str(args.output), "slides": [], "warnings": [], "errors": []}
    images = sorted(list(args.images.glob("slide_*.png")) + list(args.images.glob("slide_*.jpg")) + list(args.images.glob("slide_*.jpeg")))
    if not images:
        report["errors"].append("No slide_XX images found.")
    try:
        from pptx import Presentation  # type: ignore
        from pptx.util import Inches  # type: ignore
    except Exception as exc:
        report["warnings"].append(f"python-pptx unavailable; assembly skipped: {exc}")
        if args.report:
            args.report.parent.mkdir(parents=True, exist_ok=True)
            args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        return 0

    if not report["errors"]:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank = prs.slide_layouts[6]
        while len(prs.slides) > 0:
            break
        for image in images:
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(image), 0, 0, width=prs.slide_width, height=prs.slide_height)
            report["slides"].append(str(image))
        args.output.parent.mkdir(parents=True, exist_ok=True)
        prs.save(args.output)

    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return 0 if not report["errors"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
