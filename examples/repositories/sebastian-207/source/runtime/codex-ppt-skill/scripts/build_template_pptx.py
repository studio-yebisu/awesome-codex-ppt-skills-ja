#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches, Pt


def delete_all_slides(prs: Presentation) -> None:
    sld_id_lst = prs.slides._sldIdLst  # noqa: SLF001
    for sld_id in list(sld_id_lst):
        r_id = sld_id.rId
        prs.part.drop_rel(r_id)
        sld_id_lst.remove(sld_id)


def pick_layout(prs: Presentation, idx: int):
    names = ["Title Slide", "Title and Content", "Title Only", "Blank"]
    if idx == 0:
        for layout in prs.slide_layouts:
            if "title" in layout.name.lower() and "content" not in layout.name.lower():
                return layout
    for wanted in names[1:]:
        for layout in prs.slide_layouts:
            if wanted.lower() in layout.name.lower():
                return layout
    return prs.slide_layouts[min(idx, len(prs.slide_layouts) - 1)]


def set_text(shape, text: str, font_size: int = 22) -> None:
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(font_size)
            run.font.name = "Microsoft YaHei UI"
            run.font.color.rgb = RGBColor(17, 24, 39)
            if font_size >= 26:
                run.font.bold = True


def add_template_safe_structure(slide, idx: int, width, height) -> None:
    if idx == 0:
        return
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.08), width - Inches(1.24), Pt(0.7))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(214, 222, 232)
    line.line.color.rgb = RGBColor(214, 222, 232)
    marker = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.04), Inches(0.7), Pt(1.8))
    marker.fill.solid()
    marker.fill.fore_color.rgb = RGBColor(37, 99, 235)
    marker.line.color.rgb = RGBColor(37, 99, 235)


def first_generated_asset(asset_plan: dict) -> str | None:
    for item in asset_plan.get("assets", []) or asset_plan.get("slides", []) or []:
        if item.get("generation_required") and item.get("status") == "generated" and item.get("output_path"):
            path = Path(str(item.get("output_path")))
            if path.exists():
                return str(path)
    return None


def fill_slide(slide, slide_plan: dict, idx: int, width, height, generated_asset: str | None = None) -> None:
    title = slide_plan.get("title", f"Slide {idx + 1}")
    bullets = slide_plan.get("points") or slide_plan.get("body") or slide_plan.get("speaker_notes_seed") or ""
    if isinstance(bullets, str):
        bullets = [x.strip() for x in bullets.replace("；", ";").split(";") if x.strip()]
    placeholders = list(slide.placeholders)
    title_done = False
    body_done = False
    for ph in placeholders:
        try:
            ph_type = ph.placeholder_format.type
        except Exception:
            continue
        if ph_type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) and not title_done:
            set_text(ph, title, 28 if idx else 34)
            title_done = True
        elif ph.has_text_frame and not body_done:
            ph.text_frame.clear()
            for b_idx, item in enumerate(bullets[:5] or [slide_plan.get("purpose", "")]):
                para = ph.text_frame.paragraphs[0] if b_idx == 0 else ph.text_frame.add_paragraph()
                para.text = str(item)
                para.level = 0
                for run in para.runs:
                    run.font.size = Pt(17)
                    run.font.name = "Microsoft YaHei UI"
                    run.font.color.rgb = RGBColor(71, 85, 105)
            body_done = True
    if not title_done:
        box = slide.shapes.add_textbox(Inches(0.65), Inches(0.45), width - Inches(1.3), Inches(0.6))
        set_text(box, title, 28 if idx else 34)
    if not body_done and idx > 0:
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.55), width - Inches(1.6), Inches(3.8))
        box.text_frame.clear()
        for b_idx, item in enumerate(bullets[:5] or [slide_plan.get("purpose", "")]):
            para = box.text_frame.paragraphs[0] if b_idx == 0 else box.text_frame.add_paragraph()
            para.text = str(item)
            para.level = 0
            for run in para.runs:
                run.font.size = Pt(17)
                run.font.name = "Microsoft YaHei UI"
                run.font.color.rgb = RGBColor(71, 85, 105)
    if idx == 0 and generated_asset:
        try:
            slide.shapes.add_picture(generated_asset, Inches(8.05), Inches(1.35), width=Inches(4.2))
        except Exception:
            pass
    add_template_safe_structure(slide, idx, width, height)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--slides", type=Path, required=True)
    parser.add_argument("--asset-plan", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--profile", type=Path)
    parser.add_argument("--authoring-copy", type=Path)
    args = parser.parse_args()

    plan = json.loads(args.slides.read_text(encoding="utf-8-sig"))
    asset_plan = json.loads(args.asset_plan.read_text(encoding="utf-8-sig")) if args.asset_plan and args.asset_plan.exists() else {}
    generated_asset = first_generated_asset(asset_plan)
    prs = Presentation(str(args.template))
    width, height = prs.slide_width, prs.slide_height
    delete_all_slides(prs)
    for idx, slide_plan in enumerate(plan.get("slides", [])):
        slide = prs.slides.add_slide(pick_layout(prs, idx))
        fill_slide(slide, slide_plan, idx, width, height, generated_asset)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.output)
    if args.profile:
        args.profile.parent.mkdir(parents=True, exist_ok=True)
        profile = {
            "template": str(args.template),
            "output": str(args.output),
            "slide_width": width,
            "slide_height": height,
            "layouts_available": [layout.name for layout in prs.slide_layouts],
            "preserved": ["slide size", "theme", "masters", "layouts", "placeholder system where available"],
        }
        args.profile.write_text(json.dumps(profile, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if args.authoring_copy:
        args.authoring_copy.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(Path(__file__), args.authoring_copy)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
