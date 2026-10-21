"""Actual PPT-Design-Skill Build Mode; no API or office app needed to generate."""
from pathlib import Path
import argparse, hashlib, html, importlib.metadata, json, shutil, subprocess, sys

ROOT=Path(__file__).resolve().parents[1]
DEFAULT=ROOT
COMMIT='cbf7ef1324a026b29e2b77b71e2034554ddd3e22'
REPO='sunchaokun/PPT-Design-Skill'
DATA={'title':'地域カフェの朝営業を試す','cafe':'朝凪カフェ','notice':'架空の店舗の企画案。すべて提案値・仮定値で、実績は未収集。','days':10,'weeks':2,'hours':'8:00–10:00','price':650,'scenarios':[15,25,35], 'titles':['朝の2時間を、2週間だけ試す','メニューと役割を、小さく始める','売上は仮定を変えて比べる','毎日、需要と負担を同じ表で残す','2週間後に、続ける条件を確かめる'], 'notes':['架空の朝凪カフェで、通常10時開店前の8時から10時まで朝営業を試す提案です。2週間の平日10日間に限り、開始前に担当と準備量を決めます。','商品はコーヒーとトーストのセット1種類、仮価格650円です。調理と会計の役割を分け、通常営業の仕込みを圧迫しない方法を試します。','1日の販売数を15・25・35セットと仮定して売上を比較しています。価格650円を掛けた試算で、原価や追加人件費を引いた利益ではありません。','来店数・販売数・待ち時間・廃棄数を毎日記録する計画です。表の空欄は未収集を意味し、ゼロや達成済みという意味ではありません。','10日分を振り返り、需要と現場負担の両方を見て継続・変更・終了を選びます。数値の採否基準は開始前に店内で決め、今回の資料では未合意と明記します。']}
PROMPT='原典PPT-Design-SkillのBuild Modeを使い、「地域カフェの朝営業を試す」を日本語5枚にしてください。架空店舗の企画案と実績未収集を明記。2週間・平日10日・8〜10時・セット650円、15/25/35セットの売上試算を示し、利益と混同しないでください。原典テーマ合成と公開helperを使い、編集可能な文字・図形と講稿を残してください。'
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def save(p,obj):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def write(p,s):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(s,encoding='utf-8')
def command_records(out):
 records=[]
 for p in sorted((out/'source/logs').glob('*.log')):
  text=p.read_text(encoding='utf-8-sig')
  # Subprocess receipts start with their JSON argv. The in-process library
  # build receipt and the separate reproduction JSON are retained as evidence.
  if not text.lstrip().startswith('['):continue
  argv=json.loads(text.splitlines()[0])
  if not isinstance(argv,list):continue
  records.append({'argv':argv,'exit_code':0,'log':p.relative_to(out).as_posix(),'log_sha256':sha(p)})
 return records
def run(argv,out,name):
 p=subprocess.run(list(map(str,argv)),capture_output=True,text=True,encoding='utf-8',errors='replace')
 write(out/'source/logs'/name,json.dumps(list(map(str,argv)),ensure_ascii=False)+'\n'+p.stdout+'\n'+p.stderr)
 if p.returncode:raise RuntimeError(f'{name}: exit {p.returncode}')
 return p.stdout
def prepare(up,out,data):
 out.mkdir(parents=True,exist_ok=True)
 for d in ['input','source/logs','source/task','render']: (out/d).mkdir(parents=True,exist_ok=True)
 # Initializer must run against an absent destination; the task is preserved on regeneration.
 task=out/'source/task/cafe-pilot'
 if not task.exists():run([sys.executable,up/'skill/scripts/init_presentation_task.py','--project',out,'--name','cafe-pilot','--destination',task],out,'initialize.log')
 save(out/'input/data.json',data);write(out/'PROMPT.txt',PROMPT)
 write(out/'input/BRIEF.md',f"# {data['title']}\n\n{data['notice']}\n\n対象: 店主・朝担当スタッフ。目的: 小規模な実証営業の実施条件を共有する。通常営業は10時開始。朝営業は{data['hours']}、{data['weeks']}週間の平日{data['days']}日間。コーヒー＋トーストを{data['price']}円で販売する仮案。\n\n実績・写真・外部データなし。採否基準は店内で事前合意が必要。\n")
 shutil.copyfile(out/'input/BRIEF.md',task/'brief.md')
 write(task/'page-plan.md','# Page Plan\n\n'+'\n'.join(f'{i}. {t} — {data["notes"][i-1]}' for i,t in enumerate(data['titles'],1)))
 write(task/'acceptance-contract.md','# Acceptance Contract\n\n## MUST\n- 日本語5ページ・16:9・頁別講稿・編集可能な文字と図形。\n- 価格×15/25/35セットの計算と、仮定・未収集・利益ではない旨を明示。\n- 全頁PNGの視覚効果確認後に重大欠陥を検査し、未確認なら合格にしない。\n\n## SHOULD\n- 役割と受け渡しを図示し、10日間の記録欄を見せる。\n- 通常営業への切替時刻を明記。\n\n## NICE_TO_HAVE\n- 暖色と和文明朝を使う。外部写真は不要。\n')
 write(task/'visual-direction.md','# Visual Direction\n\n候補professionalは業務帳票寄り、vibrant-startupは彩度が高い。warm-elegantは近いが、試行を落ち着いて読む紙面にはterracottaの暖色とserif-editorialの見出しを選んだ。これはrecommend_stylesの候補そのものではなく、公開ThemeComposerの別atomsを比較して選ぶ自律判断である。\n\n小さな地域カフェの現場共有用。terracotta / serif-editorial / accent-bar / standardを選択し、見出しをShippori Mincho、本文をBIZ UDPGothicの公式静的字体へ明示変更。暖色の紙面と大きな数字。表紙→役割→試算→記録→判断のリズム。写真は使わない。ユーザーの自律制作依頼に基づく選定であり、人間のデザイン採用は未申請。\n')
 dest=out/'source/upstream'
 if dest.resolve()!=up.resolve():
  for name in ['LICENSE','README.md','skill/SKILL.md']:
   if(up/name).exists():(dest/name).parent.mkdir(parents=True,exist_ok=True);shutil.copyfile(up/name,dest/name)
  for name in ['skill/references','skill/scripts','skill/templates/task-init']:shutil.copytree(up/name,dest/name,dirs_exist_ok=True,ignore=shutil.ignore_patterns('__pycache__'))
 shutil.copyfile(up/'LICENSE',out/'LICENSE')
 write(out/'source/requirements.txt','pptx-designer==1.0.0b11\npython-pptx==1.0.2\n')
 return task
def generate(out,task,data):
 if importlib.metadata.version('pptx-designer')!='1.0.0b11':raise RuntimeError('Install pinned pptx-designer==1.0.0b11')
 from pptx_designer import Presentation,validate_resolved_theme,recommend_styles
 from pptx_designer.renderer.theme import ThemeComposer
 from pptx_designer.tools.text import text
 from pptx_designer.tools.shapes import rect as original_rect,rrect as original_rrect,oval as original_oval
 from pptx_designer.tools.charts import bar_chart
 from pptx.util import Inches
 from pptx.oxml.xmlchemy import OxmlElement
 from types import SimpleNamespace
 save(out/'source/style-candidates.json',recommend_styles('warm local coffee shop editorial community pilot'))
 theme=ThemeComposer().compose(palette='terracotta',fonts='serif-editorial',decoration='accent-bar',layout='standard',seed=42)
 theme['typography']={'heading':'Shippori Mincho','body':'BIZ UDPGothic'}
 theme['locale_adjustment']='Japanese heading/body fonts; original palette and other atoms unchanged'
 validate_resolved_theme(theme)
 if theme['source']['fallbacks']:raise RuntimeError('Unexpected theme fallback')
 save(out/'source/resolved-theme.json',theme)
 write(task/'theme-lock.yaml','version: 1\nstatus: LOCKED_FOR_AUTONOMOUS_SAMPLE\nuser_confirmed_at: ""\nconfirmation_basis: user_authorized_autonomous_sample\nvisual_thesis: "朝の小さな実験を暖色の紙面で共有する"\nseed: 42\nintended_tokens:\n  palette: terracotta\n  typography: serif-editorial\n  decoration: accent-bar\n  layout: standard\nresolved_theme_path: ../../resolved-theme.json\nresolved_theme_fingerprint: '+sha(out/'source/resolved-theme.json')+'\n')
 prs=Presentation(theme=theme,strict_theme=True)
 def token(v):return theme['semantic_roles'].get(v,v)
 def rect(s,x,y,w,h,fill,line=None):return original_rect(s,x,y,w,h,token(fill),line=token(line))
 def rrect(s,x,y,w,h,fill,line=None):return original_rrect(s,x,y,w,h,token(fill),line=token(line))
 def oval(s,x,y,w,h,fill,line=None):return original_oval(s,x,y,w,h,token(fill),line=token(line))
 def tx(s,x,y,w,h,v,size=20,color='ink',bold=False,serif=False):
  font='Shippori Mincho' if serif else 'BIZ UDPGothic'
  box=text(s,x,y,w,h,str(v),font_size=size,color=token(color),bold=bold,font_name=font)
  # Engine's CJK fallback takes precedence over Latin font_name; preserve chosen Japanese heading face explicitly.
  for el in box._element.iter():
   if el.tag.endswith('}ea'):el.set('typeface',font)
  return box
 def slide(n):
  s=prs.slides.add_slide(prs.slide_layouts[6]);rect(s,0,0,13.333,7.5,'background');rect(s,.6,.55,.08,.33,'accent-secondary')
  tx(s,.85,.5,10,.35,'ASANAGI CAFÉ  /  朝営業の試行計画',11,'accent-secondary')
  tx(s,.75,1.1,12,1.05,data['titles'][n-1],30,serif=True)
  rect(s,.85,6.92,11.6,.015,'border');tx(s,.85,7.01,10,.25,'架空の企画案  •  提案値・仮定値  •  実績未収集',13,'accent-secondary');tx(s,11.5,7.00,.9,.3,f'{n} / 5',10,'accent-secondary')
  s.notes_slide.notes_text_frame.text=data['notes'][n-1];return s
 s=slide(1)
 tx(s,.85,2.6,7,1.0,'朝営業の需要と、現場の負担を確かめる',25,serif=True)
 tx(s,.85,3.65,6.6,.8,'通常10時開店の前に試す企画案。\n通常営業と両立できるか、10日間で確認する。',20)
 for i,(v,l) in enumerate([(data['hours'],'朝の営業時間'),(f"{data['days']}日間",f"{data['weeks']}週間の平日"),(f"{data['price']}円",'セットの仮価格')]):
  x,w=[(.85,3.6),(4.4,2.1),(7.2,2.0)][i];tx(s,x,5.15,w,.65,v,25,bold=True);tx(s,x,5.85,w,.5,l,13,'accent-secondary')
 oval(s,9.0,2.75,2.7,2.7,'surface');oval(s,10.9,3.55,.75,.85,'background',line='accent-secondary');rrect(s,9.55,3.25,1.7,1.35,'accent-secondary');rect(s,9.55,3.25,1.7,.18,'data-series-1');tx(s,8.85,5.85,3.2,.45,data['cafe'],19,serif=True)
 s=slide(2)
 rect(s,.85,2.62,3.45,3.25,'surface');tx(s,1.07,2.88,3.05,.6,'商品は1種類',22,serif=True)
 tx(s,1.07,3.65,3.05,1.2,f"コーヒー＋トースト\n1セット {data['price']}円（仮）",20)
 tx(s,1.07,5.12,3.05,.5,'追加メニューは設けない',15,'accent-secondary')
 for i,(title,copy) in enumerate([('会計','注文・会計'),('調理','セットを準備'),('提供','会計担当が提供')]):
  x=4.85+i*2.5;rect(s,x,2.9,2.15,.72,'surface');tx(s,x+.12,3.02,1.9,.45,title,21,serif=True);tx(s,x,3.9,2.2,.7,copy,16)
  if i<2:tx(s,x+2.12,3.0,.4,.4,'→',19,'accent-secondary',True)
 tx(s,4.85,4.8,7.6,.5,'開始前に担当者を決める。役割は兼務も含めて確認。',17)
 tx(s,4.85,5.55,7.6,.7,'9:45 追加仕込みを判断 → 10:00 通常営業へ切替',16,'accent-secondary',True)
 s=slide(3)
 tx(s,.85,2.25,10,.5,f"セット価格 {data['price']}円 × 1日の販売数（仮定）",17,'accent-secondary')
 for i,n in enumerate(data['scenarios']):
  y=3.05+i*.94
  # Original public chart helper, one series per row. Clear Japanese labels are authored outside its narrow label boxes.
  bar_chart(s,3.0,y,[('',n/max(data['scenarios']), '')],max_width=6.0,bar_height=.42,typo=SimpleNamespace(caption=12),spacing=SimpleNamespace(bar_gap=.5))
  tx(s,.85,y-.06,2.0,.5,f'{n}セット / 日',19)
  tx(s,9.45,y-.1,3.0,.6,f'{n*data["price"]:,}円',26,bold=True)
 tx(s,.85,6.05,11.5,.55,'売上の試算です。原価・追加人件費などの費用は含まず、利益を示していません。',14,'accent-secondary')
 s=slide(4)
 tx(s,.85,2.3,11,.5,'10日間、来店・販売・待ち時間・廃棄を同じ単位で記録する。',19)
 for week in range(2):
  left=.85+week*6.0;tx(s,left,2.95,5.6,.45,f'第{week+1}週：第{week*5+1}〜{week*5+5}日',18,serif=True)
  widths=[1.1,1.1,1.1,1.2,1.1];x=left
  for w,label in zip(widths,['日付','来店数','販売数','待ち時間','廃棄数']):rect(s,x,3.5,w,.45,'surface');tx(s,x+.05,3.58,w-.1,.3,label,14);x+=w
  for row in range(5):
   y=3.98+row*.42;x=left
   for col,w in enumerate(widths):rect(s,x,y,w,.4,'background',line='border');tx(s,x+.05,y+.04,w-.1,.32,f'第{week*5+row+1}日' if col==0 else '—',15,'accent-secondary');x+=w
 tx(s,.85,6.3,11.5,.42,'— は未収集。0件ではない。待ち時間の測り方・集計単位は開始前に統一する。',14,'accent-secondary')
 s=slide(5)
 for i,(head,body) in enumerate([('需要','販売数と来店の時間帯から\n朝に開く意味があるか'),('負担','待ち時間・片づけ・仕込みから\n無理なく続けられるか'),('廃棄','余りの数と準備量から\n少量運用に調整できるか')]):
  y=2.65+i*.92;tx(s,.85,y,1.6,.55,head,24,serif=True);tx(s,2.8,y,8.9,.8,body,18);rect(s,.85,y+.78,11.6,.015,'border')
 rrect(s,.85,5.72,11.6,.86,'surface');tx(s,1.03,5.87,11.2,.6,'開始前に判断基準を合意  →  10日分を確認  →  継続・変更・終了を選ぶ',16,bold=True)
 # Japanese proofing on all runs, including chart groups. This does not replace the renderer.
 for s in prs.slides:
  for el in s._element.iter():
   if el.tag.endswith('}rPr') or el.tag.endswith('}defRPr'):el.set('lang','ja-JP')
 prs.save(out/'deck.pptx')
 dist=importlib.metadata.distribution('pptx-designer')
 files=[]
 for f in dist.files:
  p=Path(dist.locate_file(f))
  if str(f).endswith('.py') and 'pptx_designer/' in str(f).replace('\\','/'):files.append({'path':str(f),'sha256':sha(p)})
  if 'license' in str(f).lower() and p.is_file():shutil.copyfile(p,out/'pptx-designer-LICENSE.txt')
 save(out/'source/engine-runtime.json',{'package':'pptx-designer','version':dist.version,'python':sys.version,'engine_files':files,'dependencies':{n:importlib.metadata.version(n) for n in ['python-pptx','Pillow','lxml','numpy','shapely']}})
 write(out/'NOTES.md','\n\n'.join(f'## {i}. {t}\n\n{data["notes"][i-1]}' for i,t in enumerate(data['titles'],1)))
 print('Built 5 editable slides through original documented public API; no API calls.')
 write(out/'source/logs/build.log',json.dumps({'argv':sys.argv,'engine':'pptx-designer==1.0.0b11','result':'success','slide_count':5,'artifact_sha256':sha(out/'deck.pptx'),'api_calls':0},ensure_ascii=False,indent=2))
def finalize(out,data,reviewed=False):
 digest=sha(out/'deck.pptx');receipt=out/'render/powerpoint-render.json'
 if reviewed:
  if not receipt.exists() or json.loads(receipt.read_text(encoding='utf-8-sig'))['source_sha256']!=digest:raise RuntimeError('Fresh PowerPoint render required')
  vr=out/'source/visual-review.json'
  if not vr.exists() or json.loads(vr.read_text(encoding='utf-8'))['artifact_sha256']!=digest:raise RuntimeError('Current manual visual-review receipt required')
 images=[]
 for i in range(1,6):
  p=out/'render'/f'slide-{i:02}.png'
  if p.exists():images.append({'page':i,'file':p.relative_to(out).as_posix(),'sha256':sha(p)})
 save(out/'QA.json',{'status':'qa_complete' if reviewed else 'awaiting_visual_review','page_count':5,'artifact_sha256':digest,'images':images,'visual_review':{'pages_reviewed':[1,2,3,4,5] if reviewed else [],'method':'実PowerPointの全5頁PNGを個別視認し、日本語・計算・欄外・図表・講稿を確認' if reviewed else 'pending','result':'全5頁の実PowerPoint PNGを個別視認。日本語・計算・欄外・図表・講稿を確認。' if reviewed else '再生成後は全頁の視認が必要'},'data_checks':{'price':data['price'],'daily_revenue':[n*data['price'] for n in data['scenarios']],'actual_data':'not_collected'},'human_adoption':'not_requested'})
 viewer='''<!doctype html><html lang="ja"><meta charset="utf-8"><title>__TITLE__</title><style>html,body{margin:0;background:#142032;height:100%}img{display:block;width:100%;height:100%;object-fit:contain}</style><img alt="完成PowerPointの描画"><script>let n=1;function go(i){n=Math.max(1,Math.min(5,i));document.querySelector('img').src='render/slide-'+String(n).padStart(2,'0')+'.png';if(parent!==window)parent.postMessage({type:'yebisu-page',index:n-1,total:5},location.origin)}addEventListener('message',e=>{if(e.source!==parent||e.origin!==location.origin)return;if(e.data?.type==='yebisu-slide'&&Number.isInteger(e.data.index))go(e.data.index+1)});addEventListener('keydown',e=>{if(e.key==='ArrowRight')go(n+1);if(e.key==='ArrowLeft')go(n-1)});go(1)</script></html>'''
 write(out/'index.html',viewer.replace('__TITLE__',html.escape(data['title'])))
 refs=[{'path':p.relative_to(out/'source/upstream').as_posix(),'sha256':sha(p),'url':f'https://github.com/{REPO}/blob/{COMMIT}/'+p.relative_to(out/'source/upstream').as_posix()} for p in (out/'source/upstream').rglob('*') if p.is_file()]
 save(out/'PROVENANCE.json',{'repository':REPO,'upstream_commit':COMMIT,'execution_type':'original_skill_build_mode_published_public_api','generator':'tools/build_ppt_design_repository_example.py','references':refs,'input_sha256':sha(out/'input/data.json'),'artifact_sha256':digest,'output':'deck.pptx','package':'pptx-designer==1.0.0b11','source_contract':'skill/SKILL.md and references/public-api.md; init_presentation_task.py; ThemeComposer; strict Presentation; public text/shapes/bar_chart','local_changes':['Japanese content and page plan authored under original Build Mode.','Theme font atoms resolved to Shippori Mincho/BIZ UDPGothic; ja-JP proofing.','No upstream source modifications. Chart helper used unchanged.','Autonomous sample direction lock based on user request; no human adoption claimed.'],'api_calls':0,'license':'MIT original and pptx-designer package; no external photos','commands_executed':command_records(out),'library_build_receipt':{'file':'source/logs/build.log','sha256':sha(out/'source/logs/build.log')}})
 if Path(__file__).resolve()!=(out/'source/build_ppt_design_repository_example.py').resolve():shutil.copyfile(Path(__file__),out/'source/build_ppt_design_repository_example.py')
 write(out/'README.md','''# PPT-Design-Skill：地域カフェの朝営業を試す

原典の Build Mode と公開 `pptx-designer==1.0.0b11` を実行した日本語5枚です。原典 initializer → brief / page plan / theme lock → ThemeComposer → strict Presentation → 公開 text / shapes / bar_chart → 保存 → 原典 inspect の順で制作。共通テンプレートへの流し込みではありません。

架空の店舗の企画案です。販売数・価格は仮定で、売上は利益ではありません。実績は未収集。編集可能な文字・図形、図形ベースの棒グラフ、講稿をPPTXに保持しています。記録表もnative text/shapeです。

## 再生成

Python 3.10+（検証は3.13）で `pip install -r source/requirements.txt`。APIキーとPowerPointは生成に不要です。原典のソースは `source/upstream`、入力は `input/data.json`、和文フォントはBIZ UDPGothic / Shippori Minchoです。

リポジトリから: `python tools/build_ppt_design_repository_example.py --output-dir 任意の新規フォルダ --input examples/repositories/ppt-design-skill/input/data.json`。

このフォルダ単体から: `python source/build_ppt_design_repository_example.py --upstream source/upstream --input input/data.json --output-dir ../cafe-regenerated`。

入力JSONの価格・仮定販売数・見出し・講稿を編集できます。`--output-dir` は入力と異なる場所を推奨。再生成のQAは必ずawaiting_visual_review。生成成功は表示の合格を意味しません。実PowerPointによるPNG/PDFは今回同梱。再検査は原典 `source/upstream/skill/scripts/render_pptx.ps1 -InFile <deck.pptx> -OutDir <render>` でも可能です。

## テーマと根拠

terracotta / serif-editorial / accent-bar / standard / seed42。原典のThemeComposerの全resolved themeを保存し、日本語フォントのみ明示調整。写真・AI画像・外部APIなし。元リポジトリMITとエンジンMITを同梱。元リポジトリの巨大な広告画像デモは未使用・未同梱です。

PROVENANCE.jsonに固定コミット・原典参照hash・入力/出力hash、sourceに実行ログとruntime file hash、QA.jsonに表示レビュー状態があります。今回の自律制作を人間による正式採用とは扱っていません。
''')
 if False: # Catalog data is integrated separately by the parent agent.
  p='examples/repositories/ppt-design-skill/'
  row={'id':'ppt-design-skill','title':data['title'],'repository':REPO,'skill_repository':REPO,'upstream_commit':COMMIT,'fork_id':None,'use_case':'企画・店舗運営','page_count':5,'format':'pptx','preview_kind':'powerpoint-render','style':'PPT-Design-Skill / Terracotta Editorial','cover':p+'render/slide-01.png','interior':p+'render/slide-03.png','preview':p+'index.html','artifact_download':p+'deck.pptx','pdf_download':p+'render/deck.pdf','source_download':p+'source/build_ppt_design_repository_example.py','input':p+'input/BRIEF.md','input_data':p+'input/data.json','prompt':p+'PROMPT.txt','prompt_text':PROMPT,'provenance':p+'PROVENANCE.json','qa':p+'QA.json','guide':p+'README.md','titles':data['titles'],'notes':data['notes'],'pages':[p+f'render/slide-{i:02}.png'for i in range(1,6)],'input_sha256':sha(out/'input/data.json'),'method':'原典Build Mode / ThemeComposer / strict Presentation / published public helpers','sample_notice':data['notice'],'editability':'native_text_shapes'}
  save(ROOT/'data/pptdesign-repository-examples.json',[row])
  save(ROOT/'data/pptdesign-repository-guides.json',[{'repository':REPO,'execution_type':'original_skill_build_mode_published_public_api','execution_status':'generated_and_reviewed' if reviewed else 'generated_awaiting_review','case_ids':['ppt-design-skill'],'summary_ja':'原典Build Modeと公開エンジンで、架空のカフェ朝営業企画を日本語5枚に制作。テーマと根拠を保存。','requirements':['Python 3.10+ / pptx-designer1.0.0b11','BIZ UDPGothic / Shippori Mincho','PowerPointは表示QAのみ。生成はPythonだけ。'],'steps':['原典task初期化と入力/ページ計画','ThemeComposerで原典テーマを解決','strict Presentationと公開helperで生成','原典inspectとPowerPoint全頁レビュー'],'commands':['python tools/build_ppt_design_repository_example.py --output-dir cafe-regenerated'],'source_refs':refs,'limitations':['架空の企画案で、実績データなし。','ライブラリ自動デザイン生成APIは使わず、原典Build Modeで公開部品を配置。','再生成後は要レビュー。'],'guide':p+'README.md'}])
def main():
 a=argparse.ArgumentParser();a.add_argument('--output-dir',type=Path,default=DEFAULT);a.add_argument('--upstream',type=Path);a.add_argument('--input',type=Path);a.add_argument('--render',action='store_true');a.add_argument('--finalize-only',action='store_true');a.add_argument('--reviewed',action='store_true');args=a.parse_args()
 out=args.output_dir.resolve();data=json.loads(args.input.read_text(encoding='utf-8')) if args.input else DATA
 up=args.upstream or ROOT/'.build/actual-upstreams/sunchaokun__PPT-Design-Skill'
 if not up.exists():up=Path(__file__).resolve().parent/'upstream'
 if not up.exists():up=DEFAULT/'source/upstream'
 up=up.resolve()
 if not args.finalize_only:
  task=prepare(up,out,data);generate(out,task,data)
  run([sys.executable,up/'skill/scripts/inspect_pptx.py',out/'deck.pptx','--pretty'],out,'inspect.log')
  save(out/'QA.json',{'status':'awaiting_visual_review','artifact_sha256':sha(out/'deck.pptx'),'page_count':5})
  if args.render:run(['pwsh','-NoProfile','-File',ROOT/'source/render_powerpoint.ps1','-PresentationPath',out/'deck.pptx','-OutputDirectory',out/'render'],out,'powerpoint.log')
 finalize(out,data,args.reviewed)
if __name__=='__main__':main()
