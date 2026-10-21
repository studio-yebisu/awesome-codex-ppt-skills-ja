"""Inspect current native objects and actual PowerPoint PDF font resources."""
from pathlib import Path
import hashlib,json
from pptx import Presentation
from pypdf import PdfReader
CASE=Path(__file__).resolve().parents[1]
ARTIFACT=CASE/'codex-image-ppt.pptx'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def save(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def main():
 prs=Presentation(ARTIFACT);rows=[]
 for i,s in enumerate(prs.slides):
  pics=[x for x in s.shapes if x.shape_type==13]
  rows.append({'index':i,'native_shapes':len(s.shapes)-len(pics),'pictures':len(pics),'native_charts':sum(x.has_chart for x in s.shapes),'native_tables':sum(x.has_table for x in s.shapes),'texts':[x.text for x in s.shapes if x.has_text_frame and x.text],'picture_sha256':[hashlib.sha256(x.image.blob).hexdigest() for x in pics]})
 save(CASE/'source/native-elements.json',{'artifact':ARTIFACT.relative_to(CASE).as_posix(),'pptx_sha256':sha(ARTIFACT),'slides':rows})
 pdf=CASE/'render/deck.pdf';fonts=[]
 for i,page in enumerate(PdfReader(pdf).pages):
  values=[]
  for resource,ref in page['/Resources'].get_object().get('/Font',{}).items():
   f=ref.get_object();desc=f.get('/DescendantFonts');child=desc[0].get_object() if desc else f;d=child.get('/FontDescriptor');d=d.get_object() if d else {}
   values.append({'resource':str(resource),'base_font':str(f.get('/BaseFont')),'font_name':str(d.get('/FontName',f.get('/BaseFont'))),'embedded_program':any(k in d for k in ('/FontFile','/FontFile2','/FontFile3'))})
  fonts.append({'index':i,'fonts':values})
 save(CASE/'source/pdf-fonts.json',{'scope':'Actual current PowerPoint PDF resources','pdf_sha256':sha(pdf),'pages':fonts,'families':sorted({f['base_font'] for p in fonts for f in p['fonts']})})
 if (CASE/'origin_image').exists():
  images=sorted((CASE/'origin_image').glob('slide_*.png'));assert len(images)==len(rows)
  checks=[{'index':i,'source':p.relative_to(CASE).as_posix(),'source_sha256':sha(p),'embedded_sha256':rows[i]['picture_sha256'][0],'native_body_objects':rows[i]['native_shapes'],'exact_image_bytes':sha(p)==rows[i]['picture_sha256'][0]} for i,p in enumerate(images)]
  assert all(x['exact_image_bytes'] and x['native_body_objects']==0 for x in checks)
  save(CASE/'source/image-assembly-verification.json',{'artifact_sha256':sha(ARTIFACT),'checks':checks,'notes_count':sum(s.has_notes_slide for s in prs.slides)})
 print('Inspected',len(rows),'current pages:',sha(ARTIFACT))
if __name__=='__main__':main()
