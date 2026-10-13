"""Knowledge Cat's original native deck-plan builder plus narrow Japanese localization.

Without --render, generation and original editability checks run; QA remains pending.
--render additionally uses Microsoft PowerPoint. After reviewing every PNG, use
--finalize-only --reviewed. No external model/API or image generation is invoked.
Portable kits keep this script in tools/ and the case under examples/repositories/.
Pass --node and --workspace for a locally available @oai/artifact-tool workspace.
"""
from pathlib import Path
from datetime import datetime, timezone
import argparse, hashlib, importlib.metadata, json, shutil, subprocess, sys
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'examples/repositories/knowledge-cat'
UP=ROOT/'.build/actual-upstreams/gnipbao__knowledge-cat-ppt-skill'
if not UP.is_dir():UP=OUT/'source/upstream'
REPO='gnipbao/knowledge-cat-ppt-skill'
COMMIT='889c3dc00b356607fa9af935eb807056c3394886'
BUNDLE=Path('C:/Users/owner/.cache/codex-runtimes/codex-primary-runtime/dependencies/node')
REFS=['LICENSE','SKILL.md','README.md','references/engine-routing.md','references/native-pptx-recipes.md','references/story-architecture.md','references/design-systems.md','references/qa-rubric.md','assets/deck-plan.schema.json','examples/case-studies/native-editable/deck-plan.json','scripts/build_native_pptx.mjs','scripts/validate_deck_plan.py','scripts/extract_pptx_text.py','scripts/check_pptx_editability.py','scripts/probe_pptx_editability.py']
PROMPT='gnipbao/knowledge-cat-ppt-skillのnative-pptx laneと原典build_native_pptx.mjsを使い、架空の青葉企画向け「迷わない議事録のまとめ方」を日本語5枚で制作してください。決定事項・担当と期限・未決事項を分け、架空の記入例と会議末尾の確認方法を示します。入力計画を正本にし、本文・表・図形・講稿を編集可能にします。外部APIを呼ばず、原典validator・editability/probeとPowerPoint描画・目視で確かめてください。\n'

def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def rel(p):return p.relative_to(ROOT).as_posix()
def dump(p,obj):
    p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(obj,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
def run(cmd,name):
    result=subprocess.run([str(x) for x in cmd],cwd=ROOT,capture_output=True,text=True,encoding='utf8')
    log=OUT/'source'/name;log.parent.mkdir(parents=True,exist_ok=True)
    log.write_text(result.stdout+'\n'+result.stderr,encoding='utf8')
    path=OUT/'source/commands.json'
    records=json.loads(path.read_text(encoding='utf8')) if path.exists() else []
    records.append({'argv':[str(x) for x in cmd],'exit_code':result.returncode,'log':rel(log),'log_sha256':digest(log)})
    dump(path,records)
    if result.returncode:raise RuntimeError('Command failed; see '+str(log))

def localize():
    p=Presentation(OUT/'source/original-export.pptx')
    proof=next(s for s in p.slides[0].shapes if s.name=='cover-proof')
    for par,label in zip(proof.text_frame.paragraphs,['決定事項','担当と期限','未決事項']):
        par.runs[0].text=label
        for run in par.runs[1:]:run.text=''
    ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    for part in list(p.slides)+list(p.slide_layouts)+list(p.slide_masters):
        for node in part._element.findall('.//a:r',ns):
            rpr=node.find('a:rPr',ns)
            if rpr is None:rpr=OxmlElement('a:rPr');node.insert(0,rpr)
            rpr.set('lang','ja-JP')
            for tag in ('latin','ea','cs'):
                font=rpr.find('a:'+tag,ns)
                if font is None:font=OxmlElement('a:'+tag);rpr.append(font)
                font.set('typeface','Noto Sans JP')
    p.save(OUT/'deck.pptx')

def build(node,workspace):
    for f in REFS:
        target=OUT/'source/upstream'/f;target.parent.mkdir(parents=True,exist_ok=True)
        if (UP/f).resolve()!=target.resolve():shutil.copy2(UP/f,target)
    dump(OUT/'source/commands.json',[])
    dump(OUT/'QA.json',{'status':'generation_in_progress'})
    run([sys.executable,'-X','utf8',UP/'scripts/validate_deck_plan.py',OUT/'input/data.json'],'validate-plan.log')
    run([node,UP/'scripts/build_native_pptx.mjs','--plan',OUT/'input/data.json','--output',OUT/'source/original-export.pptx','--workspace',workspace,'--preview-dir',OUT/'source/artifact-previews','--inspection',OUT/'source/inspection.ndjson'],'build-native.log')
    localize()
    run([sys.executable,'-X','utf8',UP/'scripts/extract_pptx_text.py',OUT/'deck.pptx','--fail-on-placeholders','--format','json'],'extract-text.log')
    run([sys.executable,'-X','utf8',UP/'scripts/check_pptx_editability.py',OUT/'deck.pptx','--expected-slides','5','--require-native-table','--require-notes','--fail-on-image-only-slides','--json-output',OUT/'source/editability.json'],'editability.log')
    run([sys.executable,'-X','utf8',UP/'scripts/probe_pptx_editability.py',OUT/'deck.pptx','--search','記録を3つに分ける','--replacement','記録を3つに分ける・編集確認','--json-output',OUT/'source/edit-probe.json'],'edit-probe.log')
    dump(OUT/'environment.json',{'python':sys.version,'node':str(node),'artifact_workspace':str(workspace),'artifact_tool_version':json.loads((workspace/'node_modules/@oai/artifact-tool/package.json').read_text(encoding='utf8')).get('version'),'packages':{n:importlib.metadata.version(n) for n in ['python-pptx','Pillow']},'external_api_calls':0})

def finalize(rendered=False,reviewed=False):
    p=Presentation(OUT/'deck.pptx');plan=json.loads((OUT/'input/data.json').read_text(encoding='utf8'))
    assert len(p.slides)==plan['deck']['slide_count']==5
    titles=[s['action_title'] for s in plan['slides']];notes=[s['speaker_note'] for s in plan['slides']]
    (OUT/'PROMPT.txt').write_text(PROMPT,encoding='utf8')
    (OUT/'NOTES.md').write_text('\n\n'.join(f'## {i+1}\n{n}' for i,n in enumerate(notes))+'\n',encoding='utf8')
    shutil.copy2(UP/'LICENSE',OUT/'LICENSE.upstream')
    refs=[{'path':f,'sha256':digest(UP/f),'url':f'https://github.com/{REPO}/blob/{COMMIT}/{f}'} for f in REFS]
    changes=['原典build_native_pptx.mjsと全原典validatorは無変更で実行。5種類の固有レイアウトとnative table/shape/text/notesを使用。','原典cover-proofの固定技術デモ文REAL TEXT/REAL SHAPES/REAL DATAのみ、教材に合う決定事項/担当と期限/未決事項へ置換。座標と基本書式は維持。','出力PPTXの文字をNoto Sans JP/ja-JPへ局所化。原典の配置・配色・表・図形は維持。','原典plan validatorのaction-title判定は日本語5タイトルに警告。検査を抑制せずログを保持し、命題・動作を示す日本語であることを制作担当が確認。']
    dump(OUT/'source/local-changes.json',changes)
    prov={'repository':REPO,'upstream_commit':COMMIT,'references':refs,'commands_executed':json.loads((OUT/'source/commands.json').read_text(encoding='utf8')),'method':'原典deck-plan validator→build_native_pptx.mjs(@oai/artifact-tool)→再importと原典preview/inspection→出力日本語局所化→原典text/editability/probe。','execution_type':'original_native_deck_plan_builder','artifact_sha256':digest(OUT/'deck.pptx'),'input_sha256':digest(OUT/'input/data.json'),'upstream_source_modified':False,'local_changes':changes,'external_api_calls':0,'generator':'tools/build_knowledge_cat_repository_example.py','checked_at':datetime.now(timezone.utc).isoformat()}
    pages=[OUT/f'render/slide-{i:02}.png' for i in range(1,6)]
    images=[]
    if rendered:
        receipt=json.loads((OUT/'render/powerpoint-render.json').read_text(encoding='utf8'))
        assert receipt['source_sha256'].lower()==digest(OUT/'deck.pptx') and receipt['slide_count']==5
        images=[{'page':i+1,'file':rel(f),'sha256':digest(f),'dimensions':[1600,900]} for i,f in enumerate(pages)]
        prov['render_report']='render/powerpoint-render.json'
    dump(OUT/'PROVENANCE.json',prov)
    qa={'status':'qa_complete' if reviewed else ('pending_visual_review' if rendered else 'generated_awaiting_review'),'page_count':5,'artifact_sha256':digest(OUT/'deck.pptx'),'images':images,'structure':{'slides':5,'native_tables':sum(s.has_table for slide in p.slides for s in slide.shapes),'native_text_shapes':sum(s.has_text_frame and bool(s.text) for slide in p.slides for s in slide.shapes),'speaker_notes':5},'visual_review':{'pages':[1,2,3,4,5] if reviewed else [],'method':'原典再import previewの初回確認後、固定技術デモ文を修正。完成PPTXをMicrosoft PowerPoint16.0で全5枚PNG/PDFへ描画し全頁を実画像で視認。','result':'決定・担当と期限・未決の意味、native表の例、日本語文字と行間・図形・余白を確認。' if reviewed else 'pending'},'upstream_editability':json.loads((OUT/'source/editability.json').read_text(encoding='utf8')),'upstream_edit_probe':json.loads((OUT/'source/edit-probe.json').read_text(encoding='utf8')),'not_claimed':['人による正式採用','外部公開','外部情報の事実調査','LLM/画像APIによる生成']}
    dump(OUT/'QA.json',qa)
    guide='''# Knowledge Cat：迷わない議事録のまとめ方

原典のnative-pptx laneと専用 `scripts/build_native_pptx.mjs` で、架空の青葉企画向けに日本語5枚を制作。計画JSONを原典validatorで確認し、@oai/artifact-toolでnative text/shape/table/notesへ出力した。原典ビルダーは出力PPTXを再importしてpreviewとinspectionも作る。

固定原典: gnipbao/knowledge-cat-ppt-skill / 889c3dc00b356607fa9af935eb807056c3394886。MIT Copyright 2026 Knowledge Cat contributors。原文LICENSEを保持。ランタイム依存@oai/artifact-tool本体はこのキットへ複製せず、利用環境で用意する。

## 再制作

Python + python-pptx/Pillow、Node、@oai/artifact-toolを解決できるworkspaceが必要。今回の利用環境はenvironment.json。外部API・キー・画像サービスは不要。PowerPointを使う場合はNoto Sans JPも導入する。

キット内の配置は `tools/build_knowledge_cat_repository_example.py` と `examples/repositories/knowledge-cat/input/data.json`、同case内の `source/upstream/`（原典scripts、references、assets、LICENSE）。カタログの固定原典が存在しない場合は同梱source/upstreamを使う。

```text
python tools/build_knowledge_cat_repository_example.py --node <node.exe> --workspace <artifact-workspace>
```

上記だけならPPTXと原典preview・検査結果を生成し、QAはgenerated_awaiting_reviewとなる。input/data.jsonの各頁を自分の内容へ変更する。原稿は人名・日程とも架空の教材で、実在組織の規則ではない。

Windows版PowerPointとtools/render_powerpoint.ps1がある環境では `--render` を追加し、完成PPTXを描画する。全頁を目視した後だけ `--finalize-only --reviewed` で記録を確定する。既存のPNGを新しい原稿の確認証拠として流用しない。

固有の表紙・statement・table・process・closingを使い、原典ビルダー自体は無変更。表紙の固定技術デモ文と出力書体だけ日本語教材に合わせて補正した。変更範囲はsource/local-changes.json。原典plan validatorの日本語action-title警告は抑制せず保持し、内容を実読して確認した。

## 検証

原典extract_pptx_textでplaceholder検査、check_pptx_editabilityで5枚・native table・notes・image-only排除を確認。probe_pptx_editabilityは一時ファイルの文字を編集して再読込し、原本不変を確認する。構造検査は目視品質の保証ではない。PowerPoint UIでの手編集試験はしていない。

原典previewは日本語局所化前の出力、render/は局所化後のPowerPoint描画。両者の役割を分けて保持している。PPTX・原稿・講稿・原典・実行ログを同梱し、元のMIT表示を保持する。
'''
    (OUT/'README.md').write_text(guide,encoding='utf8')
    shutil.copy2(Path(__file__),OUT/'source/build_knowledge_cat_repository_example.py')
    if rendered:
        viewer='''<!doctype html><html lang="ja"><meta charset="utf-8"><title>迷わない議事録のまとめ方</title><style>html,body{margin:0;background:#12202f;height:100%}img{width:100%;height:100%;object-fit:contain}</style><img alt="PowerPointで描画した議事録ガイド"><script>let n=1;function go(i){n=Math.max(1,Math.min(5,i));document.querySelector('img').src='render/slide-'+String(n).padStart(2,'0')+'.png';parent.postMessage({type:'yebisu-page',page:n,total:5},location.origin)}addEventListener('message',e=>{if(e.source!==parent||e.origin!==location.origin)return;if(e.data.type==='yebisu-slide')go(e.data.page)});addEventListener('keydown',e=>{if(e.key==='ArrowRight')go(n+1);if(e.key==='ArrowLeft')go(n-1)});go(1)</script></html>'''
        (OUT/'index.html').write_text(viewer,encoding='utf8')
        row={'id':'knowledge-cat','title':plan['deck']['title'],'use_case':'研修・共有','repository':REPO,'skill_repository':REPO,'upstream_commit':COMMIT,'fork_id':None,'page_count':5,'format':'pptx','preview_kind':'powerpoint-render','style':'Knowledge Cat固有native builder','cover':rel(pages[0]),'interior':rel(pages[2]),'preview':rel(OUT/'index.html'),'artifact_download':rel(OUT/'deck.pptx'),'source_download':rel(OUT/'source/build_knowledge_cat_repository_example.py'),'input':rel(OUT/'input/BRIEF.md'),'input_data':rel(OUT/'input/data.json'),'prompt':rel(OUT/'PROMPT.txt'),'prompt_text':PROMPT,'provenance':rel(OUT/'PROVENANCE.json'),'qa':rel(OUT/'QA.json'),'guide':rel(OUT/'README.md'),'titles':titles,'notes':notes,'pages':[rel(f) for f in pages],'pdf_download':rel(OUT/'render/deck.pdf'),'input_sha256':digest(OUT/'input/data.json'),'method':prov['method'],'sample_notice':'架空の企業・人物・日程による研修サンプル。外部調査・実在企業の規則ではない。'}
        dump(ROOT/'data/knowledgecat-repository-examples.json',[row])
        dump(ROOT/'data/knowledgecat-repository-guides.json',[{'repository':REPO,'example_ids':['knowledge-cat'],'production_status':qa['status'],'requirements':['Python + python-pptx/Pillow','Node + @oai/artifact-tool workspace','PowerPointとNoto Sans JPは最終描画に使用'],'steps':['input/data.jsonを原典schemaへ合わせる。','原典validatorと専用native builderを実行。','固定ラベルと日本語書体を局所化し原典検査3種を実行。','PowerPoint描画を全頁目視。'],'source_refs':refs,'guide_markdown':guide,'limitations':['artifact-toolをローカルで利用できる環境が必要。','原典の日本語action-title警告は意味を別途読む必要がある。','原稿・図表を外部AIサービスが生成した実証ではない。']}])
    print(json.dumps({'status':qa['status'],'page_count':5,'artifact':rel(OUT/'deck.pptx'),'sha256':digest(OUT/'deck.pptx')},ensure_ascii=False))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--node',type=Path,default=BUNDLE/'bin/node.exe');parser.add_argument('--workspace',type=Path,default=BUNDLE);parser.add_argument('--render',action='store_true');parser.add_argument('--finalize-only',action='store_true');parser.add_argument('--reviewed',action='store_true');a=parser.parse_args()
    if a.reviewed and (not a.finalize_only or a.render):parser.error('--reviewed requires --finalize-only after actual review')
    if not a.finalize_only:build(a.node,a.workspace)
    if a.render:run(['pwsh','-NoProfile','-File',ROOT/'tools/render_powerpoint.ps1','-PresentationPath',OUT/'deck.pptx','-OutputDirectory',OUT/'render'],'powerpoint.log')
    finalize(a.render or a.finalize_only,a.reviewed)
