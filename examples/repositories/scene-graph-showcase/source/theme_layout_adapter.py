"""Bounded local repairs; upstream create/beautify remain unchanged."""
from pathlib import Path
import json

def prepare_draft(path, specs):
    from pptx import Presentation
    from pptx.util import Inches
    prs=Presentation(path);changes=[]
    for index,(slide,spec) in enumerate(zip(prs.slides,specs)):
        text=[s for s in slide.shapes if s.has_text_frame and s.text.strip()]
        edits=[]
        if spec.type=='section':
            number=next((s for s in text if s.text.strip().isdigit()),None)
            title=next((s for s in text if s is not number),None)
            edits=[(number,1.15),(title,3.75)]
        elif spec.type=='cover' and spec.subtitle:
            edits=[(next(s for s in text if s.text==spec.subtitle),4.6)]
        for shape,top in edits:
            if shape is None:continue
            before=int(shape.top);shape.top=Inches(top)
            changes.append(dict(page=index+1,text=shape.text,property='top',before=before,after=int(shape.top)))
    prs.save(path);return changes

def migrate_decor(path,old,new):
    from pptx import Presentation
    from pptx.enum.dml import MSO_FILL_TYPE
    from pptx.dml.color import RGBColor
    prs=Presentation(path);colors={str(old.primary):str(new.primary),str(old.accent):str(new.accent)};changes=[]
    for index,slide in enumerate(prs.slides):
        seen=set()
        for shape in list(slide.shapes):
            if shape.has_text_frame and shape.text.strip():continue
            try:
                if shape.fill.type!=MSO_FILL_TYPE.SOLID:continue
                before=str(shape.fill.fore_color.rgb)
            except (AttributeError,TypeError):continue
            if before in colors:
                shape.fill.fore_color.rgb=RGBColor.from_string(colors[before])
                changes.append(dict(page=index+1,shape_id=shape.shape_id,property='fill',before=before,after=colors[before]))
            if shape.left==0 and shape.top==0 and shape.height==prs.slide_height and shape.width<prs.slide_width//20:
                if shape.width!=new.left_bar_width:
                    before=int(shape.width);shape.width=new.left_bar_width
                    changes.append(dict(page=index+1,shape_id=shape.shape_id,property='width',before=before,after=int(shape.width)))
                key=(int(shape.left),int(shape.top),int(shape.width),int(shape.height),str(shape.fill.fore_color.rgb))
                if key in seen:
                    shape._element.getparent().remove(shape._element)
                    changes.append(dict(page=index+1,shape_id=shape.shape_id,property='duplicate_bar',action='removed'))
                else:seen.add(key)
    prs.save(path);return changes

def record(path,geometry,decor):
    Path(path).write_text(json.dumps(dict(author='Codex',scope='Local layout repair and old-decoration color migration around original APIs; not unmodified upstream behavior.',geometry_changes=geometry,decoration_changes=decor),ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
