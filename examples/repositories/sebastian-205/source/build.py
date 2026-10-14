from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
P=Path(__file__).resolve().parents[1]
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
INK='163C36'; GREEN='09644D'; LIME='D3EF64'; PALE='E8F0E9'; WHITE='FAFCF7'; GREY='526D63'; ORANGE='CF5A34'
def box(s,x,y,w,h,color,shape=MSO_SHAPE.RECTANGLE):
 q=s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h)); q.fill.solid(); q.fill.fore_color.rgb=RGBColor.from_string(color); q.line.fill.background(); return q
def text(s,x,y,w,h,t,size=22,color=INK,bold=False):
 q=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=q.text_frame; tf.word_wrap=True; tf.margin_left=tf.margin_right=0; tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.name='Noto Sans CJK JP'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color); p.space_after=Pt(8)
 return q
slides=[]
def page(title,n,note,dark=False):
 s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(INK if dark else WHITE)
 c=WHITE if dark else INK
 text(s,.6,.27,11,.3,'PAPER2PPT  /  RESEARCH PLAN',11,LIME if dark else GREEN,True)
 if n!=1: text(s,.6,.84,12,1.25,title,32,c,True)
 text(s,.6,7.05,11.8,.24,'出典：架空の研究ノート  |  未実施・実測結果なし  |  入力 BRIEF / data.json',10,PALE if dark else GREY)
 text(s,12.1,7.02,.6,.3,f'{n:02}',12,c,True)
 s.notes_slide.notes_text_frame.text=note
 slides.append({'title':title,'note':note}); return s
# content populated below
s=page('捨てる前の一目',1,'入力§1。架空・未実施のデザイン研究計画。廃棄原因全体ではなく、何があるか分からない情報の問題に絞る。図は入力記述に基づく説明図。',True)
text(s,.7,1.4,7,1.8,'捨てる前の\n一目。',58,WHITE,True); text(s,.75,4.05,6.6,.9,'冷蔵庫の奥を、画面の手前へ。',27,LIME,True); text(s,.75,5.15,6.6,1,'食品在庫の可視化UIを考える\nデザイン研究計画',21,WHITE)
box(s,8.2,1.4,3.9,4.9,GREEN); box(s,8.45,1.68,3.4,4.34,INK)
for y in [2.9,4.3,5.7]: box(s,8.5,y,3.3,.04,PALE)
for x,y,w in [(8.8,2.2,.7),(9.8,2.1,1.3),(9,3.55,1.15),(10.7,4.8,.6)]: box(s,x,y,w,.57,LIME)
text(s,8.7,6.3,3.5,.4,'何があるか、を見渡す。',18,WHITE)
s=page('買う前と、使う前に思い出せるか',2,'入力§2。研究の問いと三つの利用場面。毎回アプリを開く前提は未検証。見る価値が生じる場面を確かめる。')
text(s,.65,2.2,11.8,.8,'在庫の所在と状態を、どう判断につなぐ？',28,GREEN,True)
box(s,.8,4.35,11.6,.05,GREEN)
for i,(a,b,c) in enumerate([('01','買い物前','重複を確認'),('02','調理前','使う候補を選ぶ'),('03','片づけ時','残量を更新')]):
 x=.9+i*4; box(s,x,3.53,.75,.75,GREEN,MSO_SHAPE.OVAL); text(s,x+.12,3.68,.55,.4,a,22,WHITE,True); text(s,x,4.7,3.8,.5,b,27,INK,True); text(s,x,5.4,3.8,.6,c,21,GREY)
text(s,.9,6.3,11.6,.4,'確認すること：どの場面で、開く価値があるか。',17,GREEN)
s=page('入力の負担が、見える在庫を支える',3,'入力§3。仮説Aは保存場所の一覧、仮説Bは粗い残量段階。記録精度と更新負担のトレードオフであり、優劣の結果ではない。')
text(s,.7,2.2,5.1,.7,'記録の細かさ',32,GREEN,True); text(s,7.15,2.2,5.4,.7,'更新の手軽さ',32,GREEN,True)
box(s,1.4,3.32,10.5,.08,INK); box(s,6.15,3.42,1.05,1.0,INK,MSO_SHAPE.ISOSCELES_TRIANGLE)
text(s,.8,3.8,4.8,1,'場所の一覧が\n探す手がかりになるか',23); text(s,7.7,3.8,4.7,1,'大まかな残量なら\n更新を続けやすいか',23)
for i,a in enumerate(['多め','半分ほど','少なめ','不明']):
 x=.85+i*3.08; box(s,x,5.35,2.82,.72,PALE if i<3 else LIME); text(s,x+.2,5.5,2.45,.45,a,23,INK,True)
text(s,.85,6.3,11.8,.4,'設計上のトレードオフ。継続と正確さの両方を評価する。',17,GREY)
s=page('場所から探し、状態で選ぶ',4,'入力§4。提案UIは説明用の架空表示。食品名・残量・更新時点・包装表示からの期限転記を区別。期限未登録は期限がないことを意味しない。安全を保証しない。')
box(s,.85,2.03,6.05,4.65,INK); box(s,1.05,2.23,5.65,4.25,PALE)
text(s,1.3,2.5,5.1,.45,'冷蔵   /   冷凍   /   常温',22,GREEN,True)
for i,(food,qty) in enumerate([('にんじん','少なめ'),('豆腐','半分ほど'),('冷凍ごはん','不明')]):
 y=3.12+i*1.04; box(s,1.3,y,5.13,.9,WHITE); text(s,1.48,y+.08,2.75,.4,food,22,INK,True); text(s,4.3,y+.14,1.85,.36,qty,17,GREEN); text(s,1.48,y+.5,4.5,.25,'更新：前回の買い物時 / 期限：未登録',10,GREY)
text(s,7.6,2.4,5,1.2,'不明な状態を、\n隠さない。',32,GREEN,True); text(s,7.6,4.05,4.8,1.9,'残量と期限は別の情報。\n更新時点を示し、\n包装の表示と区別する。',23); text(s,7.6,6.2,4.9,.35,'説明用の架空表示',15,ORANGE,True)
s=page('「使った」を、短い操作にする',5,'入力§5。選択→残量変更→保存場所確認の提案。使い切りを取り消せる設計を検討。家族の共有は同意・権限を要し、自動的には許可されない。')
for i,(label,body) in enumerate([('選択','にんじん\n豆腐\n冷凍ごはん'),('残量変更','多め\n半分ほど\n少なめ'),('確認','保存場所：冷蔵\n更新する\n取り消す')]):
 x=.8+i*4.15; text(s,x,2.22,3.7,.4,f'0{i+1}  {label}',22,GREEN,True); box(s,x,2.92,3.45,2.9,PALE); 
 for j,line in enumerate(body.split('\n')):
  yy=3.23+j*.7; active=j==i;
  if active: box(s,x+.17,yy-.05,3.1,.59,GREEN if i==2 else LIME)
  text(s,x+.3,yy,2.8,.5,line,20,WHITE if active and i==2 else INK,active)
text(s,.85,6.2,11.7,.48,'共有する前に、誰に何を見せるかを決める。',22,GREEN,True)
s=page('見つけやすさと、更新しやすさを評価する',6,'入力§6。紙試作→操作可能な試作の計画。課題、観察、操作記録、振り返りを組み合わせる。数値・結果なし。家庭で継続評価する場合はデータ最小化と参加・共有の別の同意が必要。')
text(s,.7,2.12,11.8,.5,'紙の試作 → 操作可能な試作',24,GREEN,True)
cols=[.8,4.9,8.5]; widths=[3.85,3.45,3.95]
for x,w,h in zip(cols,widths,['課題','観察する点','聞くこと']): box(s,x,2.95,w,.62,GREEN); text(s,x+.16,3.06,w-.3,.4,h,19,WHITE,True)
rows=[['食材を見つける','見落とし','探した手がかり'],['使う候補を選ぶ','選択理由','不明情報'],['残量を更新する','操作の迷い','負担感']]
for i,row in enumerate(rows):
 y=3.65+i*.8
 for x,w,a in zip(cols,widths,row): box(s,x,y,w,.73,PALE); text(s,x+.16,y+.18,w-.3,.4,a,20)
text(s,.8,6.35,11.7,.36,'観察・操作記録・振り返りを組み合わせる。評価結果はまだない。',16,GREY)
s=page('画面の在庫と、実物はずれる',7,'入力§7。未入力・更新忘れ・家族の利用でずれる。UIだけで食品ロス削減の証明はできない。期限表示は安全保証ではなく、実物と包装表示を確認する。',True)
text(s,.75,2.25,4.8,.5,'画面の一覧',27,LIME,True); text(s,8.0,2.25,4.6,.5,'実物の棚',27,LIME,True)
for i in range(3): box(s,.8,3.15+i*.68,4.2,.45,GREEN)
for y in [3.4,4.35,5.3]: box(s,8.05,y,4.4,.06,PALE)
for x,y in [(8.5,2.95),(10.4,3.8),(9.5,4.8)]: box(s,x,y,1.1,.45,LIME)
text(s,5.25,3.85,2.6,.8,'更新の\nずれ',23,WHITE,True)
text(s,.85,5.85,11.7,.65,'可食判断を、自動化しない。',32,LIME,True)
text(s,.85,6.52,11.7,.35,'未入力・更新忘れ・共有による変化。実物と包装表示の確認が必要。',15,WHITE)
s=page('一つの棚から、始める',8,'入力§8。未実施研究計画の次の一歩。棚を決める→試作→課題→修正。通知やレシピを最初から増やさず、基本一覧と更新が判断に使えるかを検証する。')
text(s,.75,2.15,11.7,.8,'「見える」を、判断に「使える」へ。',32,GREEN,True)
for i,(a,b) in enumerate([('棚を決める','範囲を小さく'),('試作する','優先順位を形に'),('課題を試す','迷いを観察'),('表示を直す','必要なら減らす')]):
 x=.85+i*3.07; y=4.5-i*.33; box(s,x,y,2.75,.13,GREEN); text(s,x,y+.3,2.8,.5,a,24,INK,True); text(s,x,y+1,2.8,.5,b,17,GREY)
text(s,.85,6.5,11.7,.35,'未実施の研究計画  /  最初の試作で確かめる',16,ORANGE,True)
prs.save(P/'deck.pptx')
(P/'input'/'data.json').write_text(json.dumps({'title':'捨てる前の一目','source_sufficiency':'notes-only','status':'fictional-unperformed-research-plan','slides':slides},ensure_ascii=False,indent=2))
