"""Audit actual edit operations, native objects, and donor-part preservation.

XML equality ignores declaration/attribute serialization and whitespace-only
formatting. Only package Relationships children are compared as an unordered
map keyed by Id: their XML order is not a relationship identity.
"""
from pathlib import Path
from collections import Counter
import hashlib
import json
import shutil
import zipfile
from lxml import etree
from pptx import Presentation
from pypdf import PdfReader

CASE = Path(__file__).resolve().parents[1]
REL_NS = 'http://schemas.openxmlformats.org/package/2006/relationships'

def sha(data): return hashlib.sha256(data).hexdigest()
def save(path, value): path.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')

def xml_content(data):
    root = etree.fromstring(data)
    def walk(node):
        children = [walk(c) for c in node]
        if node.tag == '{'+REL_NS+'}Relationships':
            children.sort(key=lambda child: dict(child[1]).get('Id', ''))
        text=node.text or ''; tail=node.tail or ''
        return (node.tag, tuple(sorted(node.attrib.items())), text if text.strip() else '', tuple(children), tail if tail.strip() else '')
    return walk(root)

def main():
    donor, result = CASE/'input/donor.pptx', CASE/'deck.pptx'
    patch = json.loads((CASE/'review/replay/patch.json').read_text('utf-8'))
    shutil.copy2(CASE/'review/replay/patch.json', CASE/'source/effective-patch.json')
    parts=[]
    with zipfile.ZipFile(donor) as before, zipfile.ZipFile(result) as after:
        for name in sorted(before.namelist()):
            if name.endswith('/') or not name.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/theme/', 'ppt/media/')): continue
            a=before.read(name); present=name in after.namelist(); b=after.read(name) if present else b''
            xml=name.endswith(('.xml','.rels'))
            parts.append({'path':name,'present':present,'before_sha256':sha(a),'after_sha256':sha(b) if present else None,'byte_identical':present and a==b,'content_equal':present and (xml_content(a)==xml_content(b) if xml else a==b),'comparison':'parsed XML; Relationships children keyed by Id' if xml else 'exact bytes'})
    preservation={'donor_sha256':sha(donor.read_bytes()),'result_sha256':sha(result.read_bytes()),'patch':'source/effective-patch.json','actual_operations':len(patch['ops']),'operations_by_type':dict(Counter(x['op'] for x in patch['ops'])),'parts':parts,'all_master_layout_theme_media_parts_byte_identical':all(x['byte_identical'] for x in parts),'all_master_layout_theme_media_parts_content_equal':all(x['content_equal'] for x in parts),'part_count':len(parts),'byte_identical_count':sum(x['byte_identical'] for x in parts),'content_equal_count':sum(x['content_equal'] for x in parts),'normalization':'XML declaration, attribute serialization, formatting-only surrounding whitespace; Relationships child order keyed by unchanged Id. All attribute values and element text remain compared. No geometry, font, or target value is normalized.','diff':'review/replay/diff.log'}
    save(CASE/'source/edit-preservation.json',preservation)
    prs=Presentation(result); slides=[]
    for i,slide in enumerate(prs.slides):
        pictures=sum(s.shape_type==13 for s in slide.shapes)
        slides.append({'index':i,'native_shapes':len(slide.shapes)-pictures,'pictures':pictures,'native_charts':sum(s.has_chart for s in slide.shapes),'native_tables':sum(s.has_table for s in slide.shapes),'texts':[s.text for s in slide.shapes if s.has_text_frame and s.text]})
    save(CASE/'source/native-elements.json',{'artifact':'deck.pptx','pptx_sha256':sha(result.read_bytes()),'slides':slides})
    pdf=CASE/'render/deck.pdf'; fonts=[]
    for i,page in enumerate(PdfReader(pdf).pages):
        values=[]
        for resource,ref in page['/Resources'].get_object().get('/Font',{}).items():
            f=ref.get_object(); descendants=f.get('/DescendantFonts'); child=descendants[0].get_object() if descendants else f
            d=child.get('/FontDescriptor'); d=d.get_object() if d else {}
            values.append({'resource':str(resource),'base_font':str(f.get('/BaseFont')),'font_name':str(d.get('/FontName',f.get('/BaseFont'))),'embedded_program':any(k in d for k in ('/FontFile','/FontFile2','/FontFile3'))})
        fonts.append({'index':i,'fonts':values})
    save(CASE/'source/pdf-fonts.json',{'scope':'Actual current PowerPoint-exported PDF resources, not PPTX family attributes alone','pdf_sha256':sha(pdf.read_bytes()),'pages':fonts,'families':sorted({f['base_font'] for p in fonts for f in p['fonts']})})
    print(json.dumps({k:v for k,v in preservation.items() if k not in ('parts','normalization')},ensure_ascii=False,indent=2))

if __name__=='__main__': main()
