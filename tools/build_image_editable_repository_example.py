"""Package and replay the recorded original image-to-editable-ppt workflow."""
from pathlib import Path
import argparse, hashlib, json, os, shutil, subprocess, sys
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'examples/repositories/image-to-editable'
UP=ROOT/'.build/actual-upstreams/ningzimu__image-to-editable-ppt-skill'
REPO='ningzimu/image-to-editable-ppt-skill'
COMMIT='b7be494e31a0ed56ef98716891db5474606b8cdf'
PROMPT='既存の日本語Git入門画像2枚を、原典image-to-editable-pptのCLIとページworkerで編集可能なPowerPointへ再構成。文字、カード、単純な矢印はnative要素。アイコンはCodex内蔵image_genによる元画像の編集で分離し、画像として配置。第三者API、購入、外部OCRは使わない。元画像と完成PowerPointの全ページを視認し、編集できる部分と画像部分を明示する。'
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def write(p,s):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(s,encoding='utf-8')
def save(p,o):write(p,json.dumps(o,ensure_ascii=False,indent=2)+'\n')
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def call(argv,out,name):
 env=os.environ.copy()
 for k in ['OPENAI_API_KEY','OPENAI_BASE_URL','PADDLE_OCR_TOKEN']:env.pop(k,None)
 cfg=out/'source/offline-config';env.update(EDITPPT_CONFIG_HOME=str(cfg),CODEX_AUTH_FILE=str(cfg/'no-auth.json'),CODEX_PPT_HOME=str(cfg/'no-codex-ppt'),PYTHONIOENCODING='utf-8')
 # Python on Windows can omit a script directory beyond MAX_PATH from sys.path[0].
 # The original CLI spawns sibling runtime scripts; give those unmodified scripts
 # their own module directory explicitly, including deeply nested ZIP extractions.
 env['PYTHONPATH']=str((out/'source/upstream/skills/image-to-editable-ppt/cli/editppt/runtime').resolve())
 p=subprocess.run([str(x)for x in argv],capture_output=True,text=True,encoding='utf-8',errors='replace',env=env)
 log=out/'source/logs'/name;write(log,json.dumps([str(x)for x in argv],ensure_ascii=False)+'\n'+p.stdout+'\n'+p.stderr)
 receipt={'argv':[str(x)for x in argv],'exit_code':p.returncode,'log':log.relative_to(out).as_posix(),'log_sha256':sha(log)}
 save(log.with_suffix('.receipt.json'),receipt)
 if p.returncode:raise RuntimeError(p.stdout+'\n'+p.stderr)
 return receipt

NOTES=[
 'このページでは、エディターでの保存、Gitへの記録、リモートへの共有を区別します。保存だけでは履歴に残らず、コミットだけでは共有先へ送られません。この作例は元画像から文字とカードを編集可能に再構成し、アイコンは画像として保持しています。この案内ノートはカタログ制作時に新規追加したもので、元画像から抽出したものではありません。',
 '変更は作業ツリー、ステージ、ローカル履歴の順に進みます。git addで次に記録する内容を選び、git commitでその内容を手元の履歴に残します。add後の追加編集は自動でステージへ入らないため、必要なら再度addして確認します。この案内ノートはカタログ制作時に新規追加したものです。'
]
def prepare_input(out):
 p=out/'input/data.json'
 if not p.exists():save(p,{'title':'画像から編集できるGit入門へ','notes_origin':'New Japanese catalog guidance written after reconstruction; input PNG images contain no speaker notes.','pages':[{'page':i,'source':f'input/slide_{i:02}.png','sha256':sha(out/f'input/slide_{i:02}.png'),'title':title,'note':NOTES[i-1]}for i,title in [(1,'Gitの保存・記録・共有'),(2,'変更は3つの場所を進む')]]})
 return load(p)
def add_notes(out,source):
 from pptx import Presentation
 from zipfile import ZipFile
 data=prepare_input(out);dest=out/'deck.pptx';receipt=out/'source/notes-addition.json'
 if receipt.exists()and dest.exists():
  old=load(receipt)
  if old.get('source_sha256')==sha(source)and old.get('input_sha256')==sha(out/'input/data.json')and old.get('artifact_sha256')==sha(dest):return
 def stable_parts(path):
  with ZipFile(path)as z:return {n:sha_bytes(z.read(n))for n in z.namelist()if (n.startswith('ppt/slides/slide')and n.endswith('.xml'))or n.startswith('ppt/media/')}
 before=stable_parts(source);prs=Presentation(source)
 for slide,page in zip(prs.slides,data['pages']):slide.notes_slide.notes_text_frame.text=page['note']
 prs.save(dest)
 # Preserve original slide bodies and media byte-for-byte; notes use package relationships.
 temp=dest.with_suffix('.notes-tmp.pptx')
 with ZipFile(source)as original,ZipFile(dest)as annotated,ZipFile(temp,'w')as merged:
  for item in annotated.infolist():merged.writestr(item,original.read(item.filename)if item.filename in before else annotated.read(item.filename))
 temp.replace(dest)
 assert before==stable_parts(dest),'Notes addition changed slide XML or embedded media'
 check=Presentation(dest);assert [x.notes_slide.notes_text_frame.text for x in check.slides]==[x['note']for x in data['pages']]
 save(receipt,{'source_sha256':sha(source),'input_sha256':sha(out/'input/data.json'),'artifact_sha256':sha(dest),'operation':'Add newly authored Japanese guidance to speaker notes using python-pptx; original PNGs contain no notes.','slide_xml_and_media_unchanged':True,'verified_parts':len(before),'notes':[x['note']for x in data['pages']]})
def sha_bytes(b):return hashlib.sha256(b).hexdigest()

def resolve_failed_receipts(receipts):
 resolutions={
  'source/logs/add-notes-and-package.log':('source/logs/add-notes-and-package-repaired.log','Same argv after repairing the notes-only transform: python-pptx rewrote original slide XML during its first roundtrip. The repaired implementation retains original slide XML and media bytes while adding note package parts and relationships, then asserts unchanged body/media hashes.'),
  'run/pages/page_001/process-sheet.log':('run/pages/page_001/process-sheet-repair.log','Corrected the split region boundary from assumed width 2048 to the actual generated sheet width 2046. The successful command replaces --asset-sheet-source assets/icon-sheet.png with --skip-chroma to reuse the already created alpha sheet and rerun only the failed split stage.')
 }
 by_log={x['log']:x for x in receipts}
 for failed,(success,reason)in resolutions.items():
  if failed not in by_log:continue
  if by_log[failed]['exit_code']==0:raise RuntimeError('Expected preserved failure receipt: '+failed)
  if success not in by_log or by_log[success]['exit_code']!=0:raise RuntimeError('Missing successful resolution receipt: '+success)
  by_log[failed]['resolved_by_log']=success;by_log[failed]['resolution_reason']=reason
 return receipts

def freeze(out,up):
 dest=out/'source/upstream';dest.mkdir(parents=True,exist_ok=True)
 for name in ['LICENSE','AGENTS.md','README.md','README_en.md']:
  if(up/name).exists():shutil.copy2(up/name,dest/name)
 shutil.copytree(up/'skills/image-to-editable-ppt',dest/'skills/image-to-editable-ppt',dirs_exist_ok=True,ignore=shutil.ignore_patterns('__pycache__','*.pyc','*.egg-info'))
 shutil.copy2(up/'LICENSE',out/'LICENSE.upstream')
 if(up/'ACQUISITION.json').exists():shutil.copy2(up/'ACQUISITION.json',out/'source/upstream-acquisition.json')
 if Path(__file__).resolve()!=(out/'source/build_image_editable_repository_example.py').resolve():shutil.copy2(__file__,out/'source/build_image_editable_repository_example.py')
 write(out/'PROMPT.txt',PROMPT+'\n')
 write(out/'source/requirements.txt','-e ./upstream/skills/image-to-editable-ppt/cli\npython-pptx==1.0.2\n')

def assemble(source,out,up):
 if source.resolve()==out.resolve():raise ValueError('Replay requires a separate output directory')
 out.mkdir(parents=True,exist_ok=False)
 shutil.copytree(source/'run',out/'run')
 shutil.copytree(source/'input',out/'input')
 freeze(out,up)
 deck=load(out/'run/deck_manifest.json');deck['job_dir']=str((out/'run').resolve());save(out/'run/deck_manifest.json',deck)
 cli=out/'source/upstream/skills/image-to-editable-ppt/cli/editppt/cli.py'
 call([sys.executable,cli,'run','finalize',out/'run'],out,'replay-finalize.log')
 add_notes(out,out/'run'/deck['output'])
 save(out/'QA.json',{'status':'awaiting_visual_review','page_count':2,'artifact_sha256':sha(out/'deck.pptx'),'visual_review':{'pages_reviewed':[],'method':'pending','result':'Replayed recorded page manifests/assets; fresh PowerPoint review required.'}})
 print(out/'deck.pptx')

def finalize(out,up,reviewed):
 freeze(out,up)
 deck=load(out/'run/deck_manifest.json');src=out/'run'/deck['output']
 if not src.exists():raise RuntimeError('Original run finalize must succeed first')
 add_notes(out,src)
 data=prepare_input(out)
 validation=load(src.parent/'validation.json')
 render=load(out/'render/powerpoint-render.json') if(out/'render/powerpoint-render.json').exists()else None
 review=load(out/'source/visual-review.json')if(out/'source/visual-review.json').exists()else None
 if reviewed and(not render or render['source_sha256']!=sha(out/'deck.pptx')or not review or review['artifact_sha256']!=sha(out/'deck.pptx')or review['pages_reviewed']!=[1,2]):raise RuntimeError('Current all-page PowerPoint visual review required')
 refs=[{'path':p.relative_to(out/'source/upstream').as_posix(),'sha256':sha(p),'url':f'https://github.com/{REPO}/blob/{COMMIT}/'+p.relative_to(out/'source/upstream').as_posix()}for p in sorted((out/'source/upstream').rglob('*'))if p.is_file()]
 receipts=[load(p)for p in sorted((out/'source/logs').glob('*.receipt.json'))]
 for receiptfile in sorted((out/'run/pages').glob('*/cli-receipts.json')):
  for item in load(receiptfile):
   item['log']=(receiptfile.parent.relative_to(out)/item['log']).as_posix();receipts.append(item)
 receipts=resolve_failed_receipts(receipts)
 save(out/'PROVENANCE.json',{'repository':REPO,'upstream_commit':COMMIT,'execution_type':'original_editppt_cli_and_page_workers_builtin_image_edit','references':refs,'commands_executed':receipts,'input_sha256':sha(out/'input/data.json'),'notes_addition':load(out/'source/notes-addition.json'),'input_provenance':'source/original-PROVENANCE.json','inputs':[{'path':p.relative_to(out).as_posix(),'sha256':sha(p)}for p in sorted((out/'input').glob('*.png'))],'output':'deck.pptx','artifact_sha256':sha(out/'deck.pptx'),'image_backend':'Codex built-in image_gen; model undisclosed','external_ocr':False,'third_party_api':False,'license':'Original MIT; Python dependencies not redistributed as binaries','changes':'Japanese page manifests and built-in edited asset sheets; upstream CLI unmodified','image_jobs':['run/pages/page_001/imagegen-jobs.json','run/pages/page_002/imagegen-jobs.json']})
 save(out/'QA.json',{'status':'qa_complete'if reviewed else'awaiting_visual_review','page_count':2,'artifact_sha256':sha(out/'deck.pptx'),'upstream_validation':validation,'visual_review':review if reviewed else{'pages_reviewed':[],'method':'pending','result':'pending'},'editability':'Native Japanese text, cards, simple arrows; separated icons remain bitmap images','human_adoption':'pending'})
 write(out/'README.md','''# 画像から編集できるGit入門へ

元画像2枚を原典 `image-to-editable-ppt` のprepare → page worker → build/validate → record → finalizeで再構成した作例です。新しい日本語原稿をゼロから生成する作例ではありません。

`deck.pptx`の文字・カード・単純な矢印はPowerPointのnative要素です。アイコンは内蔵image_genで元画像から分離した画像で、内部の線や色は個別編集できません。元画像はinput、ページmanifest・編集結果・検証はrun/pages、完成PowerPointの描画はrender。画像入力に発表者ノートは含まれません。カタログ制作時に新規作成した日本語案内ノート2件を、原典CLIの出力後にpython-pptxで追加しています。input/data.jsonとカタログの発表メモとPPTXノートは同じ文章です。元画像からの抽出ノートではありません。source/notes-addition.jsonに変更前後のSHAと本文XML・画像の不変を記録し、replayでも同じ処理を実行します。前作の説明ノートはsource/original-NOTES.mdを参考資料として保持しています。

## 再構成と再出力

新しい画像では、同梱source/upstream/skills/image-to-editable-ppt/SKILL.mdを読ませ、PROMPT.txtと入力を渡してください。画像編集を使うAIエージェントとページworkerが必要です。APIキーやOAuthの読取りはこの作例では行っていません。外部OCRは使わず、ローカルtext hintsと画像の視認を使用しました。

保存済みmanifestと分離済み画像からの再出力はPython 3.10+と原典CLI依存の導入後、次で行えます。新しい画像編集やworkerを再実行するコマンドではありません。

```powershell
python -m pip install -e source/upstream/skills/image-to-editable-ppt/cli python-pptx==1.0.2
python source/build_image_editable_repository_example.py --replay --source-dir . --output-dir ../image-editable-replayed --upstream-dir source/upstream
```

既存出力先は拒否します。再出力後は表示レビューが必要です。原典CLIのPyMuPDFなど依存パッケージは別途導入し、そのライセンスに従ってください。原典ソースはMIT、LICENSE.upstreamを保持。第三者の公式作例や、人間による採用済みという意味ではありません。
''')
 write(out/'index.html','''<!doctype html><html lang="ja"><meta charset="utf-8"><title>画像から編集可能なGit入門へ</title><style>html,body{margin:0;height:100%;background:#eee}img{width:100%;height:100%;object-fit:contain}</style><img alt="完成PowerPointの表示"><script>let n=1;function go(i){n=Math.max(1,Math.min(2,i));document.querySelector('img').src='render/slide-'+String(n).padStart(2,'0')+'.png';if(parent!==window)parent.postMessage({type:'yebisu-page',index:n-1,total:2},location.origin)}addEventListener('message',e=>{if(e.source!==parent||e.origin!==location.origin)return;if(e.data?.type==='yebisu-slide'&&Number.isInteger(e.data.index))go(e.data.index+1)});addEventListener('keydown',e=>{if(e.key==='ArrowRight')go(n+1);if(e.key==='ArrowLeft')go(n-1)});go(1)</script></html>''')
 write(out/'comparison.html','''<!doctype html><html lang="ja"><meta charset="utf-8"><title>変換前後の比較</title><style>body{margin:0;padding:24px;font-family:Yu Gothic,sans-serif;color:#123259;background:#f5f8fc}h1{font-size:24px}section{display:grid;grid-template-columns:1fr 1fr;gap:16px;margin:28px 0}figure{margin:0}img{width:100%;display:block;border:1px solid #d8e1eb}figcaption{margin:8px 0}@media(max-width:700px){section{grid-template-columns:1fr}}</style><h1>画像から編集できるGit入門へ</h1><p>元画像と完成PowerPointの描画を比較できます。文字・カード・矢印は編集可能、アイコンは画像です。</p><p><a href="deck.pptx">PowerPointを開く</a> · <a href="README.md">制作と再出力の説明</a></p>'''+''.join(f'<section><figure><figcaption>{i}ページ：元画像</figcaption><img src="input/slide_{i:02}.png" alt="元画像{i}"></figure><figure><figcaption>{i}ページ：PowerPoint</figcaption><img src="render/slide-{i:02}.png" alt="再構成{i}"></figure></section>'for i in [1,2])+'</html>')
 if out.resolve()==OUT.resolve():
  p='examples/repositories/image-to-editable/'
  save(ROOT/'data/imageeditable-repository-examples.json',[{'id':'image-to-editable','title':'画像から編集できるGit入門へ','use_case':'画像スライドの編集可能化','repository':REPO,'skill_repository':REPO,'upstream_commit':COMMIT,'fork_id':'image-to-editable-ppt-ja','page_count':2,'format':'pptx','preview_kind':'powerpoint-render','style':'Git入門 / Native text and cards','pdf_download':p+'render/deck.pdf','input_data':p+'input/data.json','input_sha256':sha(out/'input/data.json'),'notes':[x['note']for x in data['pages']],'cover':p+'render/slide-01.png','interior':p+'render/slide-02.png','preview':p+'index.html','artifact_download':p+'deck.pptx','source_download':'downloads/image-to-editable-example-ja.zip','input':p+'input/slide_01.png','prompt':p+'PROMPT.txt','prompt_text':PROMPT,'provenance':p+'PROVENANCE.json','qa':p+'QA.json','guide':p+'README.md','pages':[p+f'render/slide-{i:02}.png'for i in [1,2]],'method':'Original editppt CLI + page workers + built-in image editing','editability':'native_text_shapes_and_bitmap_icons','sample_notice':'既存の画像2枚から再構成。文字・カードは編集可能、アイコンは画像。','titles':['保存・履歴・共有を分ける','変更は3つの場所を通る'],'comparison':p+'comparison.html','before_pages':[p+f'input/slide_{i:02}.png' for i in [1,2]]}])
  save(ROOT/'data/imageeditable-repository-guides.json',[{'repository':REPO,'execution_type':'original_editppt_cli_and_page_workers_builtin_image_edit','execution_status':'generated_and_reviewed'if reviewed else'generated_awaiting_review','case_ids':['image-to-editable'],'summary_ja':'日本語Git入門画像2枚を編集可能な文字・カードと画像アイコンに分離。','requirements':['Python3.10+ / 原典editppt CLI','新規入力には内蔵画像編集とpage worker対応エージェント','PowerPoint（今回の表示QA）'],'steps':['prepareとpage prompt','workerでforeground分離とnative再構成','page build/validateとrun record','run finalizeとPowerPoint全頁視認'],'commands':['python source/build_image_editable_repository_example.py --replay --source-dir . --output-dir ../replayed --upstream-dir source/upstream'],'source_refs':refs,'limitations':['アイコンはbitmapで個別の線を編集できない','OCR外部APIは未使用','元画像の画素を完全一致させる変換ではない'],'guide':p+'README.md'}])

def main():
 a=argparse.ArgumentParser();a.add_argument('--output-dir',type=Path,default=OUT);a.add_argument('--upstream-dir',type=Path,default=UP);a.add_argument('--source-dir',type=Path,default=OUT);a.add_argument('--replay',action='store_true');a.add_argument('--reviewed',action='store_true');v=a.parse_args()
 if v.replay:assemble(v.source_dir.resolve(),v.output_dir.resolve(),v.upstream_dir.resolve())
 else:finalize(v.output_dir.resolve(),v.upstream_dir.resolve(),v.reviewed)
if __name__=='__main__':main()
