"""Synchronize actual original-runtime reconstruction evidence and safe distributions."""
from pathlib import Path
import hashlib,json,shutil,subprocess,sys,zipfile
import pymupdf
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];CASE=ROOT/'examples/repositories/slidepoise-showcase';FORK=ROOT/'skill-forks/slidepoise-ja'
IDS=['reduce-steps','branch-and-information','targets-and-order']
load=lambda p:json.loads(p.read_text(encoding='utf-8-sig'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def dump(p,v):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
SKILL='''---
name: slidepoise-ja
description: 原稿の意図を画像で設計し、計測と意味の割当を経て編集可能なPowerPointへ再構築する。
---

# 画像で設計し、文字・図形を編集できるPowerPointへ

原典 henryhyw/slidepoise / 481a028486b528e65706c4e14aa9b74a5e340f12 / MIT。非公式の日本語案内。upstream/slidepoise/SKILL.mdと必要referencesを読む。

1. 原典CLIのdoctor・profile showで環境と選択Profile/Library Setsを確認し、resolve_config.pyで今回の上書きを解決する。目的・読み手・根拠・必要内容を確定し、work/deck-outline.jsonに安定した頁IDと順序・役割を書く。
2. 各頁の意図に合う参照候補を検索し、実物を見て選ぶ。Profileと共通frameを決める。画像の生成領域はframeを除く本文領域で、常に16:9とは限らない。
3. intentと資源からprepare_generationでrequestをcompileし、その正確なpromptと参照を画像ツールへ渡す。手動生成を選んだ場合は同じrequestを書き出して画像を取り込む。採用画像と原稿を実際に見比べる。
4. semantic mapと独立したcontent obligationsを記す。文字領域、図形、接続元/先、画像の役割はAgentが判断する。原典reconstruct-slideでOpenCV計測→overlay確認→contract→constructor sceneへ進む。コードは美しさや意味を判定しない。
5. 原典render-deckで頁順に組み立てる。全頁の実描画を開き、内容の欠落、文字のかたまり、光学的位置、接続方向、頁間の統一を生成画像と同じ縮尺で比べる。複数頁では独立した最終レビューを行う。問題を元の意図・mapへ戻して直し、再描画する。
6. 編集可能PPTX、プレビュー、使用素材・字体・差分と未解決事項を渡す。汎用PptxGenJS/python-pptx/HTML/SVGの独自builderを代替にしない。画像背景だけを編集可能と称さない。

## この作例で試したこと

仮想予約サイトの3頁。原典compile済みrequestで生成した画像→normalizer→host semantic map→OpenCV→contract→constructor→原典rendererを全頁で実行。native要素は34/28/21個。数値は目標で実測ではない。現在の全3頁をPowerPointで描画して見比べ、1頁の切り欠きと文字の接触、3頁の数字と単位の離れを修復した。数字と単位は同じnative文字領域にまとめ、元のcontent obligationsを保持した。

原画像と比べやすい平面の構成を保持しているが、書体のinkや紙の陰影の完全一致は保証しない。字体解決と文字領域の補正はsource/text-layout-repair.jsonに記録。原画像の再生成はしていない。独立した最終再評価、外部Library検索・Console/global Profile導入、SAM、実在ロゴ/写真の保持は未実施。

## 再現

Python3.10+、OpenCV/numpy/Pillow/python-pptx、Nodeとpptxgenjs4.0.1/jszip3.10.1。作り直す場合は画像生成ツールまたは手動取り込み。Windows PowerPoint描画にはMicrosoft PowerPointと同梱Dela Gothic One・BIZ UDPGothic・Manropeの導入が必要。各OFLを保持する。

作例ZIPでは `python rebuild.py` が保存画像から原典再構築を再実行する。例の内容とsemantic mapを再現する処理で、任意の題材を自動企画するアプリではない。`export_case.ps1`は自分が開いたdeckだけ閉じる。全頁の実表示と比較を行う。
'''
def main():
 d=pymupdf.open(CASE/'.private/repair-render/deck.pdf');assert len(d)==3
 fonts=sorted({f[3].split('+')[-1] for pg in d for f in pg.get_fonts(full=True)})
 assert any('DelaGothicOne' in f for f in fonts) and any('BIZ' in f for f in fonts)
 shutil.copy2(CASE/'.private/repair-render/deck.pdf',CASE/'render/deck.pdf')
 dump(CASE/'render/powerpoint-render.json',load(CASE/'.private/repair-render/powerpoint-render.json'))
 for n,pg in enumerate(d,1):pg.get_pixmap(matrix=pymupdf.Matrix(1600/pg.rect.width,1600/pg.rect.width),alpha=False).save(CASE/f'render/slide-{n:02}.png')
 prs=Presentation(CASE/'deck.pptx');stats=[{'page':n,'shapes':len(s.shapes),'pictures':sum(q.shape_type==13 for q in s.shapes),'text_shapes':sum(q.has_text_frame and bool(q.text.strip()) for q in s.shapes)} for n,s in enumerate(prs.slides,1)];assert all(s['pictures']==0 for s in stats)
 comparisons=[]
 for n,sid in enumerate(IDS,1):
  out=CASE/f'qa/comparison-{n:02}.png'
  subprocess.run([sys.executable,str(CASE/'upstream/slidepoise/scripts/make_visual_comparison.py'),'--left',str(CASE/f'slides/{sid}/work/accepted-slide.png'),'--right',str(CASE/f'render/slide-{n:02}.png'),'--config',str(CASE/'work/resolved-config.json'),'--left-label','Generated target','--right-label','Native PowerPoint render','--output',str(out)],check=True,capture_output=True)
  comparisons.append({'page':n,'comparison':str(out.relative_to(CASE)),'generated_sha256':sha(CASE/f'slides/{sid}/work/accepted-slide.png'),'render_sha256':sha(CASE/f'render/slide-{n:02}.png')})
 dump(CASE/'source/current-native-proof.json',{'author':'Codex','pptx_sha256':sha(CASE/'deck.pptx'),'actual_pdf_fonts':fonts,'native_inventory':stats,'comparisons':comparisons,'observations':['p1: all eight old-stage labels stay inside their chevrons, including the four-character registration label.','p2: vacancy branches retain distinct blue/orange paths and arrow directions; native body and target differ in font ink.','p3: number and unit share each native text allocation; gaps and thin residual unit marks removed. All three target values and 24 content obligations retained.'],'independent_review':'pending'})
 qa=load(CASE/'QA.json');qa.update(status='original_reconstruction_all3_visual_checked',all_pages_detailed_visual_review=True,structural_audit=True,independent_review=False,checks=['Original reconstruction/runtime replay: source/text-layout-repair.json','All3 current PowerPoint pages and generated targets viewed; original comparisons in qa/','Actual PDF fonts and native element inventory: source/current-native-proof.json'],limits=['Generated text ink and paper shading differ from native objects.','Targets and venue details are fictional assumptions, not measured results.','Independent final recheck and global Profile/Library retrieval are not completed.']);
 if (CASE/'source/zip-replay.json').exists():qa.update(zip_reexecution=True,distribution_replay='source/zip-replay.json')
 dump(CASE/'QA.json',qa)
 prov=load(CASE/'PROVENANCE.json');prov.setdefault('initial_check_history',{'not_exercised':prov.get('not_exercised'),'visual_limitations':prov.get('visual_limitations')});prov['not_exercised']=['External Library reference retrieval, Console/global profile installation, SAM, image editing, authentic third-party photographs/logos, independent final recheck.'];prov['visual_limitations']=['Native ink and flat fills differ from generated fonts and paper shading; pixel identity is not claimed.'];prov['constructor_objects']=[34,28,21];prov['current_repair']={'author':'Codex','proof':'source/current-native-proof.json','changes':'Native metric/unit grouping and inner chevron label padding, using original runtime; original images retained.'};dump(CASE/'PROVENANCE.json',prov)
 prompt=(CASE/'PROMPT.md').read_text(encoding='utf-8');prompt=prompt.replace('全詳細目視/構造監査/独立review/ZIP再実行を追加しない。','全頁を原画像と実描画で比較し、構造・配布再実行を確認して修復する。独立レビューの結果は別記録。');(CASE/'PROMPT.md').write_text(prompt,encoding='utf-8')
 for folder in [CASE,FORK]:(folder/'SKILL.md').write_text(SKILL,encoding='utf-8')
 (CASE/'README.md').write_text('# 画像から編集可能な図解へ\n\n原典SlidePoiseの画像生成→意味の割当→OpenCV計測→contract→constructor→native PPTXの作例です。仮想予約サイトの3頁、数値は目標です。\n\n現在の全3頁を原画像・計測overlay・PowerPoint描画で比較し、文字の接触と数値の単位間隔を修復しました。source/current-native-proof.jsonとqa/に実物の証拠があります。画像の陰影や字形の完全一致は保証しません。独立した最終再評価は未了。\n\n環境と一般の使い方はSKILL.md。保存画像からの再構築はpython rebuild.py。元の生成request、semantic map、原典runtime、MIT、原字体/OFLを保持しています。\n',encoding='utf-8')
 patch=next(r for r in load(ROOT/'data/skill-effect-patches/native-03.json') if r['case_id']=='slidepoise-showcase');patch=json.loads(json.dumps(patch));patch['constraints']=[*prov['visual_limitations'],'複数頁で原典が求める独立した最終再評価は未了。','外部Library検索・Console/global Profile導入、SAM、画像編集は未実施。','数値は設計目標、会場名と120分は仮想設定。実績ではない。'];demo=patch['demonstration'];demo['summary']='原典画像→OpenCV計測→contract→constructor→native rendererを3頁で実行。全3頁の生成画像・実描画を比較し、文字の接触と単位間隔を修復。';demo['not_demonstrated']=prov['not_exercised'];
 descriptions=['8操作と4操作の各段階は編集可能なfreeform。内側の文字領域を補正し、会員登録が切り欠き線に接しない。','空席の有無に対応する接続元・先をAgentが割り当て、OpenCVの位置証拠からnative接続を作る。','4操作・90%・60秒を、それぞれ一つのnative文字領域で保持。原画像と現在の実描画を比較できる。']
 for n,p in enumerate(demo['pages']):p.update(what_to_inspect=descriptions[n],proof_path=f'examples/repositories/slidepoise-showcase/qa/comparison-{n+1:02}.png',proof_kind='comparison-image')
 demo['source_proof_paths']+=['examples/repositories/slidepoise-showcase/source/text-layout-repair.json','examples/repositories/slidepoise-showcase/source/current-native-proof.json']+[f'examples/repositories/slidepoise-showcase/qa/comparison-{n:02}.png' for n in range(1,4)]
 patch['guide_patch'].update(summary_ja=demo['summary'],steps=patch['how_to_use'],limitations=patch['constraints'],production_status='original_runtime_all3_compared_independent_pending')
 p=ROOT/'data/skill-effect-patches/repairs-native-root.json';dump(p,[r for r in load(p) if r['case_id']!='slidepoise-showcase']+[patch])
 for f in ['data/case-studies.json','data/sol-slidepoise-repository-examples.json']:
  rows=load(ROOT/f)
  for c in rows:
   if c['id']=='slidepoise-showcase':c.update(status=qa['status'],execution_status=qa['status'],preview_image_notice='全3頁のPowerPoint実描画と原画像を比較。文字の接触・単位間隔を修復。独立再評価は未了。',prompt_text=prompt)
  dump(ROOT/f,rows)
 excluded={'.private','.build','.git','node_modules','.venv','__pycache__','backups'}
 for p in CASE.rglob('*'):
  if p.is_file() and not set(p.relative_to(CASE).parts)&excluded and p.suffix not in {'.pyc','.log'}:
   q=FORK/'example'/p.relative_to(CASE);q.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,q)
 for folder,name in [(CASE,'slidepoise-example-ja'),(FORK,'slidepoise-ja')]:
  with zipfile.ZipFile(ROOT/f'downloads/{name}.zip','w',zipfile.ZIP_DEFLATED) as z:
   for p in sorted(folder.rglob('*')):
    if p.is_file() and not set(p.relative_to(folder).parts)&excluded and p.suffix not in {'.pyc','.log'}:z.write(p,(Path(name)/p.relative_to(folder)).as_posix())
 print('SlidePoise: original reconstruction, actual3pages/fonts, comparisons and2safeZIPs synchronized.')
if __name__=='__main__':main()
