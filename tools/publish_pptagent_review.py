"""Publish full workflow guidance and actual native/font/render evidence."""
from pathlib import Path
import hashlib,importlib.util,json,shutil,zipfile
import pymupdf
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];CASE=ROOT/'examples/repositories/ppt-agent-showcase';FORK=ROOT/'skill-forks/ppt-agent-ja'
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def dump(p,d):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
SKILL='''---
name: ppt-agent-ja
description: 需要調査、資料収集、構成、ページ計画、26作風での設計を通してHTMLとPPTXを作る。
---

# 調査と構成を固めてから、ページを設計する

原典 Akxan/ppt-agent-skill / 01825dee2ad40a8c719824252605396b9a570d57 / MIT。非公式の日本語案内。upstream/SKILL.mdとreferencesを読む。

1. 目的・場面・読み手・求める行動・内容・枚数・必須/禁止を調査する。原典は3問または7問のinterviewと回答待ちを必須にする。既にユーザーが指定・委任した範囲は引き継ぎ、不足と実際の変更範囲を記録する。
2. 外部情報を検索・整理し、入力資料と照合する。情報取得能力が無ければユーザーの材料に依拠する。根拠の無い数値・実績を作らない。
3. 結論と根拠、Part、頁数へ大綱を構成する。原典の章2頁以上などの自己確認を行う。8頁以下では原典の軽量化が使える。
4. 内容とカード種類・位置・layout_hintをページ計画JSONへ分け、必要な確認を行う。原典は内容頁に3カード以上・2種類以上・dataカード1つ以上を求める。根拠不足を架空データで埋めない。規則を変える場合は用途とユーザー指定に基づく差分を明示する。
5. 実在する26作風から選び、style.jsonを作る。必要な意味のある画像を用意し、原典のHTML設計規範に沿って1280×720の頁を作る。HTMLは計画の後に作り、同じカード配置だけで全頁を埋めない。
6. 原典html_packager→html2svg（Puppeteer/dom-to-svg）→svg2pptxを実行する。Node/変換依存が使えない場合の原典の降級はHTMLのみで、PPTX成功と説明しない。実装差分がある場合は原本とdiffを残す。
7. 原典quality self-checkと全頁の実描画を行い、情報・書体・配置・図の対応を確認して直す。原典文書はPowerPoint365でSVGを「図形に変換」と説明するが、この固定pinの実converterは文字・rect/circle/line/pathをnative OOXMLへ直接解析する。納品の編集範囲は実物で確認する。

## この5枚の作例

指定済みの展示企画・Bauhaus Block・3書体・5頁で制作したため、追加interviewと外部調査、原典カード数/dataカードの割当、カード数に基づくquality条件は適用していない。幾何中心の構成と実変換経路の作例で、原典の全工程への適合や全26作風の制作を示すものではない。この差分を一般の手順へ拡張しない。

文字・図形はnative、chart/table/数式は使用していない。原典3変換の実行記録はsource/executed-route-receipts.json。Windows fileURL・font待ち・cacheのadapterと原本を同梱する。欠落していたBIZ UDPMinchoを公式TTFから導入し、現在PDFの実書体を確認。全5頁を見て3頁のlabel位置を確認した。独立再評価は別記録。

## 再現と書体

Node.js24+、puppeteer25.13.0/dom-to-svg0.12.2/esbuild0.28.2、Python＋python-pptx/lxml/Pillow、Chrome。PowerPoint描画にはWindows＋Microsoft PowerPoint。同梱fontsのZen Kaku Gothic New Black、BIZ UDPMincho、Manropeと各OFLを確認し、作成者・受け手の端末へ導入する。指定名だけでは置換を防げないので、PDF/実表示の書体も確認する。

`npm ci --ignore-scripts`、既存Python依存を用意して `python source/build.py --node <node.exe> --chrome <chrome.exe>`。新出力の全頁を確認する。再現コードは現在のHTML5頁を変換する処理で、任意テーマを自動企画するアプリではない。

原典: https://github.com/Akxan/ppt-agent-skill/tree/01825dee2ad40a8c719824252605396b9a570d57
'''

def main():
 record=load(CASE/'source/BUILD.json');assert all(s['exit_code']==0 for s in record['steps'].values())
 doc=pymupdf.open(CASE/'.private/current-native-render/deck.pdf');assert len(doc)==5
 fonts=sorted({f[3].split('+')[-1] for page in doc for f in page.get_fonts(full=True)});assert any('BIZUDPMincho' in f for f in fonts) and any('ZenKakuGothicNew' in f for f in fonts)
 for n,page in enumerate(doc,1):page.get_pixmap(matrix=pymupdf.Matrix(1600/page.rect.width,1600/page.rect.width),alpha=False).save(CASE/f'render/slide-{n:02}.png')
 shutil.copy2(CASE/'.private/current-native-render/deck.pdf',CASE/'render/deck.pdf')
 render=load(CASE/'.private/current-native-render/powerpoint-render.json');render['source']='deck.pptx';dump(CASE/'render/powerpoint-render.json',render)
 prs=Presentation(CASE/'deck.pptx');stats=[]
 for n,s in enumerate(prs.slides,1):stats.append({'page':n,'shapes':len(s.shapes),'text_shapes':sum(q.has_text_frame and bool(q.text.strip()) for q in s.shapes),'pictures':sum(q.shape_type==13 for q in s.shapes)})
 assert len(stats)==5 and all(r['pictures']==0 for r in stats)
 dump(CASE/'source/current-native-font-proof.json',{'author':'Codex actual replay','pptx_sha256':sha(CASE/'deck.pptx'),'font_file':'Official unmodified BIZUDPMincho-Regular.ttf, SHA dbcea04578ac1e9d3484525e870ce491bd04361768f4d2ba4b827d96e20f891d','actual_pdf_fonts':fonts,'native_inventory':stats,'all5_actual_pages_viewed':True})
 qa=load(CASE/'QA.json');qa.update(status='original_route_all5_actual_font_and_visual_checked',whole_deck_detailed_review=True,structure_audit=True,independent_review=False,verification='source/current-native-font-proof.json',original_execution_logs='source/executed-route-receipts.json',limits=['今回の指定で追加interview・外部検索・カード数規則を適用していない。原典の全工程合格ではない。','native chart/table/数式、全26作風と配図生成は未実施。','修復後の独立再評価は未了。']);
 if (CASE/'source/zip-replay.json').exists():qa.update(zip_reexecution=True,distribution_replay='source/zip-replay.json')
 dump(CASE/'QA.json',qa)
 prov=load(CASE/'PROVENANCE.json');prov['current_replay']={'author':'Codex','route':'Original3converters, same authored HTML/native geometry, actual Windows font repair and all5PowerPoint views.','proof':'source/current-native-font-proof.json','independent_review':'pending'};dump(CASE/'PROVENANCE.json',prov)
 for folder in [CASE,FORK]:(folder/'SKILL.md').write_text(SKILL,encoding='utf-8')
 readme='''# 企画から作風をそろえたHTMLとPPTXへ

需要調査・資料収集・大綱・頁計画・26作風の選択を経て、原典3変換で編集可能PPTXを作るスキル。一般の手順はSKILL.md、固定出典はSOURCE.json、利用条件はLICENSEとfontsのOFLを参照。

## 作例の条件

指定済み展示企画・Bauhaus Block・3書体・5頁。追加interview・外部検索・原典カード数規則はこの作例に適用していない。任意テーマを自動企画するアプリではない。全26作風・配図生成・native chart/table/数式は未実施。

## 現在版の確認

原典html_packager→html2svg→svg2pptxを再実行し、全5頁をPowerPointで描画して見た。3頁のラベル位置は先行改訂で修復済み。今回、公式BIZ UDPMinchoを導入し、出力PDF内のZen Kaku Gothic New・BIZ UDPMincho・Manropeを実確認した。指定書体名だけで実使用を判断しない。

source/executed-route-receipts.jsonが実変換記録、source/current-native-font-proof.jsonが現在の書体とnative要素。source/zip-replay.jsonは配布ZIPを展開し、同梱buildで3変換を再実行した記録。5頁のnative slide XMLが一致した。独立再評価は別記録。

## 再現

SKILL.mdの環境と字体を用意して `npm ci --ignore-scripts`、`python source/build.py --node <node.exe> --chrome <chrome.exe>`。新出力の全頁・実書体・編集範囲を確認する。
'''
 (CASE/'README.md').write_text(readme,encoding='utf-8')
 for p in CASE.rglob('*'):
  if p.is_file() and not set(p.relative_to(CASE).parts)&{'.private','.build','node_modules','__pycache__'} and p.suffix not in {'.pyc','.log'}:
   q=FORK/p.relative_to(CASE);q.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,q)
 patch=next(r for r in load(ROOT/'data/skill-effect-patches/native-03.json') if r['case_id']=='ppt-agent-showcase');patch=json.loads(json.dumps(patch));steps=['目的・読み手・材料を調査し、根拠を集める。既存指定と不足を分ける。','大綱と頁計画を作ってから、26作風の選択・配図・HTML設計を進める。','原典3変換でHTML/SVG/PPTXへ出力し、全頁の情報・書体・配置・編集範囲を確認する。']
 patch.update(name='調査と構成から、作風をそろえたHTMLとPPTXへ',upstream_principle='需要調査→資料収集→大綱→頁計画→26作風での設計→HTML/SVG/PPTX出力の全工程を分けて進める。',change='企画の根拠と頁の内容を先に固め、用途に合わせた作風でスライドへ設計できる。',input='テーマ・目的・受け手・発表条件・元資料・必須/禁止・枚数',output='HTMLプレビューと頁別SVG/PPTX。変換環境が無い場合は原典の降級条件を確認する。',how_to_use=steps,constraints=qa['limits'])
 patch['demonstration'].update(summary='指定済み展示企画を5頁へ計画し、原典Bauhaus BlockのCSSと幾何で設計。原典HTML→DOM/SVG→native PPTXを実行し、全5頁と出力書体を確認。追加interview・外部検索・カード数規則は今回の指定に合わせて不適用。',not_demonstrated=['追加interview・外部調査','原典の3カード/2種/dataカード規則とカード数によるquality条件','全26作風・配図生成','native chart/table/数式','修復後の独立再評価'])
 patch['demonstration']['source_proof_paths']+=['examples/repositories/ppt-agent-showcase/source/executed-route-receipts.json','examples/repositories/ppt-agent-showcase/source/current-native-font-proof.json','examples/repositories/ppt-agent-showcase/render/powerpoint-render.json','examples/repositories/ppt-agent-showcase/upstream/scripts/svg2pptx.py']
 patch['guide_patch'].update(summary_ja=patch['change'],steps=steps,limitations=qa['limits'],production_status='original_three_converters_all5_actual_fonts_checked_independent_pending')
 p=ROOT/'data/skill-effect-patches/repairs-native-root.json';rows=load(p);dump(p,[r for r in rows if r['case_id']!='ppt-agent-showcase']+[patch])
 spec=importlib.util.spec_from_file_location('pack',ROOT/'tools/build_ppt_agent_showcase.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);m.package();print('PPT Agent: full general workflow, explicit case differences, actual font repair, all5views and2ZIPs synchronized.')
if __name__=='__main__':main()
