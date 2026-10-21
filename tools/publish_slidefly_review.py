"""Publish original-engine motion evidence and actual native export inspection."""
from pathlib import Path
import difflib,hashlib,json,shutil,zipfile
import pymupdf
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];CASE=ROOT/'examples/repositories/slidefly-showcase';FORK=ROOT/'skill-forks/slidefly-ja'
load=lambda p:json.loads(p.read_text(encoding='utf-8-sig'))
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def dump(p,v):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
SKILL='''---
name: slidefly-ja
description: 同じ図形を頁間で動かし、話をつなぐHTMLと編集可能なPPTXを作る。
---

# 同じ図形が、次のページへつながる

原典andyluu98/slidefly / 7de466df3f2e0199c4939ccbe4490688d4165fa3 / MIT。非公式の日本語案内。upstream/SKILL.mdと必要referencesを読む。

1. 目的・読み手・元資料・枚数・文字量・作風を整理する。原典は85作風、21頁型、17フレームを備える。作風指定がなければpick-styles.pyの3候補を参考にする。34の二層作風では対応.art.mdを読み、主張ごとの専用SVGを用意する。
2. 頁ごとの主張と型を計画する。同じ形がどの関係を引き継ぐかを決め、named actorとposeへ対応させる。原典engineがactorsを作るため、独自の.actorsブロックで代替しない。独自作風は原典のfree-layoutとMorph契約を読む。
3. 原典CSS/engine/nav/frames等を決められた順に読み込み、本文をreveal要素へ置く。同じ図形やdata-morph-idの要素が、移動・拡縮・形状変化で次の頁につながる。文字の配置と図形の動線を分ける。
4. inline-assets.pyで単一HTMLへまとめる。ブラウザーで全頁を開き、文字・重なり・余白・型の反復とdeck.audit()を確認し、動きの途中も見て直す。見た目の合格は監査結果だけで決めない。
5. PPTXが必要ならexport-pptx.pyで出力する。文字とactorは編集可能、inline SVGはvector画像。変換器が認識するクラス/タグ・明示改行・SVGの配置条件を守り、PowerPointでも全頁を開く。Morph設定があることと実際のアニメーション再生は分けて確認する。

## この作例で試したこと

独自walk作風の6頁。道・地点は創作で、実在の案内ではない。元の3actorが全5区間で移動・拡縮する。元の次頁ボタンでbefore/100ms中間/1400ms後を記録し、同じIDを保持することを確認。全6頁のWeb描画と原典監査も確認した。

独自h1/h2が原典exporterで除外されるため、同じ見た目のrole=heading付きdivへ変更。改行はbr、SVGは原典対応の直接配置へ移し、currentColorを同じ明示色へ解決した。原典コードは未改変。現在の全6頁をPowerPointで開き、見出しと分岐図の欠落がないことを確認。別コピーのnative文字変更・保存・再表示も実施した。Webの英数字Space GroteskはPPTXでは原典body字体のZen Kaku Gothic Newへ置換される。

## 入力と再現

Python3.10+標準ライブラリ、モダンブラウザー。PPTX表示はPowerPoint2019/365と使用字体。原典Google Fontsのオンライン読込とは異なり、この作例は公式TTF/OFLをHTMLに内包している。PPTXは字体非埋込。fonts/を利用端末へ導入する。

`python upstream/scripts/inline-assets.py source.html index.html` → `python upstream/scripts/export-pptx.py index.html deck.pptx`。別題材では入力・各頁の主張・作風・actor/pose・必要な図を先に設計する。再現コードは保存済み作例の梱包・変換で、任意の原稿を自動企画するアプリではない。

85作風の比較、二層作風の全頁専用イラスト、PowerPoint内の実Morph再生、独立した最終再評価は未実施。
'''
def main():
 rows=load(CASE/'source/motion-browser-proof.json')['rows'];assert len(rows)==5
 assert all(r['before']['actors']!=r['mid']['actors']!=r['after']['actors'] for r in rows)
 audit=json.loads(load(CASE/'source/current-original-audit.json')['result']['value']);assert len(audit)==6 and all(x['status']=='OK' for x in audit)
 doc=pymupdf.open(CASE/'.private/current-native-render/deck.pdf');assert len(doc)==6
 for n,pg in enumerate(doc,1):
  out=CASE/f'render/native/slide-{n:02}.png';out.parent.mkdir(parents=True,exist_ok=True);pg.get_pixmap(matrix=pymupdf.Matrix(1600/pg.rect.width,1600/pg.rect.width),alpha=False).save(out)
 shutil.copy2(CASE/'.private/current-native-render/deck.pdf',CASE/'render/native/deck.pdf')
 edit=load(CASE/'qa/editability/native-editability.json');assert edit['original_sha256']==sha(CASE/'deck.pptx') and edit['saved_edit_reopened']
 ed=pymupdf.open(CASE/'qa/editability/native-text-edited.pdf');ed[0].get_pixmap(matrix=pymupdf.Matrix(1600/ed[0].rect.width,1600/ed[0].rect.width),alpha=False).save(CASE/'qa/editability/native-text-edited.png')
 prs=Presentation(CASE/'deck.pptx');inventory=[{'page':n,'native_texts':[q.text for q in s.shapes if q.has_text_frame and q.text.strip()],'pictures':sum(q.shape_type==13 for q in s.shapes)} for n,s in enumerate(prs.slides,1)];assert all(len(x['native_texts'])>=4 for x in inventory)
 dump(CASE/'source/native-export-proof.json',{'pptx_sha256':sha(CASE/'deck.pptx'),'all6_PowerPoint_pages_viewed':True,'pdf_fonts':sorted({f[3].split('+')[-1] for pg in doc for f in pg.get_fonts(full=True)}),'inventory':inventory,'PowerPoint_Morph_playback':False,'native_edit_saved_reopened':True})
 before=(CASE/'.private/before-branch-label-repair/source.html').read_text(encoding='utf-8');after=(CASE/'source.html').read_text(encoding='utf-8');(CASE/'source/markup-compatibility.diff').write_text(''.join(difflib.unified_diff(before.splitlines(True),after.splitlines(True),fromfile='initial source.html',tofile='source.html')),encoding='utf-8')
 qa=load(CASE/'QA.json');qa.setdefault('initial_check_history',dict(qa));qa.update(status='all6_web_native_and5_motion_boundaries_checked',all_pages_detailed_visual_review=True,native_structure_audit=True,independent_review=False,motion_check='source/motion-browser-proof.json',original_audit='source/current-original-audit.json',native_export='source/native-export-proof.json',not_performed=['PowerPoint Morph playback','All85 styles / two-layer styles','Independent final recheck']);
 if (CASE/'source/zip-replay.json').exists():qa.update(zip_reexecution=True,distribution_replay='source/zip-replay.json')
 dump(CASE/'QA.json',qa)
 prov=load(CASE/'PROVENANCE.json');prov['current_repair']={'author':'Codex','changes':'Source markup compatibility, explicit SVG labels/ink, readable pager; original engine/inliner/exporter unchanged.','native_export':'source/native-export-proof.json','motion':'source/motion-browser-proof.json','independent_review':'pending'};dump(CASE/'PROVENANCE.json',prov)
 for p in [CASE,FORK]:(p/'SKILL.md').write_text(SKILL,encoding='utf-8')
 (CASE/'README.md').write_text('# 同じ図形を次のページへ動かす\n\n原典SlideFlyのnamed actorとposeで6頁をつなぐ作例。全6頁のWeb表示・原典監査、全5区間のbefore/mid/afterを確認した。render/motion/とsource/motion-browser-proof.jsonを参照。\n\nPPTXは原典exporterが読めるHTML構造へ補正し、全6頁をPowerPointで開いた。文字/actorはnative、分岐図はvector画像。別コピーで文字編集・保存・再表示を実施。PPTXの英数字は原典body字体へ置換され、Webと完全同一ではない。PowerPoint内の実Morph再生と独立再評価は未了。\n\n使い方と未実施範囲はSKILL.md。source.htmlとwalk.cssを編集して原典inline-assets.py→export-pptx.pyで再現する。MITとfonts/OFLを保持。\n',encoding='utf-8')
 base=CASE.relative_to(ROOT).as_posix();patch={'case_id':'slidefly-showcase','name':'同じ図形を動かして、ページ間の関係をつなぐ','job':'','problem':'頁が切り替わるたびに、同じ図形や位置関係を読み直してしまう。','change':'同じ要素のIDを保ったまま移動・拡縮し、前の頁から何が変わったかを追える。','input':'原稿、目的・聴衆・枚数・作風、頁ごとの主張、引き継ぐ要素とその状態','output':'Morph状の単一HTML。必要なら文字/actorを編集できるPPTX。','how_to_use':['目的と頁ごとの主張を整理し、作風・頁型・frameを選ぶ。','同じ関係を引き継ぐactor/poseを決め、原典HTML/CSS/runtimeで構成する。','単一HTML化し、全頁と動きの途中を確認して原典auditの指摘を直す。','必要なら原典PPTX exporterで出力し、変換後の全頁・編集・実再生を別々に確認する。'],'constraints':['独自walk作風の例。85作風の比較と二層作風の全頁専用図は未実施。','PPTXは文字とactorがnative、分岐図はvector画像。HTMLの英数字書体はPPTXで原典body字体に置換される。','PowerPoint内の実Morph再生と独立した最終再評価は未了。'],'demonstration':{'summary':'原典3actorを6poseでつなぎ、全5区間の実before/mid/after、Web6頁、PowerPoint6頁とnative文字編集を確認した。','applied_features':['同じnamed actorの移動・拡縮・形状変化','独自free-layoutと原典監査','原典単一HTML化とPPTX exporter'],'not_demonstrated':['85作風の比較・二層作風の全頁専用イラスト','PowerPoint Morph playback','独立した最終再評価'],'pages':[{'index':i,'label':f'{i+1}ページ目','title':prs.slides[i].shapes[0].name if False else title,'skill_rule':rule,'what_to_inspect':explain,'proof_path':f'{base}/render/slide-{i+1:02}.png','proof_kind':'render-image'} for i,title,rule,explain in [(1,'同じ3要素が次の配置へ動く','named actorのIDを保ってposeを変える。','丸・線・四角が頁をまたいで移動し、別の3要素へ置き換わらない。中間描画もmotion/で確認できる。'),(2,'分岐の意味を図のラベルに対応させる','free-layoutの図を内容に合わせて設計する。','出発・曲がり角・明るい通り・静かな道を図の各位置へ対応。PPTXも図の欠落を修復した。'),(5,'拡大と形状変化で終わりを示す','同じactorの大きさや形を変え、話の進行をつなぐ。','赤い丸と黄色い目印が拡大・変形。native PPTX版も文字とactorを保持する。')]],'source_proof_paths':[f'{base}/source/{x}' for x in ['motion-browser-proof.json','current-original-audit.json','native-export-proof.json','markup-compatibility.diff']]+[f'{base}/render/motion/boundary-{n}-{phase}.png' for n in range(1,6) for phase in ['before','mid','after']]+[f'{base}/render/native/slide-{n:02}.png' for n in range(1,7)]+[f'{base}/qa/editability/native-text-edited.png']},'guide_patch':{'requirements':['Python3.10+標準ライブラリとモダンブラウザー','PPTX表示はPowerPoint2019/365と使用字体。非埋込のため受け手にも導入する。'],'execution_type':'original_morph_engine_inliner_exporter_all6_native_opened','production_status':qa['status']}}
 prompt='原典SlideFlyを読み、目的・読み手・原稿から頁の主張、作風、actorとposeを設計する。元engine/inliner/exporterを使い、全頁のWeb表示と動きの途中、PPTX全頁と編集範囲を確認し、未実施を分けて渡す。'
 (CASE/'PROMPT.txt').write_text(prompt,encoding='utf-8')
 patch.update(status=qa['status'],execution_status=qa['status'],prompt_text=prompt,preview_notice='全6頁と全5遷移の実描画、原典監査、PPTX全6頁とnative文字編集を確認。PowerPoint内のMorph実再生は未了。',preview_image_notice='原典HTMLの安定状態を1280×720で全6頁実撮影。中間フレームは別の動きの証拠へ。')
 dump(CASE/'source/current-artifact-binding.json',{'html_sha256':sha(CASE/'index.html'),'pptx_sha256':sha(CASE/'deck.pptx'),'motion_proof':'source/motion-browser-proof.json','source_markup':'source/markup-compatibility.diff'})
 p=ROOT/'data/skill-effect-patches/repairs-root-slidefly.json';dump(p,[patch])
 excluded={'.private','.build','.git','node_modules','.venv','__pycache__'}
 for p in CASE.rglob('*'):
  if p.is_file() and not set(p.relative_to(CASE).parts)&excluded and p.suffix not in {'.pyc','.log'}:
   q=FORK/'example'/p.relative_to(CASE);q.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,q)
 for folder,name in [(CASE,'slidefly-example-ja'),(FORK,'slidefly-ja')]:
  with zipfile.ZipFile(ROOT/f'downloads/{name}.zip','w',zipfile.ZIP_DEFLATED) as z:
   for p in sorted(folder.rglob('*')):
    if p.is_file() and not set(p.relative_to(folder).parts)&excluded and p.suffix not in {'.pyc','.log'}:z.write(p,(Path(name)/p.relative_to(folder)).as_posix())
 print('SlideFly original motion5boundaries/Web6/native6/editability and2ZIPs synchronized.')
if __name__=='__main__':main()
