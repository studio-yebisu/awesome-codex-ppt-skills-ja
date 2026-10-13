"""Japanese teaching example using the pinned, unmodified slides_maker deckkit."""
from pathlib import Path
import argparse, hashlib, json, shutil, subprocess, sys
ROOT=Path(__file__).resolve().parents[1]
COMMIT='3767cc346d63ef1646082c87f15fbf0be582b152'
def dump(p,v):
    p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def linear_search(values,target):
    visited=[]
    for index,value in enumerate(values):
        visited.append(value)
        if value==target:return {'index':index,'visited':visited,'comparisons':len(visited)}
    return {'index':-1,'visited':visited,'comparisons':len(visited)}
def binary_search(values,target):
    lo,hi=0,len(values)-1;visited=[];steps=[]
    while lo<=hi:
        mid=(lo+hi)//2;value=values[mid];visited.append(value)
        steps.append({'low':values[lo],'high':values[hi],'mid':value})
        if value==target:return {'index':mid,'visited':visited,'comparisons':len(visited),'steps':steps}
        if value<target:lo=mid+1
        else:hi=mid-1
    return {'index':-1,'visited':visited,'comparisons':len(visited),'steps':steps}
TITLES=['並べておくと、探し方が変わる','先頭からなら、13まで13回','真ん中で分けると、4回で届く','15個すべてで試すと、最大4回','使える条件も、いっしょに覚える']
NOTES=[
'初めて探索を学ぶ人への短い教材です。同じ昇順の配列から、実際にPythonで13を探しました。ここで数える1回は、調べる要素を1つ選んで目的値と比べることです。CPU命令数や秒数ではありません。',
'線形探索では先頭から順に値を確認します。目的値が13なら1から13まで調べます。この方法は、データが未整列でも使えます。今回は両方式で同じ配列を使って条件をそろえました。',
'二分探索は昇順という前提を利用します。8より大きいので9〜15、12より大きいので13〜15、14より小さいので13だけを残し、次に13を調べて終了します。中央は添字の小さい側へ切り捨てています。',
'1から15の全てを目的値にして実行しました。線形探索は先頭の1なら1回、末尾の15なら15回。二分探索は8なら1回で、最多でも4回でした。存在する値だけを探す15通りの比較であり、未発見時や大規模データの計測は含めていません。',
'二分探索の前に整列が必要です。今回の実行では初めから昇順の配列を与えたため、並べ替えの手間は計測していません。途中で並びを変えれば前提が崩れます。配列の中央を直接参照できる条件での教材であり、検索回数の減少をそのまま何倍速いと読み替えないでください。']
def build(out,up):
    scripts=up/'skills/slide-maker/scripts';sys.path.insert(0,str(scripts))
    import deckkit as dk
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
    dk.FONT=dk.EAFONT='Yu Gothic';dk.MONO='Consolas'
    ink=RGBColor.from_string('132C3B');blue=RGBColor.from_string('116B8A');orange=RGBColor.from_string('A8471C');muted=RGBColor.from_string('4E626C');white=RGBColor(255,255,255)
    dk.DEEP=ink
    values=list(range(1,16));traces=[{'target':t,'linear':linear_search(values,t),'binary':binary_search(values,t)} for t in values]
    data={'title':TITLES[0],'titles':TITLES,'notes':NOTES,'array':values,'unit':'訪問した要素数（目的値と照合した回数）','timing_measured':False,'sorting_cost_measured':False,'traces':traces}
    assert traces[12]['linear']['comparisons']==13 and traces[12]['binary']['visited']==[8,12,14,13]
    assert max(t['binary']['comparisons'] for t in traces)==4
    dump(out/'input/data.json',data)
    prs=dk.blank_deck(13.333,7.5)
    def txt(s,x,y,w,h,t,size=24,c=ink,bold=False,center=False):
        return dk.text(s,x,y,w,h,[[(t,size,c,bold,False,dk.FONT)]],align=PP_ALIGN.CENTER if center else PP_ALIGN.LEFT,anchor=MSO_ANCHOR.MIDDLE,space_after=0)
    def start(i):
        s=dk.add_slide(prs);dk.a11y_title(s,TITLES[i],ea=dk.EAFONT);dk.title_bar(s,TITLES[i],kicker='探索のしくみ  /  1〜15で実行',accent=blue,title_c=ink)
        dk.footer(s,tag='条件：昇順の配列 1〜15 ／ 比較回数＝訪問した要素数',page=i+1)
        s.notes_slide.notes_text_frame.text=NOTES[i]
        return s
    def row(s,y,active,hit=None):
        for j,v in enumerate(values):
            fill='116B8A' if v in active else 'E8EEF0'
            if v==hit:fill='A8471C'
            sh=dk.box(s,.62+j*.805,y,.71,.65,fill=fill)
            sh._element.nvSpPr.cNvPr.set('descr',f'値 {v}'+('：確認する値' if v in active else '：対象外'))
            txt(s,.62+j*.805,y,.71,.65,str(v),20,white if v in active or v==hit else muted,True,True)
    s=start(0)
    txt(s,.65,1.65,11.7,.55,'15個の数字から「13」を見つける。',30,bold=True)
    row(s,2.65,[],13)
    txt(s,.7,4.05,5.3,.5,'先頭から、ひとつずつ',25,orange,True)
    txt(s,7,4.05,5.4,.5,'真ん中から、範囲を半分に',25,blue,True)
    txt(s,.7,4.8,5.3,.8,'13回',46,orange,True)
    txt(s,7,4.8,5.3,.8,'4回',46,blue,True)
    txt(s,.7,6.1,11.9,.4,'同じデータでも、並び方を使うと調べる場所を減らせる。',22)
    s=start(1);row(s,2.35,list(range(1,14)),13)
    txt(s,.68,1.65,11.8,.45,'線形探索：見つかるまで、左から順に照合',26,bold=True)
    txt(s,.7,3.55,8.4,.65,'1 → 2 → 3 → … → 12 → 13',32,orange,True)
    txt(s,.7,4.65,8.5,.8,'今回は13回目で見つかった。\n並んでいないデータにも使える。',25)
    txt(s,10.1,3.65,2.4,1.1,'13',66,orange,True,True);txt(s,10.1,4.8,2.4,.5,'回の比較',23,orange,True,True)
    s=start(2)
    for k,step in enumerate(traces[12]['binary']['steps']):
        y=1.7+k*1.06
        txt(s,.55,y,.7,.5,str(k+1),24,blue,True)
        txt(s,1.2,y,2.35,.5,f"範囲 {step['low']}〜{step['high']}",21)
        decision=['13 > 8：右へ','13 > 12：右へ','13 < 14：左へ','13 = 13：発見'][k]
        txt(s,1.2,y+.51,2.35,.34,decision,16,orange,True)
        for j,v in enumerate(values):
            x=3.65+j*.58;active=step['low']<=v<=step['high'];mid=v==step['mid'];fill='A8471C' if mid else ('D8EAF0' if active else 'F2F5F6')
            sh=dk.box(s,x,y,.51,.52,fill=fill);sh._element.nvSpPr.cNvPr.set('descr',f'第{k+1}回 値{v}'+(' 中央で照合' if mid else ''))
            txt(s,x,y,.51,.52,str(v),16,white if mid else ink if active else muted,mid,True)
    txt(s,.65,6.12,12,.45,'茶色＝照合する中央の値　　8 → 12 → 14 → 13 の4回',23,bold=True)
    s=start(3)
    txt(s,.65,1.53,11.9,.5,'目的値を1〜15に変えて、全15通りを実行',25,bold=True)
    chart=dk.native_chart(s,.55,2.14,9.25,3.95,[str(v) for v in values],[('線形探索',[t['linear']['comparisons'] for t in traces]),('二分探索',[t['binary']['comparisons'] for t in traces])],kind='line_markers',palette=[orange,blue],font=dk.FONT,smooth=False,value_fmt='0')
    chart.value_axis.minimum_scale=0;chart.value_axis.maximum_scale=16;chart.value_axis.major_unit=4
    txt(s,.7,2.03,2,.3,'比較回数（回）',13,muted);txt(s,7.4,6.13,2.35,.4,'探す値 →',16,muted)
    dk.takeaway_rail(s,10.15,2.45,2.6,'二分探索の最大','4回','この15通りでは、\n線形探索は最大15回。',accent=blue,ink=ink,body_c=muted)
    txt(s,.7,6.56,11.9,.3,'実行結果：input/data.json ／ 配列内に存在する値だけを比較',14,muted)
    s=start(4)
    txt(s,.7,1.7,11.8,.65,'二分探索の前提は「昇順が保たれていること」。',29,blue,True)
    txt(s,.8,2.85,5.4,.55,'今回、確かめたこと',25,bold=True)
    txt(s,.8,3.63,5.4,1.25,'1〜15の全15通り\n二分探索の比較は最大4回',25)
    txt(s,7.1,2.85,5.4,.55,'今回、測っていないこと',25,bold=True)
    txt(s,7.1,3.63,5.4,1.25,'実際にかかる秒数\n並べ替えにかかる手間',25)
    txt(s,.8,5.65,11.8,.8,'「比較が少ない」と「何倍速い」は、分けて考える。',28,orange,True)
    out.mkdir(parents=True,exist_ok=True);dk.declare_delivery(out,'selfread',builds='static');prs.core_properties.language='ja-JP';prs.save(out/'deck.pptx')
    dump(out/'roster.json',[{'title':t,'note':n} for t,n in zip(TITLES,NOTES)])
    (out/'NOTES.md').write_text('\n\n'.join(f'## {i+1}. {t}\n\n{n}' for i,(t,n) in enumerate(zip(TITLES,NOTES))),encoding='utf-8')
    print(json.dumps({'slides':len(prs.slides),'target13':traces[12],'output':str(out/'deck.pptx')},ensure_ascii=False))
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--output-dir',type=Path,default=ROOT/'examples/repositories/slide-maker');ap.add_argument('--upstream-dir',type=Path,default=ROOT/'.build/actual-upstreams/addsumtech__slides_maker');a=ap.parse_args();build(a.output_dir.resolve(),a.upstream_dir.resolve())
