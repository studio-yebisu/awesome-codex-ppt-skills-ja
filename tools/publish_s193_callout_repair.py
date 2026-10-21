"""Synchronize the inspected S193 repair and its two distributable archives."""
from pathlib import Path
from copy import deepcopy
import hashlib, json, shutil, zipfile
from pptx import Presentation
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
CASE = ROOT / 'examples/repositories/sebastian-193'
FORK = ROOT / 'skill-forks/sebastian-193-ja'
EXCLUDED = {'.private', '.build', '.git', '.venv', 'node_modules', '__pycache__'}
load = lambda p: json.loads(p.read_text(encoding='utf-8-sig'))
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()

def dump(p, value):
    p.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')

def files(base):
    return [p for p in sorted(base.rglob('*')) if p.is_file()
            and not set(p.relative_to(base).parts) & EXCLUDED
            and p.suffix not in {'.log', '.pyc'}]

def walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:
            yield from walk(shape.shapes)

def main():
    artifact = sha(CASE/'deck.pptx')
    before = CASE/'.private/round4-before-callout-fix'
    unchanged = [n for n in range(1, 9) if sha(CASE/f'render/slide-{n:02}.png') == sha(before/f'render/slide-{n:02}.png')]
    assert unchanged == [2, 4, 6], unchanged
    with zipfile.ZipFile(before/'deck.pptx') as old, zipfile.ZipFile(CASE/'deck.pptx') as new:
        protected = [n for n in old.namelist() if n.startswith(('ppt/slideMasters/', 'ppt/slideLayouts/', 'ppt/theme/', 'ppt/notesSlides/'))]
        assert all(old.read(n) == new.read(n) for n in protected)
        assert all(old.read(f'ppt/slides/slide{n}.xml') == new.read(f'ppt/slides/slide{n}.xml') for n in unchanged)
    edit = load(CASE/'qa/editability/group-move.json')
    assert edit['passed'] and edit['source_sha256'] == artifact
    qa_report = load(CASE/'source/qa-report.json')
    assert qa_report['summary'] == {'error': 0, 'warning': 0, 'info': 0}
    report = load(CASE/'source/structure-report.json')
    report['details'].update(input='deck.pptx', output='structured.pptx')
    dump(CASE/'source/structure-report.json', report)
    stats = []
    for i, slide in enumerate(Presentation(CASE/'deck.pptx').slides):
        shapes = list(walk(slide.shapes))
        stats.append({'index': i, 'native_shapes': len(shapes), 'groups': sum(x.shape_type == 6 for x in shapes),
                      'pictures': sum(x.shape_type == 13 for x in shapes), 'native_charts': sum(x.has_chart for x in shapes),
                      'native_tables': sum(x.has_table for x in shapes),
                      'texts': [x.text for x in shapes if x.has_text_frame and x.text.strip()]})
    dump(CASE/'source/native-elements.json', {'pptx_sha256': artifact, 'includes_group_children': True, 'slides': stats})
    pdf = pymupdf.open(CASE/'render/deck.pdf')
    dump(CASE/'source/pdf-fonts.json', {'scope': 'Actual PowerPoint PDF font resources', 'pages': [
        {'index': i, 'fonts': [{'base_font': f[3], 'type': f[2]} for f in p.get_fonts(full=True)]} for i, p in enumerate(pdf)]})
    pdf.close()
    repair = {'author': 'Codex', 'artifact_sha256': artifact, 'changed_pages': [1, 3, 5, 7, 8],
              'changed_pages_fullsize_visually_checked': True, 'unchanged_pages_pixel_identical': unchanged,
              'protected_master_layout_theme_notes_parts_byte_identical': len(protected),
              'original_structure_tool': 'source/upstream/scripts/structure_pptx.py', 'native_groups': 4, 'bound_connectors': 4,
              'actual_PowerPoint_group_move': 'qa/editability/group-move.json', 'independent_recheck': 'pending'}
    dump(CASE/'source/callout-repair.json', repair)
    render = load(CASE/'render/powerpoint-render.json'); render['source'] = 'deck.pptx'; dump(CASE/'render/powerpoint-render.json', render)
    replay = CASE/'source/distribution-replay.json'
    replay_data = load(replay) if replay.exists() and load(replay).get('current_pptx_sha256') == artifact else {'status': 'pending_current_callout_repair'}
    qa = load(CASE/'QA.json')
    qa.update(artifact_sha256=artifact, current_callout_repair='source/callout-repair.json',
              original_qa=qa_report['summary'], zip_replay=replay_data,
              structure={k: sum(x[k] for x in stats) for k in ['native_shapes', 'groups', 'pictures', 'native_charts', 'native_tables']},
              images=[{'page': n, 'file': f'render/slide-{n:02}.png', 'sha256': sha(CASE/f'render/slide-{n:02}.png')} for n in range(1, 9)])
    qa['not_demonstrated'] = [x for x in qa['not_demonstrated'] if x != 'Interactive move/edit test']
    qa['verified_edit_scope'] = 'One box-and-label group moved in PowerPoint, saved and reopened; four native arrows remain attached.'
    dump(CASE/'QA.json', qa)
    prov = load(CASE/'PROVENANCE.json')
    prov.pop('upstream_full_qa_pass_claim', None)
    prov.update(artifact_sha256=artifact, full_upstream_feature_coverage=False, original_final_qa_verified=True,
                current_callout_repair='source/callout-repair.json', current_distribution_replay='source/distribution-replay.json')
    prov['not_demonstrated'] = [x for x in prov['not_demonstrated'] if x != 'Interactive move/edit test']
    prov['repair'] = repair
    dump(CASE/'PROVENANCE.json', prov)
    current = {'scope': 'Final original QA, structure and real group edit', 'pptx_sha256': artifact, 'receipts': [
        {'command': 'python source/upstream/scripts/qa_pptx.py deck.pptx --profile group-meeting --report source/qa-report.json', 'exit_code': 0, 'log': 'source/original-qa.txt'},
        {'command': 'python source/upstream/scripts/verify_final_qa.py deck.pptx source/qa-report.json --require-profile group-meeting', 'exit_code': 0, 'log': 'source/original-verify-final-qa.txt'}]}
    dump(CASE/'source/current-check-receipts.json', current)
    rows = load(ROOT/'data/skill-effect-patches/repairs-native4.json')
    row = deepcopy(next(x for x in rows if x['case_id'] == 'sebastian-193'))
    row['guide_patch']['summary_ja'] = '研究資料を、参考テンプレートの書式と内容に合う図解を持つ、編集可能なPowerPointへまとめる。'
    row['guide_patch']['requirements'] = ['研究の元資料と、参考にするPPTXや画像。', 'Node.jsとPptxGenJS 4.0.1、Python 3.10以上とlxml・pypdfium2。', '同梱例の確認はWindowsのPowerPointとPowerShell 7。BIZ UDPGothicは公式OFL付きファイルを同梱。']
    row['constraints'][-1] = '原典の静的検査と完成ファイルの照合を実施。8頁では箱とラベルを一緒に移動し、保存・再表示後も4本の接続を保持することを確認。独立再評価は未了。'
    row['how_to_use'] = ['研究の元資料・聴衆・発表時間・出力形式を決め、使うテンプレートや参考画像を渡す。',
        '話の組み立てを選び、参考資料の書式と元の図・根拠を整理する。各頁の計画表を作ってから、具体的な構図と別案を決める。',
        '新規作成・既存PPTXの修正・画像からの復元のうち、入力に合う原典の方式で制作する。制作範囲の判断が必要なら先に確認する。',
        '全頁を描画して確認し、原典の検査を完成ファイルへ実行する。グループや接続線は保存したコピーを動かして確かめる。']
    row['guide_patch']['steps'] = row['how_to_use']
    row['guide_patch']['limitations'] = row['constraints']
    demo = row['demonstration']
    demo['summary'] = 'テンプレートの青・見出し・マスターを保ち、順序・比較・経路・注釈を8頁へ組む。最後の工程図には箱と文字のグループ、動かした箱に追従する接続線を使った。'
    demo['not_demonstrated'] = ['画像案の生成、画像からの復元、実論文の抽出、実施設での評価、修復後の独立再評価。']
    for p in demo['pages']:
        p['title'] = load(CASE/'input/data.json')['slides'][p['index']]['title']
    demo['pages'][-1]['what_to_inspect'] = '主行き先と補助情報を指す線の先端を見る。「内側の余白」は左下の寸法線で示す。'
    demo['pages'].append({'index': 7, 'label': '8頁・修正の循環', 'title': load(CASE/'input/data.json')['slides'][7]['title'],
        'what_to_inspect': '修正から試作へ戻る矢印をたどる。下の実編集例では、手順確認の箱を動かしても文字と接続線が追従する。',
        'skill_rule': '一緒に動く箱と文字はグループ化し、接続線を箱へ結び付ける。保存後に実際のPowerPointで確認する。',
        'proof_path': 'examples/repositories/sebastian-193/render/slide-08.png', 'proof_kind': 'powerpoint-render'})
    demo['source_proof_paths'] += ['examples/repositories/sebastian-193/'+p for p in ['source/structure-plan.json', 'source/structure-report.json', 'source/callout-repair.json', 'qa/editability/group-move.json', 'qa/editability/slide-08.png']]
    demo['editability'] = {'before_path': 'examples/repositories/sebastian-193/render/slide-08.png', 'after_path': 'examples/repositories/sebastian-193/qa/editability/slide-08.png', 'before_label': '元の工程図', 'after_label': '手順確認の箱を動かしたPPTX', 'artifact_path': 'examples/repositories/sebastian-193/qa/editability/group-moved.pptx', 'artifact_label': '動かしたPPTXを取得 ↓', 'report_path': 'examples/repositories/sebastian-193/qa/editability/group-move.json', 'description': '手順確認の箱を右へ24pt・下へ18pt動かして保存し、開き直した。ラベルは箱と一緒に動き、4本の線は接続を保っている。'}
    proposal = load(CASE/'source/catalog-repair-proposal.json')
    proposal.update(current_pptx_sha256=artifact, current_repair='source/callout-repair.json', current_independent_review='pending')
    proposal['guide_patch'] = row['guide_patch']
    proposal['page_proofs'] = [{'index': n, 'image': f'render/slide-{n+1:02}.png'} for n in range(8)]
    dump(CASE/'source/catalog-repair-proposal.json', proposal)
    path = ROOT/'data/skill-effect-patches/zz-root-consistency.json'
    dump(path, [x for x in load(path) if x['case_id'] != 'sebastian-193']+[row])
    old = 'native要素の対話的な移動試験と独立再評価は未実施。'
    new = '8頁の箱と文字をグループ化し、PowerPointで移動・保存・再表示を確認。独立再評価は未実施。'
    for p in [CASE/'README.md', FORK/'SKILL.md', FORK/'README.md']:
        p.write_text(p.read_text(encoding='utf-8').replace(old, new), encoding='utf-8')
    for p in files(CASE):
        target = FORK/p.relative_to(CASE); target.parent.mkdir(parents=True, exist_ok=True); shutil.copy2(p, target)
    for base in [CASE, FORK]:
        source = load(base/'SOURCE.json')
        source['files'] = [{'path': p.relative_to(base).as_posix(), 'sha256': sha(p)} for p in files(base) if p != base/'SOURCE.json']
        dump(base/'SOURCE.json', source)
    for base, name in [(CASE, 'sebastian-193-example-ja.zip'), (FORK, 'sebastian-193-ja.zip')]:
        with zipfile.ZipFile(ROOT/'downloads'/name, 'w', zipfile.ZIP_DEFLATED) as archive:
            for p in files(base): archive.write(p, base.name+'/'+p.relative_to(base).as_posix())
        with zipfile.ZipFile(ROOT/'downloads'/name) as archive: assert archive.testzip() is None
    print('S193: five repaired pages, three unchanged page PNGs/XML, protected parts, original QA0, actual group edit and 2ZIPs synchronized.')

if __name__ == '__main__':
    main()
