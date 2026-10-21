"""Publish the repaired evidence-first research-note example and two packages."""
from pathlib import Path
import hashlib, json, shutil, zipfile
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[1]
CASE=ROOT/'examples/repositories/sebastian-205'
FORK=ROOT/'skill-forks/sebastian-205-ja'
PIN='bb98ce95ca5942fc6551dddc723cbdfe2d637d59'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def dump(p,d):
    p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

SKILL=f'''---
name: paper2ppt-ja
description: 論文・研究ノートの主張を根拠に対応させ、編集可能な日本語の研究発表を作る。
---

# 根拠を保った研究発表を作る

原典 wp-a/paper2ppt / commit {PIN} / MIT。非公式の日本語案内。source/UPSTREAM_SKILL.mdとsource/upstream-referencesを読む。

1. 論文全文・一部・研究ノートのどれがあるかを分類し、目的、聞き手、時間、言語を決める。原典の既定は中国語12〜16枚・15〜20分。依頼の指定を優先する。一般営業資料や画像スライドの復元とは用途を分ける。
2. 原典scripts/preflight.pyで入力・出力先・制作と描画の環境を確認する。入力が欠ける場合は止めて不足を報告する。環境の不足を成功に置き換えない。
3. 作成前にclaim-to-sourceを作る。主張、原資料の節、根拠の内容、支持の程度、留保、使う頁を記す。引用、数値、因果、実験結果を推測で埋めない。ノートだけなら提案・仮説・評価計画と実証結果を分ける。
4. 問い、方法、結果、解釈、限界へ、資料に存在する内容を構成する。1頁1主張を基本に、原典のpaper-typesとevidence-first workflowで図と説明順を決める。結果がなければ結果頁を捏造しない。
5. 利用可能な専用PPTXツールを優先し、無ければ原典の許容するpython-pptx fallbackで作る。文字・表・主要図を編集可能にし、主張の出典と話者ノートを付ける。
6. 原典inspect_pptx.py --require-notesで構造を検査し、内容を台帳と照合する。利用できる描画器で全頁を実表示し、文字、図、出典、重なり、余白を直す。描画器が無ければ原典の状態structurally-verified-with-disclosure、作成不能ならblockedとし、passと混同しない。
7. 既定の納品はpresentation_ja.pptxとnotes/qa_report.md。PDF・PNG・制作一式などの追加は依頼に合わせる。QAには根拠対応、構造、描画、残る限界を分けて記す。

## この8枚の作例

架空・未実施の食品在庫UI研究計画をnotes-onlyとして扱った。BRIEF §1〜8に対応する8主張の台帳を保存し、それを生成器が読んでから全頁の出典・話者ノートを作る。原典preflightとinspectorを実行し、PowerPoint16で全8頁を描画・確認した。提案、仮説、評価計画、限界を区別し、食品ロス削減の実測値や安全保証を加えていない。

8枚・日本語・制作ZIPは今回の依頼による変更。source/build.pyは固定作例の生成器であり、任意論文を自動抽出する製品ではない。別の題材では元資料、主張台帳、構成、生成内容を更新する。配布のdeck.pptxはpresentation_ja.pptxと同じ互換名。独立評価と正式採用は別に記録する。

## 環境と書体

この再現にはPythonとpython-pptx1.0.2以上。原典preflightはPyMuPDF/Pillowも調べる。表示確認にはPowerPointまたはLibreOffice等の描画器が必要。依存を自動導入しない。
書体はNoto Sans CJK JP Regular/Bold。Noto Sans JPとは名前が異なる。公式Noto CJKのOFL配布物を作成者と受け手の端末に導入する。OSからフォントを抽出して再配布しない。

原典: https://github.com/wp-a/paper2ppt/tree/{PIN}
書体: https://github.com/notofonts/noto-cjk/tree/f8d157532fbfaeda587e826d4cd5b21a49186f7c/Sans/OTF/Japanese
'''

def package():
    excluded={'.private','.git','.venv','node_modules','__pycache__'}
    for folder,name in [(CASE,'sebastian-205-example-ja.zip'),(FORK,'sebastian-205-ja.zip')]:
        with zipfile.ZipFile(ROOT/'downloads'/name,'w',zipfile.ZIP_DEFLATED) as z:
            for p in sorted(folder.rglob('*')):
                if p.is_file() and not set(p.relative_to(folder).parts)&excluded and p.suffix!='.pyc':z.write(p,folder.name+'/'+p.relative_to(folder).as_posix())
        with zipfile.ZipFile(ROOT/'downloads'/name) as z:
            assert z.testzip() is None
            assert not any(set(Path(n).parts)&excluded for n in z.namelist())

def main():
    inspect=load(CASE/'source/structural-inspector.json');preflight=load(CASE/'source/preflight-repair.json')
    assert inspect['ok'] and inspect['slides']==8 and inspect['notes_slides']==8 and preflight['result']['ok']
    ledger=load(CASE/'notes/claim-to-source.json');prs=Presentation(CASE/'presentation_ja.pptx')
    assert len(prs.slides)==len(ledger)==8
    assert sha(CASE/'deck.pptx')==sha(CASE/'presentation_ja.pptx')
    for n,(s,c) in enumerate(zip(prs.slides,ledger),1):
        assert f'BRIEF §{n}' in '\n'.join(q.text for q in s.shapes if q.has_text_frame)
        assert c['claim'] in s.notes_slide.notes_text_frame.text
    images=[CASE/f'render/slide-{n:02}.png' for n in range(1,9)];assert all(p.exists() for p in images)
    render=load(CASE/'.private/powerpoint-caption-repair/powerpoint-render.json')
    render['source']='presentation_ja.pptx';render['pdf']='Private rendering intermediate; not distributed';dump(CASE/'source/powerpoint-render.json',render)
    qa={'status':'pass','meaning':'Original technical/content/render QA for this fixed notes-only example; independent grading pending.', 'source_sufficiency':'notes-only','fictional_unperformed':True,'content':'8 distinct claims map to BRIEF §1–8 before rebuilding. All8 notes contain the claim, source and specific caveat. No measured results or citations fabricated.', 'preflight':'source/preflight-repair.json','structural':'source/structural-inspector.json','native_editable':True,'slides':8,'notes':8,'media':0,'visual':{'renderer':'Microsoft PowerPoint16 PDF → PyMuPDF RGB PNG','all_pages_viewed':8,'scope':'All eight actual pages read and checked for source footers, hierarchy, proposal/plan/limitations distinctions and clipping.','receipt':'source/powerpoint-render.json','images':[dict(path=p.relative_to(CASE).as_posix(),sha256=sha(p)) for p in images]},'independent_review':'pending_after_repair','limits':['架空の研究ノートだけを使った未実施計画。効果・食品安全性の実証はない。','固定8頁用の生成コードで、任意論文を自動抽出するアプリではない。','原典の対象は研究発表。一般営業資料、画像スライド復元、論文検索だけの用途は対象外。','原典の既定は簡体中国語12〜16頁・15〜20分。今回は指定に合わせた日本語8頁。','他の描画環境と修復後の独立再評価は未確認。']}
    if (CASE/'source/zip-replay.json').is_file():qa['distribution_replay']='source/zip-replay.json'
    dump(CASE/'QA.json',qa)
    prov=load(CASE/'PROVENANCE.json')
    if 'initial_record' not in prov:prov={'initial_record':prov,'repository':'wp-a/paper2ppt','upstream_commit':PIN,'license':'MIT'}
    prov.update(source_sufficiency='notes-only',repair_author='Codex',repair_scope='Original preflight → source-specific claim ledger → builder consumes ledger → original structural inspector → all8 actual PowerPoint pages.',upstream_scripts_executed=True,renderer='PowerPoint16; PDF is a private rendering intermediate',font='Noto Sans CJK JP installed; no OS font files redistributed',current_limits=qa['limits']);dump(CASE/'PROVENANCE.json',prov)
    report='''# 検査結果

状態: pass（この作例の内容・構造・表示検査）。独立採点は修復後の再評価待ち。

入力はnotes-onlyの架空・未実施研究計画。原典preflight成功後、主張・節・根拠・留保を8頁別に記録して再生成した。生成器は台帳を読み、全頁の話者ノートに主張・節・固有の限界を入れる。提案、仮説、評価計画、限界を分け、実測値や引用を加えていない。

原典inspect_pptx.py --require-notes: ok=true、8頁・8notes・media0・bounds0・warning0・error0。PowerPoint16で現在PPTXを開いて全8頁を描画し、主張、出典、文字・図、重なりと余白を全頁確認した。PNGはその実PDFから作成。描画器がない場合に使うstructurally-verified-with-disclosureとは区別する。

履歴: 初回の作例では原典2検査を使わず最小QAに置換していた。今回は原典検査と全頁表示を追加した修復runであり、初回からこの手順を通したとは主張しない。食品ロス削減や使いやすさの実証、他環境互換、独立再評価、正式採用は未了。
'''
    (CASE/'notes/qa_report.md').write_text(report,encoding='utf-8')
    (CASE/'SKILL.md').write_text(SKILL,encoding='utf-8');(FORK/'SKILL.md').write_text(SKILL,encoding='utf-8')
    readme='''# 根拠から組み立てる研究発表

paper2pptは、論文・研究ノートの主張を根拠と限界に対応させ、編集できるPPTXへ整理するスキルです。この作例は、架空・未実施の食品在庫UI研究計画8頁。実測結果はありません。

## 使い方

presentation_ja.pptxをPowerPointで開きます。deck.pptxは同じ内容の互換名。render/には現在PPTXから実描画した全8頁があります。notes/claim-to-source.jsonで、各頁の主張とBRIEFの節・固有の留保を確認できます。

再現はPythonとpython-pptx1.0.2以上、Noto Sans CJK JPで `python source/build.py`。入力台帳を先に読み、PPTXとinput/data.jsonを生成します。任意論文を自動抽出するコードではないため、自分の資料ではBRIEF・台帳・構成・生成器の本文を変更してください。

原典の検査: `python source/upstream-scripts/preflight.py input/BRIEF.md --output . --json` → `python source/upstream-scripts/inspect_pptx.py presentation_ja.pptx --require-notes --json`。新しく作ったPPTXは全頁を描画し直し、台帳と照合します。元のQAを新出力の合格として引き継ぎません。

## 確認範囲と条件

今回の修復は原典preflight、inspector、8頁の内容対応とPowerPoint16の全頁表示まで確認。独立再評価、他環境の表示、食品ロス削減効果は未確認。初回の最小QAからの変更はPROVENANCE.jsonに保持しています。原典の既定納品PPTX＋qa_reportに、依頼された制作ZIPとPNGを加えました。PDFは表示用の私的中間物で配布しません。

原典 wp-a/paper2ppt、固定commit bb98ce95ca5942fc6551dddc723cbdfe2d637d59、MIT。原典SKILLとreferences、変更していない2QA scriptを同梱。フォントは環境書体で本体を同梱しません。公式取得先はSKILL.mdに記載。
'''
    (CASE/'README.md').write_text(readme,encoding='utf-8')
    prompt='原典paper2pptのpreflight→claim-to-source→構成→native PPTX→inspect_pptx→全頁描画の順で、input/BRIEF.mdの架空・未実施研究計画を日本語8頁へ。各頁の主張・節・留保を対応させる。結果・数値・引用を捏造しない。期限・更新時点・可食判断を分ける。原典が認めるpython-pptx fallbackを使い、全頁に出典と話者ノートを入れる。未実施の独立レビューを合格扱いしない。\n'
    (CASE/'PROMPT.txt').write_text(prompt,encoding='utf-8')
    for folder in [CASE,FORK]:
        stale=folder/'render/deck.pdf'
        if stale.exists():
            dst=folder/'.private/stale-initial-render.pdf';dst.parent.mkdir(parents=True,exist_ok=True)
            if not dst.exists():shutil.move(stale,dst)
    # Synchronize maintained fork, leaving its original identity/licence intact.
    for p in CASE.rglob('*'):
        if p.is_file() and not set(p.relative_to(CASE).parts)&{'.private','__pycache__'}:
            dest=FORK/p.relative_to(CASE);dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,dest)
    source=load(FORK/'SOURCE.json');source['repair']='Current native rebuild reads8source-specific claims before authoring. Original preflight/inspector and all8PowerPoint views checked; independent review pending.';source['files']=[dict(path=p.relative_to(FORK).as_posix(),sha256=sha(p)) for p in sorted(FORK.rglob('*')) if p.is_file() and p.name!='SOURCE.json' and not set(p.relative_to(FORK).parts)&{'.private','__pycache__'}];dump(FORK/'SOURCE.json',source)
    steps=['元資料を全文・一部・ノートに分類し、用途と発表条件を決める。原典preflightで環境を確認する。','主張・資料の節・支持の程度・留保を台帳にしてから、問い・方法・結果・解釈・限界を構成する。','編集可能なPPTXを作り、出典と話者ノートを付ける。原典inspector、根拠照合、全頁描画で直す。']
    rel=lambda p:p.relative_to(ROOT).as_posix()
    patch={'case_id':'sebastian-205','name':'研究資料を、主張と根拠が対応するPPTXへ','upstream_principle':'論文全文・一部・研究ノートを分類し、主張と根拠の台帳を先に作ってから、編集可能な研究発表とQA報告へ仕上げる。','change':'資料のどの節が各主張を支えるか、何が未確認かを残した発表へ整理できる。','input':'論文・研究ノート、目的、聞き手、発表時間と言語','output':'編集可能なPPTXとQA報告。出典・話者ノート・根拠台帳を保持する。','how_to_use':steps,'constraints':qa['limits'],'demonstration':{'summary':'架空の研究計画をnotes-onlyに分類。8頁別の主張・原資料の節・固有の留保を先に記録し、生成器が台帳を読んで出典・話者ノートを作る。原典2検査と全8頁のPowerPoint表示を確認した。','applied_features':['入力資料の充足度分類','制作前のclaim-to-source台帳','主張別の出典・話者ノート','原典preflight/inspectorと全頁表示'],'not_demonstrated':['査読論文の実験結果・効果の発表','任意論文の自動抽出','修復後の独立再評価・他環境の互換性'],'pages':[dict(index=2,label='3ページ目',title='仮説を、実証結果に見せない',skill_rule='資料中の仮説・支持の程度・留保を区別して構成する。',what_to_inspect='場所一覧と粗い残量は未検証の仮説。「設計上のトレードオフ」と明記し、BRIEF §3と固有の留保をnotesの台帳に対応させる。',proof_path=rel(images[2]),proof_kind='render-image'),dict(index=5,label='6ページ目',title='結果がない資料は、評価計画として示す',skill_rule='存在しない結果や数値を作らず、原資料で支持される方法を示す。',what_to_inspect='3課題と観察・質問を表へ対応。「評価結果はまだない」と明示。BRIEF §6と同じ課題・方法を保つ。',proof_path=rel(images[5]),proof_kind='render-image'),dict(index=6,label='7ページ目',title='研究の限界を主張の近くに残す',skill_rule='根拠の範囲と限界を、図・本文・話者ノートで一貫させる。',what_to_inspect='画面と実物のずれを対比し、可食判断を自動化しないと明示。BRIEF §7の未入力・更新忘れ・共有の限界を保持する。',proof_path=rel(images[6]),proof_kind='render-image')],'source_proof_paths':[rel(CASE/p) for p in ['notes/claim-to-source.json','notes/qa_report.md','notes/structural-qa.md','source/build.py','source/build-evidence.json','source/repair-order.json','source/preflight-repair.json','source/structural-inspector.json','source/powerpoint-render.json','source/upstream-script-identity.json','QA.json']]},'guide_patch':{'summary_ja':'研究資料の主張と根拠・留保を先に対応させ、編集可能なPPTXとQA報告を作る。','steps':steps,'limitations':qa['limits'],'requirements':['Python / python-pptx1.0.2+','原典preflight: PyMuPDF/Pillow','描画確認: PowerPointまたはLibreOffice等','Noto Sans CJK JP'],'production_status':'original_preflight_inspector_all8_powerpoint_checked_independent_pending'}}
    p=ROOT/'data/skill-effect-patches/repairs-native-root.json';rows=load(p) if p.exists() else [];dump(p,[r for r in rows if r['case_id']!='sebastian-205']+[patch])
    p=ROOT/'data/sol-sebastian-batch2-repository-examples.json';rows=load(p)
    for c in rows:
        if c['id']=='sebastian-205':c.update(artifact_download=rel(CASE/'presentation_ja.pptx'),preview_kind='powerpoint-pdf-rgb-render',preview_image_count=8,preview_image_notice='原典2検査とPowerPoint16の全8頁表示を確認。架空・未実施研究計画。独立再評価は未了。',sample_notice='架空・未実施。効果の実測値なし。原典2検査と全8頁のPowerPoint表示を確認。',prompt_text=prompt.strip(),execution_status='original_qa_all8_powerpoint_checked')
    dump(p,rows);package();print('Paper2ppt: evidence-first native rebuild,8source-specific claims, original QA, all8actual pages and2ZIPs synchronized.')
if __name__=='__main__':main()
