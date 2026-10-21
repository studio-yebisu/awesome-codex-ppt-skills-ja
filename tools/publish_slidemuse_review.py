"""Publish genuine image-first workflow and original validators/packer/render proof."""
from pathlib import Path
import hashlib,json,shutil,zipfile
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[1];CASE=ROOT/'examples/repositories/sebastian-192';FORK=ROOT/'skill-forks/sebastian-192-ja'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def load(p):return json.loads(p.read_text(encoding='utf-8-sig'))
def dump(p,d):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
SKILL='''---
name: slidemuse-ja
description: 本・論文・報告・原稿から、四つのスタイルを比較して画像主体のスライドを作る。明示依頼時は画像版から編集可能なPPTXへ再構築する。
---

# 内容と見た目を決めてから、ページを画像で作る

原典 helloo1568/slidemuse / 7034412b4c19915c17f8d7206e3723e9e53cf8d2 / MIT。非公式の日本語案内。source/upstream/SKILL.mdと参照文書を読む。

1. 原資料、参考方向、用途と受け手、頁数、納品範囲を確認する。既存の指定・委任を引き継ぎ、不足だけを確認する。任意の編集済みPPTXを直接修正する用途とは分ける。
2. 原資料を読み、全体の答えと頁ごとの主張・根拠・説明・限界・話者ノートを含む内容稿を作る。内容自検を行い、内容を確認するか、既存の代定権限を記録する。
3. 同じ内容と頁順の4スタイルを、独立した4回の宿主imagegenで総覧にする。8頁以下は全頁、それより多い場合は6〜8代表頁を同じ組で示す。総覧は横長、縮略頁も指定比率・等大・余白・外側頁番号を保つ。validate_style_options.pyと目視を行い、選択または委任による代選を記録する。明示的な固定参照・省略は原典の条件に従う。
4. 確認済み文字・データ・出典・要素意図・安定ID・書体/色/余白をPage Specへ固定し、validate_page_spec.py --strictを実行する。画像版を飛ばして編集用PPTXへ直行しない。
5. 頁別imagegenで画像を作り、直後に文字・数字・図の意味・欠け・指定比率・全体の統一を確認する。誤りは該当頁だけ直す。原典の内容監査と画像/台帳のhash対応を使い、機械チェックを目視の代わりにしない。
6. 全頁が確認できたら、validate_page_spec.py --require-images --strict → build_image_ppt.py → render_deck.py。画像版PPTXを実描画して全頁を見比べる。画像の文字を個別編集できるとは説明しない。生図能力がない場合はその段階で止める。
7. 明示的に編集版が必要な場合だけS6へ進む。既知の文字/数値は再OCRせず、画像を層分けしScene v1へ記す。文字・単純図形・表・可靠データ図をnative、複雑な絵を分離画像にする。全面画像を編集用の底絵として残さない。
8. 編集版はbuild_editable_ppt.pyとaudit_editability.py、全頁の原図比較、文字変更・図表編集・移動で確認する。リリースの完全合格はevaluate_delivery.pyの現出力に紐づく記録で判断する。未実施を合格と記載しない。

## この8頁の作例

架空・未実施の日陰と動線の研究計画。内容稿→4独立総覧→委任による案3選択→Page Spec→8画像の制作を実行した。完成頁の画像は16:9の丸め1px内。画像本体の文字・図はラスタ、話者ノートは編集可能。宿主のモデル名は非露出で、特定のImage2.5子モデルを指定したとは記さない。

今回の修復ではPage Specを原典schemaへ整え、原典の総覧検査、画像付きstrict検査、画像PPTX packer、PowerPoint render_deckを実行し、全8頁を確認した。画像は再生成・再圧縮せず、承認済みPNGをそのまま埋めた。総覧の内部縮略頁は指定16:9に未適合で、外枠の構造検査合格と全視覚条件の合格は分ける。S6編集版と完全なevaluate_deliveryは未実施。原稿を変える場合は内容稿からやり直し、旧画像を新しい原稿の出力として扱わない。

## 再現

Python、python-pptx、Pillow、jsonschema。PowerPoint描画にはWindowsとMicrosoft PowerPoint。新しい画像を作るには宿主imagegenが必要で、原典script自体は生図/OCR/課金APIを提供しない。`python source/build_images.py reproduced.pptx` は承認済み8画像を原典validator/packerで再包装する。画像を作り直す処理ではない。

原典: https://github.com/helloo1568/slidemuse/tree/7034412b4c19915c17f8d7206e3723e9e53cf8d2
'''

def package():
 for folder,name in [(CASE,'sebastian-192-example-ja.zip'),(FORK,'sebastian-192-ja.zip')]:
  with zipfile.ZipFile(ROOT/'downloads'/name,'w',zipfile.ZIP_DEFLATED) as z:
   for p in sorted(folder.rglob('*')):
    if p.is_file() and not set(p.relative_to(folder).parts)&{'.private','__pycache__','node_modules','.git','.venv'} and p.suffix!='.pyc':z.write(p,folder.name+'/'+p.relative_to(folder).as_posix())
  with zipfile.ZipFile(ROOT/'downloads'/name) as z:assert z.testzip() is None

def main():
 gates=load(CASE/'source/original-gates.json');assert len(gates)==3 and all(r['exit_code']==0 for r in gates)
 render=load(CASE/'reviews/original-render/render-report.json');assert render['backend']=='powerpoint' and len(render['slides'])==8
 spec=load(CASE/'page-spec.json');prs=Presentation(CASE/'deck.pptx');assert len(prs.slides)==8
 with zipfile.ZipFile(CASE/'deck.pptx') as z:
  media=[z.read(n) for n in z.namelist() if n.startswith('ppt/media/')]
  assert len(media)==8
  for s,row in zip(prs.slides,spec['slides']):assert len(s.shapes)==1 and (CASE/row['image_file']).read_bytes() in media and s.notes_slide.notes_text_frame.text==row['speaker_notes']
 old=load(CASE/'QA.json');qa={'status':'original_validators_packer_powerpoint_all8_checked','initial_record':old.get('initial_record',old),'original_gates':'source/original-gates.json','original_render':'reviews/original-render/render-report.json','page_count':8,'image_only_slide_objects':True,'speaker_notes_exact_match':True,'embedded_media_byte_match':True,'all8_actual_powerpoint_pages_viewed':True,'all4_style_overviews_viewed':True,'image_generation_calls_added':0,'independent_review':'pending_after_repair','limits':['総覧内部の縮略頁は指定16:9に未適合。外枠の構造検査と全視覚条件の適合を分ける。','S6のnative再構築は未実施。文字・図は画像、話者ノートのみ編集可能。','完全なevaluate_delivery、厳密なOCR文字照合、独立再評価は未実施。','原典は宿主imagegenが必要。修復コマンドは既存8画像の再包装で、新画像を生成しない。']}
 if (CASE/'source/zip-replay.json').exists():qa['distribution_replay']='source/zip-replay.json'
 dump(CASE/'QA.json',qa)
 prov=load(CASE/'PROVENANCE.json')
 if 'initial_record' not in prov:prov={'initial_record':prov,'repository':'helloo1568/slidemuse','upstream_commit':'7034412b4c19915c17f8d7206e3723e9e53cf8d2','license':'MIT'}
 prov.update(repair_author='Codex',current_route='Original strict Page Spec validator and image PPTX packer; unchanged approved8PNGs; original PowerPoint render_deck; all8pages visually checked.',style_selection='Delegated option3 from4 independent historical imagegen overviews; no new generation.',limits=qa['limits']);dump(CASE/'PROVENANCE.json',prov)
 (CASE/'source/build_images.py').write_text('''"""Repackage approved images with unchanged original strict validator/packer."""
from pathlib import Path
import subprocess,sys
p=Path(__file__).resolve().parents[1]
output=Path(sys.argv[1]).resolve() if len(sys.argv)>1 else p/'reproduced.pptx'
scripts=p/'source/upstream/scripts'
subprocess.run([sys.executable,str(scripts/'validate_page_spec.py'),str(p/'page-spec.json'),'--require-images','--strict'],check=True)
subprocess.run([sys.executable,str(scripts/'build_image_ppt.py'),str(p/'page-spec.json'),str(output),'--fit','contain','--max-width','0','--jpeg-quality','0'],check=True)
''',encoding='utf-8')
 (CASE/'SKILL.md').write_text(SKILL,encoding='utf-8')
 (CASE/'README.md').write_text('''# 内容と見た目を決めて画像スライドを作る

SlideMuseは、内容稿から同じ内容の4スタイルを比較し、選んだ見た目で頁別画像を作るスキルです。明示依頼があれば、その画像版から編集可能なPPTXへ再構築します。

この8頁は、架空・未実施の日陰と動線の研究計画。source/styles/の4独立総覧から案3を委任選択し、Page Specと頁別生成で作りました。文字・図は画像、話者ノートのみ編集可能です。S6のnative再構築はしていません。

現在PPTXは原典strict validatorとbuild_image_ppt.pyで承認済みPNGを無加工で再包装し、原典render_deck.pyのPowerPoint backendで全8頁を実描画して見比べました。render/slide-NN.pngはPowerPoint実PDFのRGB描画、pages/は元の生成PNGです。reviews/original-renderに原典の比較記録があります。

再現はPython＋python-pptx/Pillow/jsonschemaで `python source/build_images.py reproduced.pptx`。新画像の生成には宿主imagegenが必要です。原稿を変える場合は内容稿とスタイルから作り直してください。このコマンドは既存画像の再包装です。自分の原稿から始める5段の手順と総覧依頼例はsource/REUSE.mdへ。

総覧の外枠と候補数・hashは原典検査を通りましたが、内部縮略頁の16:9は未適合です。完全なevaluate_delivery、S6、厳密OCR照合、独立再評価は未実施。初回の独自包装・未実描画からの変更はPROVENANCEに保持しています。原典MITとモデル名非露出を明記し、特定の生図モデルを使ったとは説明しません。
''',encoding='utf-8')
 s=(CASE/'pages/index.html').read_text(encoding='utf-8');s=s.replace('プレビューはPPTX内の画像とバイト一致した生成PNGで、PPTXレンダリングではありません。','ここはPPTX内と同一の生成PNG。修復後のPPTX全8頁は別途PowerPointで実描画・確認済みです。');(CASE/'pages/index.html').write_text(s,encoding='utf-8')
 for p in CASE.rglob('*'):
  if p.is_file() and not set(p.relative_to(CASE).parts)&{'.private','__pycache__'}:
   q=FORK/p.relative_to(CASE);q.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,q)
 source=load(FORK/'SOURCE.json');source['repair_scope']='Original validators/packer/render_deck used; all8actual PowerPoint views, no new image generation.';source['files']=[dict(path=p.relative_to(FORK).as_posix(),sha256=sha(p)) for p in sorted(FORK.rglob('*')) if p.is_file() and p.name!='SOURCE.json' and not set(p.relative_to(FORK).parts)&{'.private','__pycache__'}];dump(FORK/'SOURCE.json',source)
 patch=next(r for r in load(ROOT/'data/skill-effect-patches/native-03.json') if r['case_id']=='sebastian-192');patch=json.loads(json.dumps(patch))
 patch.update(name='同じ内容で4スタイルを比べ、画像スライドを作る',change='内容を先に固め、四つの見た目から選んで、頁ごとの画像へ仕上げられる。',output='画像主体のPPTX。明示依頼時は、画像版の後に編集可能PPTXへ再構築する。',constraints=qa['limits'])
 patch['demonstration']['summary']='同じ8頁内容の4独立総覧から案3を委任選択し、Page Specを固定して頁別に生成。修復では原典strict検査・画像packer・PowerPoint rendererを実行し、8画像の埋込と全8頁の表示を確認した。'
 patch['demonstration']['not_demonstrated']=['S6の編集可能native再構築','総覧内部の16:9適合','完全なevaluate_delivery','独立再評価']
 patch['demonstration']['source_proof_paths']+=['examples/repositories/sebastian-192/'+f for f in ['page-spec.json','source/original-gates.json','source/original-image-receipts.json','reviews/original-render/render-report.json',*[f'source/styles/{n:02}-sorter.png' for n in range(1,5)]]]
 patch['guide_patch'].update(summary_ja=patch['change'],limitations=qa['limits'],production_status='original_strict_packer_powerpoint_all8_checked_independent_pending')
 p=ROOT/'data/skill-effect-patches/repairs-native-root.json';rows=load(p);dump(p,[r for r in rows if r['case_id']!='sebastian-192']+[patch])
 p=ROOT/'data/sol-sebastian-batch2-images-repository-examples.json';rows=load(p)
 for r in rows:
  if r['id']=='sebastian-192':r.update(input_data='examples/repositories/sebastian-192/page-spec.json',preview_kind='generated_png_plus_all8_powerpoint_checked',preview_image_notice='元の生成PNG。現在PPTXを原典PowerPoint rendererで全8頁確認済み。文字・図は画像です。',sample_notice='架空・未実施。4スタイルから画像主体8頁を制作。総覧内部16:9とnative編集版は未達。',method='Content outline →4independent style overviews→delegated option3→page spec→8AI PNGs; repaired with original strict validator/image packer/PowerPoint render_deck.',execution_status='original_strict_packer_powerpoint_all8_checked')
 dump(p,rows);package();print('SlideMuse: original strict validators/packer/PowerPoint renderer,8actual pages,4style sources and2ZIPs synchronized.')
if __name__=='__main__':main()
